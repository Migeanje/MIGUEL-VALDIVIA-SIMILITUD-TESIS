"""Manual verification of the objectives locator (O05): the sample, its workbook, the import.

The acceptance criterion O05 asks for at least 90% correct general objectives on a manual,
stratified sample of about 60 theses. This module:

1. draws the sample (:func:`make_sample`): ``per_program`` theses of each program, from the
   theses with a valid own PDF. The wrong-file item (:data:`WRONG_FILE_HANDLES`) is left
   out, and of each exact-duplicate pair (D24: same title and abstract) only the record
   with the lexicographically smaller uuid is kept, a placeholder until the fichas (T12)
   fix the canonical record. Within a program, ``not_found`` documents enter in proportion
   to their share, so the accuracy covers the whole pipeline. Each (program, status)
   stratum is drawn with its own seed, derived from the configured one, and the rows are
   then shuffled, so the workbook does not group them;
2. writes the workbook (:func:`make_verification_workbook`): a guide sheet and a check
   sheet in Spanish, one row per thesis, with links to the repository item and to the
   local PDF, the extracted texts, and two drop-down verdict columns. Person names of the
   snapshot are replaced by ``[nombre omitido]``, and no cell is ever a formula. Beside it,
   ``sample.json`` records the draw, without any text;
3. reads the filled workbook back (:func:`import_verification`), refusing it with every
   problem listed when a verdict is blank or outside its domain, or a row was lost;
4. computes the accuracy (:func:`verification_report`): ``Correcto`` over the judged rows,
   with a Wilson 95% interval, overall and per program. ``Parcial`` counts as not correct
   (strict O05 rule). The report holds numbers only.

The workbook and ``sample.json`` hold thesis text and live in
``objectives.verification_sample.workbook_dir/<snapshot_id>/``, inside the gitignored
``data/``. Once the workbook is filled, import it::

    uv run python -m thematic_redundancy.labeling.objectives_check [--project-root PATH]
        [--snapshot-id ID] [--report-out results/objectives/<snapshot_id>/verification.json]
"""

import argparse
import hashlib
import math
import os
import random
import re
import sys
import zipfile
from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal
from uuid import UUID

import openpyxl
import yaml
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, JsonValue, ValidationError

from thematic_redundancy.corpus.pdf_manifest import PDF_DIR_NAME, pdf_file_name
from thematic_redundancy.corpus.profile import ABSTRACT_KEY, PERSON_KEYS, TITLE_KEY
from thematic_redundancy.corpus.snapshot import (
    RAW_DIR_NAME,
    Count,
    SnapshotRecord,
    UtcDatetime,
    check_snapshot_id,
    latest_snapshot_id,
)
from thematic_redundancy.extraction.objectives import LocatorStatus, ObjectivesRecord
from thematic_redundancy.extraction.text_manifest import write_json_atomically
from thematic_redundancy.preprocessing.light_cleaner import normalize_text
from thematic_redundancy.shared.config import AppConfig, load_config

CONFIG_PATH = Path("config") / "default.yaml"
"""Configuration that the command line loads, relative to the project root."""

WORKBOOK_FILE_NAME = "objectives_check.xlsx"
SAMPLE_FILE_NAME = "sample.json"

WRONG_FILE_HANDLES = frozenset({"20.500.12920/11777"})
"""Items whose PDF in the repository is another thesis's file (data card, section 7). Their
objectives fall back to title and abstract, so they stay out of the sample."""

TARGET_ACCURACY = 0.9
"""O05: the share of correct general objectives that the locator must reach."""

Z_95 = 1.959963984540054
"""Two-sided 95% quantile of the standard normal distribution."""

GeneralVerdict = Literal["correct", "partial", "incorrect", "missed", "absent"]
SpecificVerdict = Literal["correct", "partial", "incorrect", "not_applicable"]

GENERAL_VERDICTS: dict[str, GeneralVerdict] = {
    "Correcto": "correct",
    "Parcial": "partial",
    "Incorrecto": "incorrect",
    "No encontrado (existe)": "missed",
    "No existe en la tesis": "absent",
}
"""Labels of the general-objective verdict, as the workbook offers them, and their tokens."""

SPECIFIC_VERDICTS: dict[str, SpecificVerdict] = {
    "Correcto": "correct",
    "Parcial": "partial",
    "Incorrecto": "incorrect",
    "No aplica": "not_applicable",
}
"""Labels of the specific-objectives verdict, as the workbook offers them, and their tokens."""

_ALLOWED_GENERAL: dict[LocatorStatus, frozenset[GeneralVerdict]] = {
    "extracted": frozenset({"correct", "partial", "incorrect", "absent"}),
    "not_found": frozenset({"missed", "absent"}),
}

STATUS_LABELS: dict[LocatorStatus, str] = {"extracted": "extraído", "not_found": "no encontrado"}

COLUMNS = (
    "N.º",
    "Programa",
    "Handle",
    "PDF local",
    "Páginas del PDF",
    "Estado",
    "Objetivo general extraído",
    "Objetivos específicos extraídos",
    "Veredicto objetivo general",
    "Veredicto objetivos específicos",
    "Notas",
)
"""Header of the check sheet; the import refuses a workbook whose header changed."""

_WIDTHS = (6, 24, 20, 11, 10, 14, 60, 70, 24, 24, 40)
_WRAPPED = frozenset({"Programa", "Objetivo general extraído", "Objetivos específicos extraídos"})
_CHECK_SHEET = "Verificación"
_GUIDE_SHEET = "Guía"
_REDACTED = "[nombre omitido]"
_MINUTES_PER_ROW = (4, 6)
_STATUSES: tuple[LocatorStatus, ...] = ("extracted", "not_found")


class WorkbookValidationError(ValueError):
    """The filled workbook cannot be imported; ``problems`` lists every reason."""

    def __init__(self, problems: Sequence[str]) -> None:
        self.problems = tuple(problems)
        listed = "\n".join(f"- {problem}" for problem in self.problems)
        super().__init__(f"{len(self.problems)} problem(s) in the workbook:\n{listed}")


class _Model(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid", hide_input_in_errors=True)


class SampleItem(_Model):
    """One sampled thesis, in workbook order."""

    n: Count
    item_uuid: UUID
    handle: str
    program: str
    status: LocatorStatus


class VerificationSample(_Model):
    """The draw: its settings, the eligible population, the exclusions and the items.

    Saved as ``sample.json`` beside the workbook; it holds no thesis text.
    """

    snapshot_id: str | None = None
    created_at: UtcDatetime | None = None
    provenance: dict[str, JsonValue] = {}
    """Locator version, settings fingerprint and table digest the sample was drawn from."""
    seed: int
    per_program: Count
    population: dict[str, dict[str, Count]]
    """Eligible theses per program and status."""
    excluded_wrong_file: tuple[str, ...]
    """Handles left out because their PDF is another thesis's file."""
    excluded_duplicates: tuple[str, ...]
    """Handles left out as the larger-uuid record of an exact-duplicate pair (D24)."""
    items: tuple[SampleItem, ...]


@dataclass(frozen=True)
class VerificationWorkbook:
    """What :func:`make_verification_workbook` wrote."""

    workbook_path: Path
    sample_path: Path
    sample: VerificationSample
    redacted: int
    """Cells in which a person's name was replaced."""


@dataclass(frozen=True)
class Judgement:
    """The verdicts of one row of a filled workbook."""

    n: int
    handle: str
    program: str
    status: LocatorStatus
    general: GeneralVerdict
    specific: SpecificVerdict


# Sampling


def make_sample(
    objectives: Sequence[ObjectivesRecord],
    snapshot_records: Sequence[SnapshotRecord],
    *,
    programs: Sequence[str],
    per_program: int,
    seed: int,
) -> VerificationSample:
    """Draw the verification sample from the objectives table; see the module docstring.

    ``programs`` gives the program order; programs with no eligible thesis are left out.
    """
    eligible, wrong_file, duplicates = _eligible(objectives, snapshot_records)
    population: dict[str, dict[str, int]] = {}
    drawn: list[ObjectivesRecord] = []
    for program in programs:
        strata = {
            status: sorted(
                (r for r in eligible if r.program == program and r.status == status),
                key=lambda record: str(record.item_uuid),
            )
            for status in _STATUSES
        }
        sizes = {status: len(records) for status, records in strata.items()}
        if not sum(sizes.values()):
            continue
        population[program] = sizes
        for status, quota in _allocate(sizes, per_program).items():
            stratum_rng = random.Random(f"{seed}:{program}:{status}")
            drawn.extend(stratum_rng.sample(strata[status], quota))
    random.Random(f"{seed}:order").shuffle(drawn)
    return VerificationSample(
        seed=seed,
        per_program=per_program,
        population=population,
        excluded_wrong_file=wrong_file,
        excluded_duplicates=duplicates,
        items=tuple(
            SampleItem(
                n=n,
                item_uuid=record.item_uuid,
                handle=record.handle or "",
                program=record.program,
                status=record.status,
            )
            for n, record in enumerate(drawn, start=1)
        ),
    )


def _eligible(
    objectives: Sequence[ObjectivesRecord], snapshot_records: Sequence[SnapshotRecord]
) -> tuple[list[ObjectivesRecord], tuple[str, ...], tuple[str, ...]]:
    """Return the eligible records, then the handles of the wrong-file and duplicate records
    left out."""
    keys = {record.uuid: _duplicate_key(record) for record in snapshot_records}
    wrong_file = tuple(sorted(r.handle for r in objectives if r.handle in WRONG_FILE_HANDLES))
    candidates = [r for r in objectives if r.handle not in WRONG_FILE_HANDLES]
    groups: dict[tuple[str, str], list[ObjectivesRecord]] = {}
    for record in candidates:
        key = keys.get(record.item_uuid)
        if key is not None and all(key):
            groups.setdefault(key, []).append(record)
    dropped = {
        record.item_uuid
        for group in groups.values()
        for record in sorted(group, key=lambda record: str(record.item_uuid))[1:]
    }
    duplicates = tuple(
        sorted(
            record.handle or str(record.item_uuid)
            for record in candidates
            if record.item_uuid in dropped
        )
    )
    return [r for r in candidates if r.item_uuid not in dropped], wrong_file, duplicates


def _duplicate_key(record: SnapshotRecord) -> tuple[str, str]:
    """Title and abstract with whitespace collapsed and case folded, accents kept (D24)."""
    return (
        normalize_text(_first_value(record, TITLE_KEY)).casefold(),
        normalize_text(_first_value(record, ABSTRACT_KEY)).casefold(),
    )


def _first_value(record: SnapshotRecord, key: str) -> str:
    for entry in record.metadata.get(key, ()):
        value = entry.get("value")
        if isinstance(value, str) and value.strip():
            return value
    return ""


def _allocate(sizes: Mapping[LocatorStatus, int], per_program: int) -> dict[LocatorStatus, int]:
    """Split ``per_program`` rows over the statuses in proportion to their sizes, by largest
    remainder (ties go to the earlier status); a smaller program gives all its theses."""
    total = sum(sizes.values())
    if total <= per_program:
        return dict(sizes)
    exact = {status: per_program * size / total for status, size in sizes.items()}
    quotas = {status: math.floor(share) for status, share in exact.items()}
    by_remainder = sorted(sizes, key=lambda status: -(exact[status] - quotas[status]))
    for status in by_remainder[: per_program - sum(quotas.values())]:
        quotas[status] += 1
    return quotas


# Workbook


def make_verification_workbook(
    config: AppConfig,
    project_root: Path,
    snapshot_id: str,
    objectives: Sequence[ObjectivesRecord],
    snapshot_records: Sequence[SnapshotRecord],
    *,
    provenance: Mapping[str, JsonValue],
    clock: Callable[[], datetime] = lambda: datetime.now(UTC).replace(microsecond=0),
) -> VerificationWorkbook:
    """Draw the sample and write the workbook and ``sample.json`` of ``snapshot_id``.

    Raises:
        FileExistsError: if the workbook already exists; it may hold verdicts, so it is
            never overwritten.
        ValueError: if a configured location resolves outside the project root.
    """
    settings = config.objectives.verification_sample
    directory = workbook_directory(config, project_root, snapshot_id)
    workbook_path = directory / WORKBOOK_FILE_NAME
    if workbook_path.exists():
        raise FileExistsError(
            f"the verification workbook {workbook_path} already exists and may hold verdicts; "
            "move it away first to draw a new one"
        )
    sample = make_sample(
        objectives,
        snapshot_records,
        programs=[program.key for program in config.snapshot.programs],
        per_program=settings.per_program,
        seed=settings.seed,
    ).model_copy(
        update={"snapshot_id": snapshot_id, "created_at": clock(), "provenance": dict(provenance)}
    )
    data_dir = config.paths.resolve_against(project_root).data_dir
    pdf_dir = data_dir / RAW_DIR_NAME / snapshot_id / PDF_DIR_NAME
    by_item = {record.item_uuid: record for record in objectives}
    redactor = _NameRedactor(_person_names(snapshot_records))
    workbook = _build_workbook(
        sample,
        by_item,
        program_names={program.key: program.name for program in config.snapshot.programs},
        handle_url=f"{str(config.repository.base_url).rstrip('/')}/handle/",
        pdf_link=lambda item: Path(
            os.path.relpath(pdf_dir / pdf_file_name(item), directory)
        ).as_posix(),
        redactor=redactor,
    )
    directory.mkdir(parents=True, exist_ok=True)
    sample_path = directory / SAMPLE_FILE_NAME
    write_json_atomically(sample_path, sample.model_dump_json(indent=2))
    _save_atomically(workbook, workbook_path)
    return VerificationWorkbook(workbook_path, sample_path, sample, redactor.redacted)


def workbook_directory(config: AppConfig, project_root: Path, snapshot_id: str) -> Path:
    """Return the directory that holds the workbook and ``sample.json`` of a snapshot."""
    workbook_dir = config.objectives.verification_sample.resolve_workbook_dir(
        project_root, config.paths.data_dir
    )
    return workbook_dir / check_snapshot_id(snapshot_id)


_LETTER = r"[^\W\d_¹²³⁰⁴-⁹]"
"""A letter, accented letters and ñ included. ``\\w`` also matches digits, ``_`` and superscript
digits such as ``²``; this class leaves them out."""


class _NameRedactor:
    """Replaces every form of the given person names in a text, ignoring letter case.

    A form is replaced only where it stands as whole words: a letter right before or after it
    means it is part of a longer word. A digit, such as a footnote number, or punctuation
    does not, so the name is still replaced.
    """

    def __init__(self, names: Iterable[str]) -> None:
        forms = sorted({form for form in names if form}, key=len, reverse=True)
        alternatives = "|".join(r"\s+".join(map(re.escape, form.split())) for form in forms)
        self._pattern = (
            re.compile(rf"(?<!{_LETTER})(?:{alternatives})(?!{_LETTER})", re.IGNORECASE)
            if forms
            else None
        )
        self.redacted = 0

    def __call__(self, text: str) -> str:
        if self._pattern is None:
            return text
        redacted, count = self._pattern.subn(_REDACTED, text)
        self.redacted += bool(count)
        return redacted


def _person_names(records: Iterable[SnapshotRecord]) -> set[str]:
    """Return the ways in which each author, advisor and juror name may be written: as in
    the metadata (``Family, Given``), as ``Given Family``, and its family part when that has
    two words or more."""
    forms: set[str] = set()
    for record in records:
        for key in PERSON_KEYS:
            for entry in record.metadata.get(key, ()):
                name = entry.get("value")
                if not isinstance(name, str) or not name.strip():
                    continue
                flat = normalize_text(name)
                family, comma, given = (part.strip() for part in flat.partition(","))
                forms.add(flat)
                if comma and given and family:
                    forms.add(f"{given} {family}")
                if comma and len(family.split()) >= 2:
                    forms.add(family)
    return forms


def _build_workbook(
    sample: VerificationSample,
    by_item: Mapping[UUID, ObjectivesRecord],
    *,
    program_names: Mapping[str, str],
    handle_url: str,
    pdf_link: Callable[[UUID], str],
    redactor: Callable[[str], str],
) -> openpyxl.Workbook:
    workbook = openpyxl.Workbook()
    guide = workbook.active
    assert guide is not None
    guide.title = _GUIDE_SHEET
    _write_guide(guide, sample)
    sheet = workbook.create_sheet(_CHECK_SHEET)
    for column, (name, width) in enumerate(zip(COLUMNS, _WIDTHS, strict=True), start=1):
        cell = sheet.cell(1, column, name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[cell.column_letter].width = width
    verdict_fill = PatternFill("solid", fgColor="FFF6D5")
    for row, item in enumerate(sample.items, start=2):
        record = by_item[item.item_uuid]
        values = {
            "N.º": item.n,
            "Programa": program_names.get(item.program, item.program),
            "Handle": item.handle,
            "PDF local": "Abrir PDF",
            "Páginas del PDF": _page_range(record),
            "Estado": STATUS_LABELS[item.status],
            "Objetivo general extraído": redactor(record.objective_general or ""),
            "Objetivos específicos extraídos": redactor(record.objectives_specific or ""),
        }
        for column, name in enumerate(COLUMNS, start=1):
            cell = _text_cell(sheet, row, column, values.get(name))
            cell.alignment = Alignment(wrap_text=name in _WRAPPED, vertical="top")
            if name.startswith("Veredicto"):
                cell.fill = verdict_fill
        _link(sheet.cell(row, COLUMNS.index("Handle") + 1), f"{handle_url}{item.handle}")
        _link(sheet.cell(row, COLUMNS.index("PDF local") + 1), pdf_link(item.item_uuid))
    last = len(sample.items) + 1
    for name, labels in (
        ("Veredicto objetivo general", GENERAL_VERDICTS),
        ("Veredicto objetivos específicos", SPECIFIC_VERDICTS),
    ):
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(labels) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Valor no permitido",
            error="Elija un valor de la lista.",
        )
        letter = sheet.cell(1, COLUMNS.index(name) + 1).column_letter
        validation.add(f"{letter}2:{letter}{last}")
        sheet.add_data_validation(validation)
    sheet.freeze_panes = "A2"
    workbook.active = 1
    return workbook


def _text_cell(sheet: Worksheet, row: int, column: int, value: object) -> Cell:
    """Write ``value``; text is always stored as text, never as a formula."""
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value  # type: ignore[assignment]
    return cell


def _link(cell: Cell, target: str) -> None:
    cell.hyperlink = target
    cell.font = Font(color="0563C1", underline="single")


def _page_range(record: ObjectivesRecord) -> str:
    """Return the pages of the objectives, counted from 1 as PDF viewers do."""
    if record.page_start is None or record.page_end is None:
        return "—"
    start, end = record.page_start + 1, record.page_end + 1
    return str(start) if start == end else f"{start}–{end}"


def _write_guide(sheet: Worksheet, sample: VerificationSample) -> None:
    rows = len(sample.items)
    low, high = (rows * minutes / 60 for minutes in _MINUTES_PER_ROW)
    lines: list[tuple[str, str]] = [
        ("Verificación manual de los objetivos extraídos", ""),
        (
            "Propósito",
            f"Estimar con qué exactitud el sistema extrae el objetivo general de cada tesis. "
            f"Esta muestra tiene {rows} tesis ({sample.per_program} por programa, elegidas al "
            "azar con una semilla fija). Si al menos el 90 % de los objetivos generales es "
            "«Correcto», la extracción automática se usa; si no, se reemplaza por el título y "
            "el resumen.",
        ),
        ("Cómo trabajar", ""),
        ("1.", "Vaya a la hoja «Verificación». Cada fila es una tesis."),
        (
            "2.",
            "Abra la tesis con el enlace «Abrir PDF» de la columna «PDF local» (o con el "
            "enlace del repositorio en «Handle») y vaya a las páginas de «Páginas del PDF». "
            "Son páginas del archivo PDF contadas desde 1, no los números impresos en la tesis.",
        ),
        (
            "3.",
            "Compare el «Objetivo general extraído» con el objetivo general que declara la "
            "tesis. Si el estado es «no encontrado», busque usted los objetivos en la tesis; "
            "suelen estar en el capítulo I, en el planteamiento del problema.",
        ),
        (
            "4.",
            "Elija un valor en «Veredicto objetivo general» y otro en «Veredicto objetivos "
            "específicos» con las listas desplegables. Use «Notas» para cualquier comentario.",
        ),
        (
            "5.",
            "Guarde el archivo con el mismo nombre y en la misma carpeta, para que los enlaces "
            "a los PDF funcionen. Puede trabajar por partes: lo ya juzgado se conserva.",
        ),
        ("Veredicto del objetivo general", ""),
        (
            "Correcto",
            "El texto extraído es el objetivo general de la tesis, completo. Se admite un ruido "
            "menor en los bordes, como una palabra del título de la sección.",
        ),
        (
            "Parcial",
            "El texto está truncado o incluye texto ajeno al objetivo. Cuenta como no correcto.",
        ),
        ("Incorrecto", "El texto es otro pasaje, no el objetivo general."),
        (
            "No encontrado (existe)",
            "El estado es «no encontrado», pero la tesis sí declara un objetivo general.",
        ),
        ("No existe en la tesis", "La tesis realmente no declara objetivos."),
        ("Veredicto de los objetivos específicos", ""),
        ("Correcto", "Los objetivos específicos extraídos son los de la tesis, completos."),
        ("Parcial", "Falta alguno, alguno está truncado, o se incluye texto ajeno."),
        (
            "Incorrecto",
            "El texto es otro pasaje, o está vacío aunque la tesis sí declara objetivos "
            "específicos.",
        ),
        (
            "No aplica",
            "La tesis no declara objetivos específicos, o el estado es «no encontrado».",
        ),
        (
            "Tiempo estimado",
            f"De {_MINUTES_PER_ROW[0]} a {_MINUTES_PER_ROW[1]} minutos por fila, unas "
            f"{low:.0f} a {high:.0f} horas en total; puede hacerse en varias sesiones. Anote "
            "el inicio y el fin de cada sesión: ese tiempo se registra como costo del proceso.",
        ),
        (
            "No modifique",
            "Las columnas «N.º», «Programa», «Handle» y «Estado», ni el orden ni los nombres de "
            "las columnas: la importación los usa para validar el libro.",
        ),
        (
            "Privacidad",
            "El libro contiene fragmentos de tesis con licencia CC BY-NC-ND 4.0, solo para esta "
            "investigación: no lo comparta ni lo suba a otros servicios. Los nombres de "
            f"personas aparecen como «{_REDACTED}».",
        ),
    ]
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 100
    sections = {"Cómo trabajar", "Veredicto del objetivo general"}
    sections.add("Veredicto de los objetivos específicos")
    for row, (label, text) in enumerate(lines, start=1):
        label_cell = _text_cell(sheet, row, 1, label)
        text_cell = _text_cell(sheet, row, 2, text or None)
        label_cell.alignment = Alignment(wrap_text=True, vertical="top")
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row == 1:
            label_cell.font = Font(bold=True, size=14)
        elif label in sections or (text and not label.endswith(".")):
            label_cell.font = Font(bold=True)


def _save_atomically(workbook: openpyxl.Workbook, path: Path) -> None:
    """Save ``workbook`` to a temporary file beside ``path``, then move it onto ``path``."""
    staged = path.with_name(f".{path.name}.part")
    try:
        workbook.save(staged)
        os.replace(staged, path)
    except BaseException:
        staged.unlink(missing_ok=True)
        raise


# Import


def import_verification(path: Path, sample: VerificationSample) -> tuple[Judgement, ...]:
    """Read the verdicts of a filled workbook, in workbook order.

    Raises:
        WorkbookValidationError: listing every problem, if the header changed, a sampled
            thesis is missing or repeated, a row names an unknown thesis or changed its
            status, or a verdict is blank or not allowed. The program comes from the sample,
            through the handle.
        FileNotFoundError: if there is no workbook at ``path``.
        ValueError: if the file at ``path`` is not a readable xlsx workbook.
    """
    workbook = _load_workbook(path)
    if _CHECK_SHEET not in workbook.sheetnames:
        raise WorkbookValidationError([f"the workbook has no sheet '{_CHECK_SHEET}'"])
    sheet = workbook[_CHECK_SHEET]
    header = tuple(_text(cell.value) for cell in sheet[1])[: len(COLUMNS)]
    if header != COLUMNS:
        raise WorkbookValidationError(
            [f"the header row of '{_CHECK_SHEET}' changed; it must be: {', '.join(COLUMNS)}"]
        )
    expected = {item.handle: item for item in sample.items}
    seen: Counter[str] = Counter()
    problems: list[str] = []
    judgements: list[Judgement] = []
    for row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = dict(zip(COLUMNS, (_text(value) for value in values), strict=False))
        if not any(cells.values()):
            continue
        handle = cells.get("Handle", "")
        item = expected.get(handle)
        if item is None:
            problems.append(f"row {row}: handle '{handle}' is not in the sample")
            continue
        seen[handle] += 1
        where = f"row {row} ({handle})"
        if cells.get("Estado") != STATUS_LABELS[item.status]:
            problems.append(
                f"{where}: 'Estado' was changed; it must be '{STATUS_LABELS[item.status]}'"
            )
        general = _verdict(cells, "Veredicto objetivo general", GENERAL_VERDICTS, where, problems)
        specific = _verdict(
            cells, "Veredicto objetivos específicos", SPECIFIC_VERDICTS, where, problems
        )
        if general is not None and general not in _ALLOWED_GENERAL[item.status]:
            allowed = [
                label
                for label, token in GENERAL_VERDICTS.items()
                if token in _ALLOWED_GENERAL[item.status]
            ]
            problems.append(
                f"{where}: '{cells['Veredicto objetivo general']}' is not allowed for a thesis "
                f"whose status is '{STATUS_LABELS[item.status]}'; use one of: {', '.join(allowed)}"
            )
            general = None
        if item.status == "not_found" and specific not in (None, "not_applicable"):
            problems.append(
                f"{where}: 'Veredicto objetivos específicos' must be 'No aplica' for a thesis "
                f"whose status is '{STATUS_LABELS[item.status]}'"
            )
            specific = None
        if general is not None and specific is not None:
            judgements.append(
                Judgement(item.n, handle, item.program, item.status, general, specific)
            )
    for handle, count in sorted(seen.items()):
        if count > 1:
            problems.append(f"handle {handle} appears {count} times")
    for item in sample.items:
        if item.handle not in seen:
            problems.append(f"handle {item.handle} (N.º {item.n}) is missing")
    if problems:
        raise WorkbookValidationError(problems)
    return tuple(judgements)


_UNREADABLE = (zipfile.BadZipFile, InvalidFileException, KeyError, SyntaxError, TypeError)
"""What openpyxl raises for a file that is not a readable xlsx workbook: no zip archive, an
extension it does not support, an archive without the workbook parts, broken XML (whose
``ParseError`` is a ``SyntaxError``), or a value of the wrong type in a part."""


def _load_workbook(path: Path) -> openpyxl.Workbook:
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except _UNREADABLE as error:
        raise ValueError(
            f"{path} is not a readable xlsx workbook ({type(error).__name__}: {error})"
        ) from error


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _verdict[T](
    cells: Mapping[str, str],
    name: str,
    labels: Mapping[str, T],
    where: str,
    problems: list[str],
) -> T | None:
    value = cells.get(name, "")
    if not value:
        problems.append(f"{where} has no '{name}'")
        return None
    folded = " ".join(value.split()).casefold()
    for label, token in labels.items():
        if label.casefold() == folded:
            return token
    problems.append(
        f"{where}: '{value}' is not an allowed '{name}'; use one of: {', '.join(labels)}"
    )
    return None


# Report


def wilson_interval(successes: int, total: int, z: float = Z_95) -> tuple[float, float]:
    """Return the Wilson score interval of a proportion of ``successes`` in ``total``.

    Raises:
        ValueError: if ``total`` is not at least one, or ``successes`` is outside 0..total.
    """
    if total < 1:
        raise ValueError(f"a proportion needs at least one row, got {total}")
    if not 0 <= successes <= total:
        raise ValueError(f"successes must lie in 0..{total}, got {successes}")
    observed = successes / total
    denominator = 1 + z**2 / total
    center = (observed + z**2 / (2 * total)) / denominator
    half = z * math.sqrt(observed * (1 - observed) / total + z**2 / (4 * total**2)) / denominator
    # The bounds are exactly 0 and 1 at the ends, which rounding would otherwise miss.
    low = 0.0 if successes == 0 else max(0.0, center - half)
    high = 1.0 if successes == total else min(1.0, center + half)
    return low, high


class ProgramAccuracy(_Model):
    judged: Count
    correct: Count
    partial: Count
    accuracy: float
    ci95: tuple[float, float]


class GeneralAccuracy(_Model):
    judged: Count
    correct: Count
    counts: dict[GeneralVerdict, Count]
    accuracy: float
    """``Correcto`` over the judged rows; ``Parcial`` counts as not correct."""
    ci95: tuple[float, float]
    partial_share: float
    accuracy_with_true_absences: float
    """Informative: also counts a ``not_found`` thesis judged ``No existe en la tesis`` as
    correct, since the locator rightly found nothing."""


class SpecificAccuracy(_Model):
    applicable: Count
    """Rows judged other than ``No aplica``."""
    correct: Count
    counts: dict[SpecificVerdict, Count]
    accuracy: float | None
    ci95: tuple[float, float] | None


class VerificationReport(_Model):
    """Numbers of a filled verification workbook: no text, no handle."""

    snapshot_id: str | None
    imported_at: UtcDatetime
    workbook_sha256: str | None
    rule: str
    target: float
    sample_size: Count
    general: GeneralAccuracy
    by_program: dict[str, ProgramAccuracy]
    specific: SpecificAccuracy
    meets_target: bool
    """The point accuracy of the general objective reaches :data:`TARGET_ACCURACY`."""


RULE = (
    "accuracy = rows judged 'Correcto' / judged rows; 'Parcial' counts as not correct (strict "
    "O05 rule). accuracy_with_true_absences also counts a not_found row judged 'No existe en "
    "la tesis' as correct."
)


def verification_report(
    judgements: Sequence[Judgement],
    sample: VerificationSample,
    *,
    imported_at: datetime,
    workbook_sha256: str | None = None,
) -> VerificationReport:
    """Compute the accuracy of the general and the specific objectives from ``judgements``.

    Raises:
        ValueError: if there are no judgements.
    """
    if not judgements:
        raise ValueError("there are no judged rows")
    general_counts = Counter(j.general for j in judgements)
    correct = general_counts["correct"]
    true_absences = sum(j.status == "not_found" and j.general == "absent" for j in judgements)
    by_program: dict[str, ProgramAccuracy] = {}
    for program in dict.fromkeys(j.program for j in sorted(judgements, key=lambda j: j.n)):
        rows = [j for j in judgements if j.program == program]
        hits = sum(j.general == "correct" for j in rows)
        by_program[program] = ProgramAccuracy(
            judged=len(rows),
            correct=hits,
            partial=sum(j.general == "partial" for j in rows),
            accuracy=hits / len(rows),
            ci95=wilson_interval(hits, len(rows)),
        )
    specific_counts = Counter(j.specific for j in judgements)
    applicable = len(judgements) - specific_counts["not_applicable"]
    specific_correct = specific_counts["correct"]
    total = len(judgements)
    return VerificationReport(
        snapshot_id=sample.snapshot_id,
        imported_at=imported_at,
        workbook_sha256=workbook_sha256,
        rule=RULE,
        target=TARGET_ACCURACY,
        sample_size=len(sample.items),
        general=GeneralAccuracy(
            judged=total,
            correct=correct,
            counts={token: general_counts[token] for token in GENERAL_VERDICTS.values()},
            accuracy=correct / total,
            ci95=wilson_interval(correct, total),
            partial_share=general_counts["partial"] / total,
            accuracy_with_true_absences=(correct + true_absences) / total,
        ),
        by_program=dict(
            sorted(by_program.items(), key=lambda entry: _program_rank(sample, entry[0]))
        ),
        specific=SpecificAccuracy(
            applicable=applicable,
            correct=specific_correct,
            counts={token: specific_counts[token] for token in SPECIFIC_VERDICTS.values()},
            accuracy=specific_correct / applicable if applicable else None,
            ci95=wilson_interval(specific_correct, applicable) if applicable else None,
        ),
        meets_target=correct / total >= TARGET_ACCURACY,
    )


def _program_rank(sample: VerificationSample, program: str) -> int:
    order = list(sample.population)
    return order.index(program) if program in order else len(order)


# Command line


def _parse_arguments(argv: Sequence[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="python -m thematic_redundancy.labeling.objectives_check",
        description=(
            "Import the filled objectives verification workbook of a snapshot, validate it, "
            "and compute the accuracy of the objectives locator (O05)."
        ),
    )
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(),
        help=f"project root that holds {CONFIG_PATH.as_posix()} (default: current directory)",
    )
    parser.add_argument(
        "--snapshot-id", help="snapshot of the workbook (default: the snapshot harvested last)"
    )
    parser.add_argument(
        "--report-out",
        type=Path,
        metavar="PATH",
        help="also write the accuracy report, numbers only, as JSON to PATH",
    )
    return parser.parse_args(argv)


def _percent(value: float) -> str:
    return f"{100 * value:.1f}%"


def _describe(report: VerificationReport) -> str:
    general = report.general
    low, high = general.ci95
    lines = [
        f"Accuracy of the general objective (strict): {general.correct}/{general.judged} = "
        f"{_percent(general.accuracy)} (95% CI {_percent(low)} to {_percent(high)}); target "
        f"{_percent(report.target)}: {'met' if report.meets_target else 'not met'}",
        f"Verdicts: {', '.join(f'{token} {count}' for token, count in general.counts.items())}",
        f"Counting true absences as correct: {_percent(general.accuracy_with_true_absences)}",
    ]
    for program, stats in report.by_program.items():
        program_low, program_high = stats.ci95
        lines.append(
            f"  {program}: {stats.correct}/{stats.judged} = {_percent(stats.accuracy)} "
            f"(95% CI {_percent(program_low)} to {_percent(program_high)}), "
            f"{stats.partial} partial"
        )
    specific = report.specific
    if specific.accuracy is not None:
        lines.append(
            f"Specific objectives: {specific.correct}/{specific.applicable} correct = "
            f"{_percent(specific.accuracy)}"
        )
    return "\n".join(lines)


def main(argv: Sequence[str] | None = None) -> int:
    """Run the import command and return its exit code: 0 when the workbook was imported,
    1 when it is invalid or unreadable, or a file is missing."""
    arguments = _parse_arguments(argv)
    project_root = arguments.project_root.resolve()
    try:
        config = load_config(project_root / CONFIG_PATH)
        snapshot_id = arguments.snapshot_id
        if snapshot_id is None:
            data_dir = config.paths.resolve_against(project_root).data_dir
            snapshot_id = latest_snapshot_id(data_dir / RAW_DIR_NAME)
        directory = workbook_directory(config, project_root, snapshot_id)
        sample_path = directory / SAMPLE_FILE_NAME
        try:
            sample = VerificationSample.model_validate_json(sample_path.read_bytes())
        except ValidationError as error:
            raise ValueError(f"{sample_path} is not a valid verification sample") from error
        workbook_path = directory / WORKBOOK_FILE_NAME
        judgements = import_verification(workbook_path, sample)
        report = verification_report(
            judgements,
            sample,
            imported_at=datetime.now(UTC).replace(microsecond=0),
            workbook_sha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        )
        if arguments.report_out is not None:
            arguments.report_out.parent.mkdir(parents=True, exist_ok=True)
            write_json_atomically(arguments.report_out, report.model_dump_json(indent=2))
    except WorkbookValidationError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    except (OSError, ValueError, yaml.YAMLError) as error:
        print(f"error: {' '.join(str(error).split())}", file=sys.stderr)
        return 1
    print(_describe(report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
