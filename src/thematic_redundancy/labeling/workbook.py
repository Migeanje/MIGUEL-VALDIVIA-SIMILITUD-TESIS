"""The labeled set's workbooks (T20b, D34 (1), (2), (3), (7)): the scale, the blind pair codes, the
display of a ficha, the order and sides of the pairs, the layout of a workbook, the name guard,
and the reader of a returned workbook.

The pilot workbook, the annotators' main workbooks (:mod:`.pair_workbooks`) and the adjudication
workbook of the disputed pairs (T20c) share this layout, and T20c's importer reads the returned
workbooks with :func:`read_labels`.

- **Scale** (D34 (3)): ``Redundante`` (mismo problema y mismo enfoque), ``Relacionado`` (mismo
  tema, distinto problema o enfoque) and ``No relacionado``; T22's binary label is 1 for
  ``Redundante`` only (D31 (b)), :data:`BINARY`.
- **Blind code** (D34 (2)): each drawn pair gets a code of two letters and three digits, such as
  ``KT482``, drawn at random from :mod:`secrets` (:func:`new_blind_codes`); it is not the pair
  id, and nothing tracked derives it. Its letters are never hexadecimal digits, so no SHA-256 in
  a manifest can hold a code by chance.
- **A pair on the sheet**: a block of four rows, one per field of the ficha (``Título``,
  ``Resumen``, ``Objetivos``, ``Palabras clave``), with the two theses side by side in ``Tesis A``
  and ``Tesis B``; ``N.º``, ``Código``, the label and the note span the block. The label is a
  drop-down list of the three levels. A thesis without objectives, abstract or keywords shows a
  fixed text in that field (:data:`NO_OBJECTIVES` and the like). No score, pool, stratum, rank,
  link, doc code, pair id or person is shown.
- **Order and sides**: a workbook's pairs are sorted by pair id and shuffled with
  ``random.Random("<seed>:order:<workbook>")``; a pair's two theses swap sides when
  ``random.Random("<seed>:sides:<pair id>").random()`` is below 0.5, the same in every workbook.
- **Name guard** (D34 (1)): every text of a workbook, its sheet names and drop-down lists
  included, is looked for the snapshot's person-name forms and ORCID iDs
  (:func:`~thematic_redundancy.corpus.profile.person_strings`) with
  :func:`~thematic_redundancy.corpus.profile.leaked_strings`: without word boundaries, in Unicode
  NFC, ignoring letter case and runs of whitespace. A hit refuses every workbook, naming where
  (workbook, sheet, cell and pair code), never the string.
"""

import random
import re
import secrets
from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from thematic_redundancy.corpus.ficha import Ficha
from thematic_redundancy.corpus.profile import leaked_strings
from thematic_redundancy.labeling.objectives_check import (
    _UNREADABLE,
    _add_list_validation,
    _load_workbook,
    _save_atomically,
    _text,
    _text_cell,
    _verdict,
    _write_guide_lines,
    _write_header,
)

Level = Literal["redundant", "related", "unrelated"]

LEVELS: dict[str, Level] = {
    "Redundante": "redundant",
    "Relacionado": "related",
    "No relacionado": "unrelated",
}
"""The three levels of the scale, as the workbook offers them, and their tokens (D34 (3))."""

LEVEL_DEFINITIONS: dict[Level, str] = {
    "redundant": "Mismo problema y mismo enfoque.",
    "related": "Mismo tema, distinto problema o enfoque.",
    "unrelated": "Distinto tema.",
}
"""What each level means, as the guide states it."""

BINARY: dict[Level, int] = {"redundant": 1, "related": 0, "unrelated": 0}
"""T22's binary label of each level: only ``Redundante`` counts as redundant (D31 (b))."""

FIELDS = ("Título", "Resumen", "Objetivos", "Palabras clave")
"""The rows of a pair's block, one per field of the ficha, in this order."""

NO_ABSTRACT = "(La ficha no registra resumen.)"
NO_OBJECTIVES = "(La ficha no registra objetivos.)"
NO_KEYWORDS = "(La ficha no registra palabras clave.)"
KEYWORD_SEPARATOR = "; "

GUIDE_SHEET = "Guía"
LABELING_SHEET = "Etiquetado"
EXAMPLES_SHEET = "Ejemplos"

CODE_LETTERS = "GHJKMNPQRSTVWXYZ"
"""The letters of a blind code: no vowel, no letter that reads as a digit (I, L, O) and no
hexadecimal digit (A to F)."""
CODE_DIGITS = "23456789"
"""The digits of a blind code, without 0 and 1."""
CODE_PATTERN = re.compile(rf"[{CODE_LETTERS}]{{2}}[{CODE_DIGITS}]{{3}}")

UNREADABLE = _UNREADABLE
"""What reading a file that is not a readable xlsx workbook raises, as the objectives workbook's
safe reader catches it, ``zlib.error`` and ``EOFError`` included (T10c)."""

_BLOCK = len(FIELDS)
_WIDTHS = (6, 9, 15, 70, 70, 18, 40)
_KEY_COLUMNS = ("N.º", "Código", "Campo")
_EDITABLE = PatternFill("solid", fgColor="FFF6D5")
_BLOCK_TOP = Border(top=Side(style="medium", color="8EA2BF"))


class LabelingWorkbookError(ValueError):
    """A returned workbook cannot be imported; ``problems`` lists every reason."""

    def __init__(self, problems: Sequence[str]) -> None:
        self.problems = tuple(problems)
        listed = "\n".join(f"- {problem}" for problem in self.problems)
        super().__init__(f"{len(self.problems)} problem(s) in the workbook:\n{listed}")


class NameGuardError(ValueError):
    """A workbook holds a person's name or an ORCID iD of the snapshot; ``locations`` says where,
    never what."""

    def __init__(self, locations: Sequence[str]) -> None:
        self.locations = tuple(locations)
        super().__init__(
            f"the name guard found a person's name or an ORCID iD of the snapshot in "
            f"{len(self.locations)} place(s), so no workbook was written: "
            + "; ".join(self.locations)
        )


@dataclass(frozen=True)
class SheetLayout:
    """The columns of a labeling sheet: the key columns, the two theses, the label and a note."""

    label_column: str
    note_column: str

    @property
    def columns(self) -> tuple[str, ...]:
        return (*_KEY_COLUMNS, "Tesis A", "Tesis B", self.label_column, self.note_column)


PILOT_LAYOUT = SheetLayout(label_column="Etiqueta acordada", note_column="Justificación")
"""The pilot's labeling sheet, and the examples sheet of the main workbooks."""
MAIN_LAYOUT = SheetLayout(label_column="Etiqueta", note_column="Notas (opcional)")
"""An annotator's labeling sheet."""


# Blind codes, order and sides


def new_blind_codes(
    pair_ids: Iterable[str], *, choice: Callable[[str], str] = secrets.choice
) -> dict[str, str]:
    """Return a new, distinct blind code for each pair id, in pair-id order.

    ``choice`` picks one character of a string; by default from :mod:`secrets`, so the codes can
    be neither derived again nor told from the pair ids. A code already given is drawn again.
    """
    codes: dict[str, str] = {}
    taken: set[str] = set()
    for identifier in sorted(set(pair_ids)):
        code = _new_code(choice)
        while code in taken:
            code = _new_code(choice)
        codes[identifier] = code
        taken.add(code)
    return codes


def _new_code(choice: Callable[[str], str]) -> str:
    return "".join(choice(CODE_LETTERS) for _ in range(2)) + "".join(
        choice(CODE_DIGITS) for _ in range(3)
    )


def check_blind_codes(codes: Mapping[str, str], pair_ids: Iterable[str]) -> None:
    """Check that ``codes`` gives each of ``pair_ids``, and nothing else, a distinct valid code.

    Raises:
        ValueError: naming the number of pairs at fault, never a code.
    """
    expected = set(pair_ids)
    problems = []
    if set(codes) != expected:
        problems.append(
            f"{len(expected - set(codes))} drawn pairs have no code and "
            f"{len(set(codes) - expected)} codes belong to no drawn pair"
        )
    if len(set(codes.values())) != len(codes):
        problems.append("two pairs share a code")
    malformed = sum(not CODE_PATTERN.fullmatch(code) for code in codes.values())
    if malformed:
        problems.append(f"{malformed} codes are not two code letters and three code digits")
    if problems:
        raise ValueError("the blind codes do not fit the draw: " + "; ".join(problems))


def shuffled(pair_ids: Iterable[str], *, seed: int, workbook: str) -> list[str]:
    """Return ``pair_ids`` in the order of ``workbook``: sorted, then shuffled with
    ``random.Random("<seed>:order:<workbook>")``."""
    ordered = sorted(pair_ids)
    random.Random(f"{seed}:order:{workbook}").shuffle(ordered)
    return ordered


def swapped(pair_id: str, *, seed: int) -> bool:
    """Tell whether the pair's two theses swap sides: the later doc code shown as ``Tesis A``.
    It depends on the pair alone, so a pair has the same sides in every workbook."""
    return random.Random(f"{seed}:sides:{pair_id}").random() < 0.5


# What a pair shows


@dataclass(frozen=True)
class FichaView:
    """The four fields of a ficha as a workbook shows them."""

    title: str
    abstract: str
    objectives: str
    keywords: str

    @property
    def texts(self) -> tuple[str, str, str, str]:
        """The fields in the order of :data:`FIELDS`."""
        return (self.title, self.abstract, self.objectives, self.keywords)


def ficha_view(ficha: Ficha, redact: Callable[[str], str] = lambda text: text) -> FichaView:
    """Return the four fields of ``ficha``, each passed through ``redact`` (the name redactor of
    the objectives workbook), and a fixed text for a field the ficha lacks."""
    keywords = KEYWORD_SEPARATOR.join(ficha.keywords)
    return FichaView(
        title=redact(ficha.title),
        abstract=redact(ficha.abstract) if ficha.abstract else NO_ABSTRACT,
        objectives=redact(ficha.objectives) if ficha.objectives else NO_OBJECTIVES,
        keywords=redact(keywords) if keywords else NO_KEYWORDS,
    )


@dataclass(frozen=True)
class PairView:
    """One pair of a sheet: its number, its code, its two theses and, on the examples sheet, its
    agreed label and justification."""

    n: int
    code: str
    left: FichaView
    right: FichaView
    label: str | None = None
    note: str | None = None


def label_text(level: Level) -> str:
    """Return the label of ``level`` as the workbook offers it."""
    return next(label for label, token in LEVELS.items() if token == level)


# Building


@dataclass(frozen=True)
class SheetContent:
    """One labeling or examples sheet: its name, its columns, its pairs and whether its label
    column offers the drop-down list."""

    title: str
    layout: SheetLayout
    pairs: Sequence[PairView]
    editable: bool


def build_workbook(
    guide: Sequence[tuple[str, str]], sheets: Sequence[SheetContent]
) -> openpyxl.Workbook:
    """Return a workbook with the guide sheet first, then ``sheets``; the last one opens first.

    ``guide`` gives one ``(label, text)`` row each; a row without text is a section heading. No
    cell is ever a formula.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = GUIDE_SHEET
    _write_guide_lines(sheet, guide, {label for label, text in guide if not text})
    for content in sheets:
        _write_pairs(workbook.create_sheet(content.title), content)
    workbook.active = len(sheets)
    return workbook


def _write_pairs(sheet: Worksheet, content: SheetContent) -> None:
    columns = content.layout.columns
    _write_header(sheet, columns, _WIDTHS)
    label_column = columns.index(content.layout.label_column) + 1
    for index, pair in enumerate(content.pairs):
        first = 2 + _BLOCK * index
        values = {1: pair.n, 2: pair.code, label_column: pair.label, label_column + 1: pair.note}
        for offset, (field, left, right) in enumerate(
            zip(FIELDS, pair.left.texts, pair.right.texts, strict=True)
        ):
            row = first + offset
            for column, value in ((3, field), (4, left), (5, right)):
                cell = _text_cell(sheet, row, column, value)
                cell.alignment = Alignment(wrap_text=column > 3, vertical="top")
                cell.font = Font(bold=offset == 0 and column > 3)
        for column in range(1, len(columns) + 1):
            if column in values:
                cell = _text_cell(sheet, first, column, values[column])
                cell.alignment = Alignment(wrap_text=column > 3, vertical="top")
                if column == 2:
                    cell.font = Font(bold=True)
                if content.editable and column >= label_column:
                    cell.fill = _EDITABLE
                sheet.merge_cells(
                    start_row=first,
                    start_column=column,
                    end_row=first + _BLOCK - 1,
                    end_column=column,
                )
            sheet.cell(first, column).border = _BLOCK_TOP
    if content.editable and content.pairs:
        last = 1 + _BLOCK * len(content.pairs)
        _add_list_validation(sheet, label_column, LEVELS, 2, last)
    sheet.freeze_panes = "D2"


def save_workbook(workbook: openpyxl.Workbook, path: Path) -> None:
    """Save ``workbook`` at ``path`` through a temporary file beside it, so it appears whole or
    not at all."""
    _save_atomically(workbook, path)


# The name guard


class _Guard:
    """Looks for the guard's strings in a text, once per distinct text."""

    def __init__(self, strings: frozenset[str]) -> None:
        self._strings = strings
        self.seen: dict[str, bool] = {}

    def holds(self, text: str) -> bool:
        if text not in self.seen:
            self.seen[text] = bool(leaked_strings(text, self._strings))
        return self.seen[text]


def name_guard_hits(workbook: openpyxl.Workbook, guard: frozenset[str], *, name: str) -> list[str]:
    """Return where the workbook ``name`` holds a string of ``guard``: each cell, drop-down list
    or sheet name, with the code of the pair whose block holds the cell. The strings are never
    named."""
    return _hits(workbook, _Guard(guard), name)


def guard_workbooks(workbooks: Mapping[str, openpyxl.Workbook], guard: frozenset[str]) -> int:
    """Look for the strings of ``guard`` in every text of every workbook, by name, and return
    the number of distinct texts looked at.

    Raises:
        NameGuardError: listing every place of every workbook that holds one.
    """
    checker = _Guard(guard)
    found = [hit for name, book in workbooks.items() for hit in _hits(book, checker, name)]
    if found:
        raise NameGuardError(found)
    return len(checker.seen)


def _hits(workbook: openpyxl.Workbook, checker: _Guard, name: str) -> list[str]:
    found = []
    for number, sheet in enumerate(workbook.worksheets, start=1):
        if checker.holds(sheet.title):
            where = f"{name}, sheet {number}"
            found.append(f"{where}, its name")
        else:
            where = f"{name}, sheet '{sheet.title}'"
        codes = _codes_by_row(sheet)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and checker.holds(cell.value):
                    code = codes.get(cell.row)
                    pair = f" (pair {code})" if code else ""
                    found.append(f"{where}, cell {cell.coordinate}{pair}")
        for rule in sheet.data_validations.dataValidation:
            texts = (rule.formula1, rule.error, rule.errorTitle, rule.prompt, rule.promptTitle)
            if any(isinstance(text, str) and checker.holds(text) for text in texts):
                found.append(f"{where}, the drop-down list of {rule.sqref}")
    return found


def _codes_by_row(sheet: Worksheet) -> dict[int, str]:
    """Map each row of a pair's block to the pair's code, on a sheet with a ``Código`` column."""
    header = [cell.value for cell in sheet[1]]
    if "Código" not in header:
        return {}
    column = header.index("Código") + 1
    codes: dict[int, str] = {}
    current: str | None = None
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, column).value
        if isinstance(value, str) and value:
            current = value
        if current is not None:
            codes[row] = current
    return codes


# Reading a returned workbook


@dataclass(frozen=True)
class ReadLabel:
    """The label of one pair of a returned workbook."""

    n: int
    code: str
    level: Level
    note: str | None


def load_workbook(path: Path) -> openpyxl.Workbook:
    """Open the workbook at ``path`` with its values as last saved, through the objectives
    workbook's safe reader, which this module shares.

    Raises:
        FileNotFoundError: if there is no file at ``path``.
        ValueError: if it is not a readable xlsx workbook, in one line.
    """
    return _load_workbook(path)


def read_labels(
    path: Path, layout: SheetLayout, expected: Mapping[str, int], *, sheet: str = LABELING_SHEET
) -> tuple[ReadLabel, ...]:
    """Read the label and the note of every pair of the labeling sheet at ``path``, in sheet
    order; ``expected`` maps each code the workbook was made with to its ``N.º``.

    Raises:
        LabelingWorkbookError: listing every problem: the sheet or its header changed; a code is
            unknown, repeated or missing, or its ``N.º`` or its block's ``Campo`` rows changed; a
            label is blank or not one of :data:`LEVELS`, or stands outside the first row of its
            block.
        FileNotFoundError: if there is no workbook at ``path``.
        ValueError: if the file at ``path`` is not a readable xlsx workbook.
    """
    workbook = load_workbook(path)
    if sheet not in workbook.sheetnames:
        raise LabelingWorkbookError([f"the workbook has no sheet '{sheet}'"])
    rows = [
        tuple(_text(value) for value in values)
        for values in workbook[sheet].iter_rows(values_only=True)
    ]
    columns = layout.columns
    if not rows or rows[0][: len(columns)] != columns:
        raise LabelingWorkbookError(
            [f"the header row of '{sheet}' changed; it must be: {', '.join(columns)}"]
        )
    problems: list[str] = []
    labels: list[ReadLabel] = []
    seen: Counter[str] = Counter()
    index = {name: number for number, name in enumerate(columns)}
    blocks = _blocks(rows[1:], index["Código"], problems)
    for start, block in blocks:
        cells = dict(zip(columns, block[0], strict=False))
        code = cells["Código"]
        where = f"row {start} (pair {code})"
        if code not in expected:
            problems.append(f"row {start}: the code '{code}' is not one of this workbook's pairs")
            continue
        seen[code] += 1
        if cells["N.º"] != str(expected[code]):
            problems.append(f"{where}: 'N.º' was changed; it must be {expected[code]}")
        fields = tuple(row[index["Campo"]] if len(row) > index["Campo"] else "" for row in block)
        if fields != FIELDS:
            problems.append(
                f"{where}: the 'Campo' rows of the pair were changed; they must be "
                f"{', '.join(FIELDS)}"
            )
        for offset, row in enumerate(block[1:], start=1):
            extra = [
                name
                for name in (layout.label_column, layout.note_column)
                if len(row) > index[name] and row[index[name]]
            ]
            if extra:
                problems.append(
                    f"row {start + offset} (pair {code}): '{extra[0]}' must be in the first row "
                    "of the pair"
                )
        level = _verdict(cells, layout.label_column, LEVELS, where, problems)
        if level is not None:
            labels.append(
                ReadLabel(int(expected[code]), code, level, cells.get(layout.note_column) or None)
            )
    for code, count in sorted(seen.items()):
        if count > 1:
            problems.append(f"the pair {code} appears {count} times")
    for code, n in sorted(expected.items(), key=lambda entry: entry[1]):
        if code not in seen:
            problems.append(f"the pair {code} (N.º {n}) is missing")
    if problems:
        raise LabelingWorkbookError(problems)
    return tuple(labels)


def _blocks(
    rows: Sequence[tuple[str, ...]], code_column: int, problems: list[str]
) -> list[tuple[int, list[tuple[str, ...]]]]:
    """Group the data rows into blocks, each starting at a row with a code, as ``(sheet row,
    rows)``; a non-empty row before the first code is a problem."""
    blocks: list[tuple[int, list[tuple[str, ...]]]] = []
    for number, row in enumerate(rows, start=2):
        code = row[code_column] if len(row) > code_column else ""
        if code:
            blocks.append((number, [row]))
        elif blocks and len(blocks[-1][1]) < _BLOCK:
            blocks[-1][1].append(row)
        elif any(row):
            problems.append(f"row {number} belongs to no pair")
    return blocks


def screen_notes(
    labels: Sequence[ReadLabel],
    redact: Callable[[str], str],
    guard: frozenset[str],
    *,
    column: str,
) -> tuple[tuple[ReadLabel, ...], list[str]]:
    """Return ``labels`` with each note passed through ``redact`` (the name redactor of every
    displayed text), and a problem for each note that still holds a string of ``guard`` (a name
    glued to another word, or an ORCID iD), naming the pair and the ``column``, never the
    string. A note may be shown again, as a justification on the ``Ejemplos`` sheet, so it gets
    the same name guard as every text of a workbook (D34 (1))."""
    checker = _Guard(guard)
    screened, problems = [], []
    for label in labels:
        note = None if label.note is None else redact(label.note)
        if note is not None and checker.holds(note):
            problems.append(
                f"the '{column}' of pair {label.code} holds a person's name or an ORCID iD of the "
                "snapshot; write it without them"
            )
        screened.append(replace(label, note=note))
    return tuple(screened), problems


def missing_levels(labels: Iterable[ReadLabel]) -> tuple[str, ...]:
    """Return the labels of the levels that no pair of ``labels`` has, in scale order."""
    present = {label.level for label in labels}
    return tuple(label for label, level in LEVELS.items() if level not in present)
