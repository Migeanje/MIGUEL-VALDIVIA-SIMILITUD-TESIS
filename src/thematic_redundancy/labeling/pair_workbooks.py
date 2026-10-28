"""Make the labeled set's workbooks (T20b, D34) and import the pilot's agreed labels::

    uv run python -m thematic_redundancy.labeling.pair_workbooks pilot [--dry-run]
    uv run python -m thematic_redundancy.labeling.pair_workbooks import-pilot [--dry-run]
    uv run python -m thematic_redundancy.labeling.pair_workbooks main [--dry-run]

each with ``[--project-root PATH] [--snapshot-id ID]``. They read the draw of T20a (its tracked
``manifest.json`` and the gitignored ``pairs.parquet`` it vouches for), the included fichas and
the snapshot, for the name redactor, the name guard and the privacy guard. The layout, the scale
and the guard are :mod:`.workbook`'s.

- ``pilot`` makes the pilot workbook, its 10 pairs in a random order with random sides, and gives
  every drawn pair, pilot and labeled set, its blind code, once (D34 (2), (7)).
- ``import-pilot`` reads the pilot workbook back from the calibration meeting and stores its
  agreed labels and justifications, once it holds the workbook's 10 pairs unchanged, a label of
  the scale on every pair and at least one pair of each level (D34 (7)).
- ``main`` makes one workbook per annotator with the pairs of the annotator's pairs of
  annotators (:mod:`.assignment`, D34 (8)), whose ``Ejemplos`` sheet shows the pilot's agreed
  labels. It refuses until ``labeling.annotators`` is set and the pilot's labels are stored.

Each step happens once: a real run refuses to overwrite its records and a dirty git tree
(changes outside ``results/labeled_set/<snapshot>/workbooks/``), since its manifest records the
commit. It holds the run lock of that directory and writes the gitignored records first, under
``labeling.records_dir/<snapshot>/`` (the workbooks in its ``workbooks/`` folder), and its tracked
manifest last, as the commit point: counts and SHA-256 values, no doc code, pair id or blind
code. Before anything is written, every workbook passes the name guard and the manifest the
privacy guard; a hit refuses every write. ``--dry-run`` computes everything and applies both
guards, writing nothing. Progress goes to stderr, a summary without thesis text to stdout; the
exit code is 0 on success, 1 after an error, and 130 after an interruption.
"""

import argparse
import hashlib
import secrets
import sys
import time
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from pydantic import BaseModel, ConfigDict, ValidationError

from thematic_redundancy.corpus.build_fichas import FICHAS_MANIFEST_FILE_NAME, load_fichas
from thematic_redundancy.corpus.ficha import Ficha
from thematic_redundancy.corpus.profile import person_name_forms, person_strings, sensitive_strings
from thematic_redundancy.corpus.snapshot import (
    METADATA_FILE_NAME,
    RAW_DIR_NAME,
    check_snapshot_id,
    latest_snapshot_id,
    read_snapshot,
)
from thematic_redundancy.extraction.text_manifest import INTERIM_DIR_NAME
from thematic_redundancy.gaps.results_store import check_privacy
from thematic_redundancy.labeling.assignment import Assignment, annotator_pairs, assign_pairs
from thematic_redundancy.labeling.draw_pairs import (
    LABELED_SET_DIR_NAME,
    MANIFEST_FILE_NAME,
    PAIRS_FILE_NAME,
    DrawManifest,
    PairRow,
)
from thematic_redundancy.labeling.draw_pairs import records_dir as draw_records_dir
from thematic_redundancy.labeling.draw_pairs import results_dir as draw_results_dir
from thematic_redundancy.labeling.objectives_check import _NameRedactor
from thematic_redundancy.labeling.pairs import LABELED, PILOT, POOLS
from thematic_redundancy.labeling.workbook import (
    BINARY,
    EXAMPLES_SHEET,
    LABELING_SHEET,
    LEVEL_DEFINITIONS,
    LEVELS,
    MAIN_LAYOUT,
    NO_ABSTRACT,
    NO_KEYWORDS,
    NO_OBJECTIVES,
    PILOT_LAYOUT,
    LabelingWorkbookError,
    Level,
    PairView,
    SheetContent,
    build_workbook,
    check_blind_codes,
    ficha_view,
    guard_workbooks,
    label_text,
    missing_levels,
    new_blind_codes,
    read_labels,
    save_workbook,
    shuffled,
    swapped,
)
from thematic_redundancy.shared.config import AppConfig, load_config
from thematic_redundancy.shared.filesystem import exclusive_run, write_text_atomically
from thematic_redundancy.shared.storage import ChecksumMismatchError, ParquetTableStore
from thematic_redundancy.topic_map.results_store import library_versions
from thematic_redundancy.topic_map.run_experiment import GitState, read_git_state

CONFIG_PATH = Path("config") / "default.yaml"
RUN_COMMAND = "uv run python -m thematic_redundancy.labeling.pair_workbooks"
WORKBOOKS_DIR_NAME = "workbooks"
"""The folder of the tracked manifests, under ``results/labeled_set/<snapshot>/``, and of the
workbooks, under ``labeling.records_dir/<snapshot>/``."""
BLIND_CODES_FILE_NAME = "blind_codes.parquet"
PILOT_LAYOUT_FILE_NAME = "pilot_layout.parquet"
PILOT_LABELS_FILE_NAME = "pilot_labels.parquet"
ASSIGNMENT_FILE_NAME = "assignment.parquet"
MAIN_LAYOUT_FILE_NAME = "main_layout.parquet"
PILOT_WORKBOOK_FILE_NAME = "etiquetado_piloto.xlsx"
PILOT_MANIFEST_FILE_NAME = "pilot_manifest.json"
PILOT_LABELS_MANIFEST_FILE_NAME = "pilot_labels_manifest.json"
MAIN_MANIFEST_FILE_NAME = "main_manifest.json"
PILOT_WORKBOOK = "pilot"
"""The pilot workbook's key in the layout records and in the order's seed."""
DRAW_MANIFEST_KEY = "draw/manifest.json"


def annotator_workbook(annotator: int) -> str:
    """Return the key of an annotator's workbook in the layout records and the order's seed."""
    return f"annotator-{annotator}"


def annotator_file_name(annotator: int) -> str:
    """Return the file name of an annotator's workbook, such as ``etiquetado_anotador_1.xlsx``."""
    return f"etiquetado_anotador_{annotator}.xlsx"


class WorkbooksError(ValueError):
    """A step cannot go on: a record it needs is missing, or the configuration or the git tree
    stops it."""


class PartialWorkbooksError(WorkbooksError):
    """Writing a step stopped after some of its records were written, before its manifest."""


# --- Records --------------------------------------------------------------------------------


class _Row(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid", hide_input_in_errors=True)


class BlindCodeRow(_Row):
    """One row of the gitignored ``blind_codes.parquet``: a drawn pair's blind code."""

    pair_id: str
    code: str


class LayoutRow(_Row):
    """One row of the gitignored ``pilot_layout.parquet`` or ``main_layout.parquet``: one pair of
    one workbook, where it stands and which thesis it shows on each side."""

    workbook: str
    """``pilot``, or ``annotator-<k>`` for annotator ``k``'s workbook."""
    annotator: int | None
    n: int
    """The pair's ``N.º`` in the workbook, from 1 in the workbook's order."""
    code: str
    pair_id: str
    doc_code_left: str
    """The thesis shown as ``Tesis A``."""
    doc_code_right: str
    """The thesis shown as ``Tesis B``."""
    swapped: bool
    """Whether ``Tesis A`` is the later doc code of the two (:func:`~.workbook.swapped`)."""


class PilotLabelRow(_Row):
    """One row of the gitignored ``pilot_labels.parquet``: a pilot pair's agreed label."""

    n: int
    code: str
    pair_id: str
    level: Level
    binary: int
    """T22's binary label: 1 for ``Redundante`` only (D31 (b))."""
    justification: str | None


class AssignmentRow(_Row):
    """One row of the gitignored ``assignment.parquet``: the pair of annotators of one pair."""

    pair_id: str
    code: str
    pool: str
    source: str | None
    group: int
    annotator_a: int
    annotator_b: int


class WorkbookManifest(BaseModel):
    """Content of a step's tracked manifest: provenance, values, counts and digests, without a
    doc code, a pair id or a blind code."""

    model_config = ConfigDict(frozen=True, extra="forbid", hide_input_in_errors=True)

    step: str
    """``pilot``, ``import-pilot`` or ``main``."""
    snapshot_id: str
    git_commit: str
    seed: int
    settings: dict[str, Any]
    counts: dict[str, int]
    details: dict[str, Any]
    """The step's aggregates: the levels of the pilot's labels, or the pairs of each pair of
    annotators by pool."""
    inputs: dict[str, str]
    """SHA-256 of every input read, by name."""
    data_files: dict[str, str]
    """SHA-256 of each file written under the records directory, by its path there."""
    versions: dict[str, str]
    started_at: datetime
    finished_at: datetime
    seconds: float


# --- Places and inputs ----------------------------------------------------------------------


@dataclass(frozen=True)
class Places:
    """Where a snapshot's workbooks and their records live."""

    root: Path
    snapshot_id: str
    results: Path
    """``results/labeled_set/<snapshot>/workbooks``: the tracked manifests."""
    draw: Path
    """``results/labeled_set/<snapshot>/draw``: the draw's tracked records."""
    records: Path
    """``labeling.records_dir/<snapshot>``: the gitignored records."""

    @classmethod
    def of(cls, config: AppConfig, root: Path, snapshot_id: str) -> "Places":
        results_root = config.paths.resolve_against(root).results_dir
        return cls(
            root=root,
            snapshot_id=snapshot_id,
            results=results_root / LABELED_SET_DIR_NAME / snapshot_id / WORKBOOKS_DIR_NAME,
            draw=draw_results_dir(results_root, snapshot_id),
            records=draw_records_dir(config, root, snapshot_id),
        )

    @property
    def workbooks(self) -> Path:
        return self.records / WORKBOOKS_DIR_NAME

    @property
    def pilot_workbook(self) -> Path:
        return self.workbooks / PILOT_WORKBOOK_FILE_NAME

    def relative(self, path: Path) -> str:
        return path.relative_to(self.root).as_posix()


@dataclass(frozen=True)
class WorkbookInputs:
    """What the workbooks read besides the configuration and their own records."""

    pairs: tuple[PairRow, ...]
    """The drawn pairs, as the draw recorded them."""
    fichas: Mapping[str, Ficha]
    """The included fichas, by doc code."""
    names: frozenset[str]
    """The person-name forms that the name redactor replaces (``person_name_forms``)."""
    guard: frozenset[str]
    """The strings of the name guard (``person_strings``)."""
    sensitive: frozenset[str]
    """The snapshot's strings that no tracked output may hold (``sensitive_strings``)."""
    input_sha256: dict[str, str]


def load_workbook_inputs(config: AppConfig, root: Path, snapshot_id: str) -> WorkbookInputs:
    """Load the draw of ``snapshot_id``, its pairs checked against the draw's manifest, the
    included fichas and the snapshot, each through its own loader.

    Raises:
        FileNotFoundError: if an input is missing, the draw included.
        ValueError: if an input is damaged, the pairs differ from the draw's manifest, or a drawn
            thesis is not an included ficha.
    """
    places = Places.of(config, root, snapshot_id)
    draw_path = places.draw / MANIFEST_FILE_NAME
    if not draw_path.is_file():
        raise FileNotFoundError(
            f"there is no draw of snapshot {snapshot_id} ({draw_path}); run "
            "python -m thematic_redundancy.labeling.draw_pairs first"
        )
    draw_text = draw_path.read_bytes()
    try:
        draw = DrawManifest.model_validate_json(draw_text)
    except ValidationError as error:
        raise ValueError(f"{draw_path} is not a valid draw manifest") from error
    pairs_path = places.records / PAIRS_FILE_NAME
    expected = draw.data_files.get(PAIRS_FILE_NAME)
    if expected is None:
        raise ValueError(f"{draw_path} records no SHA-256 of {PAIRS_FILE_NAME}")
    pairs = _read_verified(PairRow, pairs_path, expected, f"the draw's {MANIFEST_FILE_NAME}")
    data_dir = config.paths.resolve_against(root).data_dir
    interim = data_dir / INTERIM_DIR_NAME / snapshot_id
    fichas = {ficha.doc_code: ficha for ficha in load_fichas(interim) if ficha.include}
    absent = {code for row in pairs for code in (row.doc_code_a, row.doc_code_b)} - set(fichas)
    if absent:
        raise ValueError(f"{len(absent)} drawn theses are not included fichas of the snapshot")
    snapshot, records = read_snapshot(data_dir / RAW_DIR_NAME / snapshot_id)
    return WorkbookInputs(
        pairs=pairs,
        fichas=fichas,
        names=person_name_forms(records),
        guard=person_strings(records),
        sensitive=sensitive_strings(records),
        input_sha256={
            DRAW_MANIFEST_KEY: hashlib.sha256(draw_text).hexdigest(),
            PAIRS_FILE_NAME: expected,
            FICHAS_MANIFEST_FILE_NAME: _sha256(interim / FICHAS_MANIFEST_FILE_NAME),
            METADATA_FILE_NAME: snapshot.metadata_sha256,
        },
    )


def _read_verified[M: BaseModel](model: type[M], path: Path, sha256: str, by: str) -> tuple[M, ...]:
    try:
        return ParquetTableStore().read_verified(model, path, sha256=sha256)
    except ChecksumMismatchError as error:
        raise ValueError(f"{path} differs from what {by} records") from error


# --- Building -------------------------------------------------------------------------------


def layout_rows(
    workbook: str,
    annotator: int | None,
    pairs: Sequence[PairRow],
    codes: Mapping[str, str],
    *,
    seed: int,
) -> tuple[LayoutRow, ...]:
    """Return the rows of one workbook: its pairs in the workbook's order, each with its code and
    its sides (:func:`~.workbook.shuffled`, :func:`~.workbook.swapped`)."""
    by_id = {pair.pair_id: pair for pair in pairs}
    rows = []
    for n, identifier in enumerate(shuffled(by_id, seed=seed, workbook=workbook), start=1):
        pair = by_id[identifier]
        turned = swapped(identifier, seed=seed)
        left, right = (
            (pair.doc_code_b, pair.doc_code_a) if turned else (pair.doc_code_a, pair.doc_code_b)
        )
        rows.append(
            LayoutRow(
                workbook=workbook,
                annotator=annotator,
                n=n,
                code=codes[identifier],
                pair_id=identifier,
                doc_code_left=left,
                doc_code_right=right,
                swapped=turned,
            )
        )
    return tuple(rows)


@dataclass(frozen=True)
class Lacking:
    """The distinct theses that some pairs show, and those of them that lack a field."""

    theses: int
    without_abstract: int
    without_objectives: int
    without_keywords: int

    @classmethod
    def of(cls, rows: Sequence[LayoutRow], fichas: Mapping[str, Ficha]) -> "Lacking":
        theses = [fichas[c] for c in {c for r in rows for c in (r.doc_code_left, r.doc_code_right)}]
        return cls(
            theses=len(theses),
            without_abstract=sum(not ficha.abstract for ficha in theses),
            without_objectives=sum(not ficha.objectives for ficha in theses),
            without_keywords=sum(not ficha.keywords for ficha in theses),
        )

    def counts(self) -> dict[str, int]:
        return {
            "theses": self.theses,
            "without_abstract": self.without_abstract,
            "without_objectives": self.without_objectives,
            "without_keywords": self.without_keywords,
        }

    def line(self) -> str:
        return (
            f"  shown without objectives: {self.without_objectives} theses; without an abstract: "
            f"{self.without_abstract}; without keywords: {self.without_keywords}"
        )


def pair_views(
    rows: Sequence[LayoutRow],
    fichas: Mapping[str, Ficha],
    redactor: Callable[[str], str],
    labels: Mapping[str, PilotLabelRow] | None = None,
) -> tuple[PairView, ...]:
    """Return the pairs of ``rows`` as a sheet shows them, each field through ``redactor``, with
    the agreed label and justification of ``labels``, by code, on an examples sheet."""
    views = []
    for row in rows:
        label = None if labels is None else labels[row.code]
        views.append(
            PairView(
                n=row.n,
                code=row.code,
                left=ficha_view(fichas[row.doc_code_left], redactor),
                right=ficha_view(fichas[row.doc_code_right], redactor),
                label=None if label is None else label_text(label.level),
                note=None if label is None else label.justification,
            )
        )
    return tuple(views)


def _scale_lines() -> list[tuple[str, str]]:
    return [
        ("La escala", ""),
        *((label, LEVEL_DEFINITIONS[level]) for label, level in LEVELS.items()),
        (
            "Etiqueta binaria",
            "Para evaluar el sistema, solo «Redundante» cuenta como par redundante; «Relacionado» "
            "y «No relacionado» cuentan como pares no redundantes.",
        ),
    ]


def _shared_lines() -> list[tuple[str, str]]:
    return [
        ("Reglas del etiquetado ciego", ""),
        (
            "Solo este libro",
            "Juzgue cada par solo con lo que muestra este libro. No busque las tesis en el "
            "repositorio, en internet ni en otras fuentes.",
        ),
        (
            "Sin identidades",
            "El libro no muestra autores, asesores, jurados, enlaces ni datos de cómo se eligieron "
            "los pares. El código de cada par no identifica las tesis.",
        ),
        (
            "Nombres omitidos",
            "Si un texto muestra «[nombre omitido]», allí figuraba el nombre de una persona.",
        ),
        (
            "Campos ausentes",
            f"Algunas tesis solo tienen título y resumen en el repositorio: su campo «Objetivos» "
            f"muestra «{NO_OBJECTIVES}». Del mismo modo, una ficha sin resumen muestra "
            f"«{NO_ABSTRACT}» y una sin palabras clave, «{NO_KEYWORDS}». Juzgue esos pares con "
            "los campos disponibles.",
        ),
        (
            "No modifique",
            "Las columnas «N.º», «Código» y «Campo», ni el orden ni los nombres de las columnas: "
            "la importación los usa para validar el libro.",
        ),
        (
            "Tiempo",
            "Anote el inicio y el fin de cada sesión de etiquetado: ese tiempo se registra como "
            "costo del proceso.",
        ),
        (
            "Privacidad",
            "El libro contiene fragmentos de tesis con licencia CC BY-NC-ND 4.0, solo para esta "
            "investigación: no lo comparta ni lo suba a otros servicios.",
        ),
    ]


_COMPARED = (
    "Qué se compara",
    "Cada par muestra dos tesis, «Tesis A» y «Tesis B», con cuatro campos de su ficha: título, "
    "resumen, objetivos y palabras clave. Cada par ocupa cuatro filas, una por campo. El orden "
    "de los pares y el lado de cada tesis se eligieron al azar.",
)
_JUDGE = (
    "2.",
    "Lea los cuatro campos de ambas tesis. Juzgue el problema que cada tesis aborda y el enfoque "
    "con que lo aborda, no solo las palabras que comparten.",
)


def pilot_guide(pairs: int) -> list[tuple[str, str]]:
    """Return the lines of the pilot workbook's guide."""
    return [
        ("Etiquetado de pares de tesis: piloto", ""),
        (
            "Propósito",
            f"Este libro reúne los {pairs} pares del piloto. Los anotadores los etiquetan juntos "
            "en la reunión de calibración y acuerdan una etiqueta por par. Las etiquetas "
            "acordadas serán los ejemplos de la guía de los libros individuales. Los pares del "
            "piloto no forman parte del conjunto etiquetado ni del cálculo del acuerdo entre "
            "anotadores.",
        ),
        _COMPARED,
        *_scale_lines(),
        ("Cómo etiquetar", ""),
        ("1.", f"Vaya a la hoja «{LABELING_SHEET}». Cada par tiene un número y un código."),
        _JUDGE,
        (
            "3.",
            "Elija en «Etiqueta acordada», con la lista desplegable, la etiqueta en que coinciden "
            "los anotadores, y anote en «Justificación» por qué, en una o dos frases.",
        ),
        (
            "4.",
            f"Etiquete los {pairs} pares y guarde el archivo con el mismo nombre y en la misma "
            "carpeta. Como las etiquetas acordadas servirán de ejemplos, la importación del libro "
            "exige al menos un par de cada nivel de la escala.",
        ),
        *_shared_lines(),
    ]


def main_guide(annotator: int, pairs: int, examples: int) -> list[tuple[str, str]]:
    """Return the lines of annotator ``annotator``'s workbook guide."""
    return [
        (f"Etiquetado de pares de tesis: anotador {annotator}", ""),
        (
            "Propósito",
            f"Este libro reúne los {pairs} pares que le corresponden del conjunto etiquetado. "
            "Cada par lo etiquetan dos anotadores por separado; con sus etiquetas se mide el "
            "acuerdo entre anotadores y se evalúa el sistema.",
        ),
        _COMPARED,
        *_scale_lines(),
        (
            "Ejemplos",
            f"La hoja «{EXAMPLES_SHEET}» muestra los {examples} pares del piloto con la etiqueta "
            "que se acordó en la reunión de calibración y su justificación. Úselos como "
            "referencia.",
        ),
        ("Cómo etiquetar", ""),
        ("1.", f"Vaya a la hoja «{LABELING_SHEET}». Cada par tiene un número y un código."),
        _JUDGE,
        (
            "3.",
            "Elija una etiqueta en «Etiqueta» con la lista desplegable. «Notas (opcional)» "
            "admite cualquier comentario.",
        ),
        (
            "4.",
            "Etiquete todos los pares: la importación rechaza un libro con pares sin etiqueta. "
            "Guarde el archivo con el mismo nombre y en la misma carpeta. Puede trabajar por "
            "partes: lo ya etiquetado se conserva.",
        ),
        (
            "Trabajo individual",
            "No consulte con otros anotadores los pares de este libro ni compare sus etiquetas "
            "hasta que se hayan entregado todos los libros.",
        ),
        *_shared_lines(),
    ]


@dataclass(frozen=True)
class PilotBuild:
    """The pilot workbook and its records, ready to write."""

    codes: dict[str, str]
    new_codes: bool
    """Whether the codes are new; the codes of a pilot run that stopped before its manifest are
    kept."""
    layout: tuple[LayoutRow, ...]
    workbook: openpyxl.Workbook
    lacking: Lacking
    redacted: int
    """Cells in which the name redactor replaced a person's name."""
    texts_checked: int
    """Distinct texts that the name guard looked at."""


def build_pilot(
    config: AppConfig,
    places: Places,
    inputs: WorkbookInputs,
    *,
    choice: Callable[[str], str] = secrets.choice,
) -> PilotBuild:
    """Give each drawn pair its blind code (or keep the codes stored), lay out and build the
    pilot workbook, and pass it through the name guard.

    Raises:
        NameGuardError: if the workbook holds a person's name or an ORCID iD of the snapshot.
        ValueError: if stored codes do not fit the draw.
    """
    identifiers = [pair.pair_id for pair in inputs.pairs]
    path = places.records / BLIND_CODES_FILE_NAME
    if path.exists():
        codes = {row.pair_id: row.code for row in ParquetTableStore().read(BlindCodeRow, path)}
        check_blind_codes(codes, identifiers)
        new = False
    else:
        codes, new = new_blind_codes(identifiers, choice=choice), True
    pilot = [pair for pair in inputs.pairs if pair.phase == PILOT]
    layout = layout_rows(PILOT_WORKBOOK, None, pilot, codes, seed=config.labeling.seed)
    redactor = _NameRedactor(inputs.names)
    workbook = build_workbook(
        pilot_guide(len(layout)),
        [
            SheetContent(
                LABELING_SHEET,
                PILOT_LAYOUT,
                pair_views(layout, inputs.fichas, redactor),
                editable=True,
            )
        ],
    )
    checked = guard_workbooks({PILOT_WORKBOOK_FILE_NAME: workbook}, inputs.guard)
    return PilotBuild(
        codes=codes,
        new_codes=new,
        layout=layout,
        workbook=workbook,
        lacking=Lacking.of(layout, inputs.fichas),
        redacted=redactor.redacted,
        texts_checked=checked,
    )


@dataclass(frozen=True)
class MainBuild:
    """The annotators' workbooks and their records, ready to write."""

    assignments: tuple[AssignmentRow, ...]
    layout: tuple[LayoutRow, ...]
    workbooks: dict[str, openpyxl.Workbook]
    """By file name, annotator 1 first."""
    lacking: Lacking
    """Over the labeled set's theses."""
    redacted: int
    texts_checked: int
    details: dict[str, Any]


def build_main(
    config: AppConfig,
    inputs: WorkbookInputs,
    *,
    annotators: int,
    codes: Mapping[str, str],
    pilot_layout: Sequence[LayoutRow],
    pilot_labels: Sequence[PilotLabelRow],
) -> MainBuild:
    """Assign the labeled set to the pairs of annotators, lay out and build each annotator's
    workbook with the pilot's pairs and agreed labels as examples, and pass them all through the
    name guard.

    Raises:
        NameGuardError: if a workbook holds a person's name or an ORCID iD of the snapshot.
    """
    seed = config.labeling.seed
    labeled = [pair for pair in inputs.pairs if pair.phase == LABELED]
    by_id = {pair.pair_id: pair for pair in labeled}
    assigned = assign_pairs(labeled, annotators, seed=seed)
    redactor = _NameRedactor(inputs.names)
    by_code = {row.code: row for row in pilot_labels}
    examples = pair_views(pilot_layout, inputs.fichas, redactor, by_code)
    layout: list[LayoutRow] = []
    workbooks: dict[str, openpyxl.Workbook] = {}
    for annotator in range(1, annotators + 1):
        mine = [by_id[a.pair_id] for a in assigned if annotator in a.annotators]
        rows = layout_rows(annotator_workbook(annotator), annotator, mine, codes, seed=seed)
        layout += rows
        workbooks[annotator_file_name(annotator)] = build_workbook(
            main_guide(annotator, len(rows), len(examples)),
            [
                SheetContent(EXAMPLES_SHEET, PILOT_LAYOUT, examples, editable=False),
                SheetContent(
                    LABELING_SHEET,
                    MAIN_LAYOUT,
                    pair_views(rows, inputs.fichas, redactor),
                    editable=True,
                ),
            ],
        )
    checked = guard_workbooks(workbooks, inputs.guard)
    return MainBuild(
        assignments=tuple(_assignment_row(a, by_id, codes) for a in assigned),
        layout=tuple(layout),
        workbooks=workbooks,
        lacking=Lacking.of(layout, inputs.fichas),
        redacted=redactor.redacted,
        texts_checked=checked,
        details=_main_details(assigned, by_id, annotators),
    )


def _assignment_row(
    assignment: Assignment, by_id: Mapping[str, PairRow], codes: Mapping[str, str]
) -> AssignmentRow:
    pair = by_id[assignment.pair_id]
    first, second = assignment.annotators
    return AssignmentRow(
        pair_id=pair.pair_id,
        code=codes[pair.pair_id],
        pool=pair.pool,
        source=pair.source,
        group=assignment.group,
        annotator_a=first,
        annotator_b=second,
    )


def _main_details(
    assigned: Sequence[Assignment], by_id: Mapping[str, PairRow], annotators: int
) -> dict[str, Any]:
    """The pairs of each pair of annotators, by pool and by model, and of each annotator: counts
    only."""
    groups = []
    for number, pair in enumerate(annotator_pairs(annotators), start=1):
        mine = [by_id[a.pair_id] for a in assigned if a.group == number]
        pools = Counter(p.pool for p in mine)
        groups.append(
            {
                "group": number,
                "annotators": list(pair),
                "pairs": len(mine),
                "pools": {pool: pools[pool] for pool in POOLS},
                "sources": dict(sorted(Counter(p.source for p in mine if p.source).items())),
            }
        )
    return {
        "groups": groups,
        "annotators": {
            str(k): sum(k in a.annotators for a in assigned) for k in range(1, annotators + 1)
        },
    }


# --- Writing --------------------------------------------------------------------------------


Table = tuple[str, Sequence[BaseModel], type[BaseModel]]
"""A gitignored table to write: its file name, its rows and their model."""


def write_records(
    places: Places,
    tables: Sequence[Table],
    workbooks: Mapping[str, openpyxl.Workbook],
    manifest: WorkbookManifest,
    manifest_name: str,
    guarded: frozenset[str],
    *,
    write_text: Callable[[Path, str], None] = write_text_atomically,
) -> WorkbookManifest:
    """Write ``tables`` and ``workbooks`` (by file name) under the records directory, then the
    manifest last, and return it with the SHA-256 of every file written.

    The privacy guard reads the manifest without its files' digests before anything is written,
    then the final manifest once more, after the other records are written and before it is.

    Raises:
        PrivacyError: if the manifest holds a string of ``guarded``; nothing is written then. If
            only the final manifest does, every record but the manifest is written already.
        OSError: if writing fails; the records written before stay.
    """
    check_privacy({manifest_name: manifest.model_dump_json(indent=2)}, guarded)
    store = ParquetTableStore()
    places.workbooks.mkdir(parents=True, exist_ok=True)
    names: list[str] = []
    for name, rows, model in tables:
        store.write(rows, places.records / name, model_cls=model)
        names.append(name)
    for name, book in workbooks.items():
        save_workbook(book, places.workbooks / name)
        names.append(f"{WORKBOOKS_DIR_NAME}/{name}")
    final = manifest.model_copy(
        update={"data_files": {name: _sha256(places.records / name) for name in names}}
    )
    text = final.model_dump_json(indent=2)
    check_privacy({manifest_name: text}, guarded)
    places.results.mkdir(parents=True, exist_ok=True)
    # The commit point: without it, the step is not complete.
    write_text(places.results / manifest_name, text)
    return final


# --- The steps ------------------------------------------------------------------------------


@dataclass(frozen=True)
class Prepared:
    """A step computed and checked, ready to describe and to write."""

    counts: dict[str, int]
    details: dict[str, Any]
    tables: Sequence[Table]
    workbooks: Mapping[str, openpyxl.Workbook]
    codes: Mapping[str, str]
    lines: list[str]
    """The step's summary, without thesis text."""
    written: str
    """What a real run writes, in words."""
    inputs: dict[str, str] = field(default_factory=dict)
    """The SHA-256 of the earlier steps' records the step read, by name."""


def prepare_pilot(run: "Run", inputs: WorkbookInputs) -> Prepared:
    """Build the pilot workbook and say what it holds."""
    built = build_pilot(run.config, run.places, inputs, choice=run.choice)
    counts = {
        "drawn_pairs": len(inputs.pairs),
        "pairs": len(built.layout),
        **built.lacking.counts(),
        "swapped": sum(row.swapped for row in built.layout),
        "redacted_cells": built.redacted,
        "guard_strings": len(inputs.guard),
        "texts_checked": built.texts_checked,
    }
    codes = [BlindCodeRow(pair_id=i, code=c) for i, c in sorted(built.codes.items())]
    lines = [
        f"Pilot workbook: {len(built.layout)} pairs over {built.lacking.theses} theses, in a "
        f"random order with random sides (seed {run.config.labeling.seed}; {counts['swapped']} "
        "pairs swapped)",
        built.lacking.line(),
        f"  names redacted in {built.redacted} cells",
        f"  blind codes: {len(built.codes)}, one per drawn pair, "
        + ("new" if built.new_codes else f"kept from {BLIND_CODES_FILE_NAME}"),
        _guard_line(built.texts_checked, inputs),
    ]
    return Prepared(
        counts=counts,
        details={},
        tables=[
            (BLIND_CODES_FILE_NAME, codes, BlindCodeRow),
            (PILOT_LAYOUT_FILE_NAME, built.layout, LayoutRow),
        ],
        workbooks={PILOT_WORKBOOK_FILE_NAME: built.workbook},
        codes=built.codes,
        lines=lines,
        written=f"{run.places.relative(run.places.pilot_workbook)} and the gitignored records",
    )


def pilot_records(places: Places) -> tuple[tuple[LayoutRow, ...], dict[str, str]]:
    """Return the pilot's layout and the blind codes, each checked against the pilot's manifest.

    Raises:
        WorkbooksError: if the pilot workbook was never made.
        ValueError: if a record differs from the manifest.
    """
    path = places.results / PILOT_MANIFEST_FILE_NAME
    if not path.is_file():
        raise WorkbooksError(
            f"the pilot workbook of snapshot {places.snapshot_id} was never made ({path} is "
            f"missing); run {RUN_COMMAND} pilot first"
        )
    manifest = _manifest_at(path)
    layout = _read_recorded(LayoutRow, places.records, PILOT_LAYOUT_FILE_NAME, manifest, path)
    codes = _read_recorded(BlindCodeRow, places.records, BLIND_CODES_FILE_NAME, manifest, path)
    return layout, {row.pair_id: row.code for row in codes}


def pilot_labels(places: Places) -> tuple[PilotLabelRow, ...]:
    """Return the pilot's agreed labels, checked against their manifest.

    Raises:
        FileNotFoundError: if they are not stored.
        ValueError: if the record differs from the manifest.
    """
    path = places.results / PILOT_LABELS_MANIFEST_FILE_NAME
    manifest = _manifest_at(path)
    return _read_recorded(PilotLabelRow, places.records, PILOT_LABELS_FILE_NAME, manifest, path)


def prepare_import(run: "Run", inputs: WorkbookInputs) -> Prepared:
    """Read the returned pilot workbook, check it, and turn its labels into records.

    Raises:
        LabelingWorkbookError: listing every problem of the workbook, a level without a pair
            included.
    """
    places = run.places
    layout, codes = pilot_records(places)
    pilot = {pair.pair_id for pair in inputs.pairs if pair.phase == PILOT}
    if {row.pair_id for row in layout} != pilot:
        raise WorkbooksError("the pilot workbook's pairs are not the draw's pilot pairs")
    path = places.pilot_workbook
    labels = read_labels(path, PILOT_LAYOUT, {row.code: row.n for row in layout})
    lacking = missing_levels(labels)
    if lacking:
        raise LabelingWorkbookError(
            [
                f"no pair is labeled '{label}'; the pilot needs at least one pair of each level "
                "(D34 (7))"
                for label in lacking
            ]
        )
    by_code = {row.code: row for row in layout}
    rows = tuple(
        PilotLabelRow(
            n=label.n,
            code=label.code,
            pair_id=by_code[label.code].pair_id,
            level=label.level,
            binary=BINARY[label.level],
            justification=label.note,
        )
        for label in sorted(labels, key=lambda label: label.n)
    )
    levels = Counter(row.level for row in rows)
    digest = _sha256(path)
    counts = {
        "pairs": len(rows),
        **{f"level_{level}": levels[level] for level in LEVEL_DEFINITIONS},
        "binary_redundant": sum(row.binary for row in rows),
        "with_justification": sum(row.justification is not None for row in rows),
    }
    lines = [
        f"Pilot labels: {len(rows)} pairs; "
        + ", ".join(f"{label} {levels[level]}" for label, level in LEVELS.items())
        + f"; {counts['with_justification']} with a justification",
        f"  {PILOT_WORKBOOK_FILE_NAME} SHA-256 {digest}",
    ]
    return Prepared(
        counts=counts,
        details={"levels": {label: levels[level] for label, level in LEVELS.items()}},
        tables=[(PILOT_LABELS_FILE_NAME, rows, PilotLabelRow)],
        workbooks={},
        codes=codes,
        lines=lines,
        written=f"the gitignored {PILOT_LABELS_FILE_NAME}",
        inputs={
            PILOT_MANIFEST_FILE_NAME: _sha256(places.results / PILOT_MANIFEST_FILE_NAME),
            PILOT_WORKBOOK_FILE_NAME: digest,
        },
    )


def prepare_main(run: "Run", inputs: WorkbookInputs) -> Prepared:
    """Build the annotators' workbooks and say what they hold."""
    places = run.places
    annotators = run.config.labeling.annotators
    assert annotators is not None
    pilot_layout, codes = pilot_records(places)
    check_blind_codes(codes, [pair.pair_id for pair in inputs.pairs])
    labels = pilot_labels(places)
    built = build_main(
        run.config,
        inputs,
        annotators=annotators,
        codes=codes,
        pilot_layout=pilot_layout,
        pilot_labels=labels,
    )
    groups = built.details["groups"]
    counts = {
        "annotators": annotators,
        "groups": len(groups),
        "pairs": len(built.assignments),
        **built.lacking.counts(),
        "examples": len(labels),
        "redacted_cells": built.redacted,
        "guard_strings": len(inputs.guard),
        "texts_checked": built.texts_checked,
    }
    lines = [
        f"Main workbooks: {annotators} annotators, {len(built.assignments)} pairs over "
        f"{len(groups)} pairs of annotators, each pair labeled by both annotators of its pair",
        *(
            f"  annotators {group['annotators'][0]} and {group['annotators'][1]}: "
            f"{group['pairs']} pairs ("
            + ", ".join(f"{pool} {count}" for pool, count in group["pools"].items())
            + ")"
            for group in groups
        ),
        "  pairs per annotator: "
        + ", ".join(f"{k} {count}" for k, count in built.details["annotators"].items()),
        f"  the {len(labels)} pilot pairs, with their agreed labels, as examples in every workbook",
        built.lacking.line(),
        f"  names redacted in {built.redacted} cells",
        _guard_line(built.texts_checked, inputs),
    ]
    return Prepared(
        counts=counts,
        details=built.details,
        tables=[
            (ASSIGNMENT_FILE_NAME, built.assignments, AssignmentRow),
            (MAIN_LAYOUT_FILE_NAME, built.layout, LayoutRow),
        ],
        workbooks=built.workbooks,
        codes=codes,
        lines=lines,
        written=f"{len(built.workbooks)} workbooks under {places.relative(places.workbooks)} and "
        "the gitignored records",
        inputs={
            PILOT_MANIFEST_FILE_NAME: _sha256(places.results / PILOT_MANIFEST_FILE_NAME),
            PILOT_LABELS_MANIFEST_FILE_NAME: _sha256(
                places.results / PILOT_LABELS_MANIFEST_FILE_NAME
            ),
        },
    )


@dataclass(frozen=True)
class Step:
    """One step of the command: what it writes, its refusal, its manifest and how it prepares."""

    name: str
    manifest_name: str
    outputs: Callable[["Run"], tuple[Path, ...]]
    """The records whose presence means that the step already ran, or began to."""
    exists: str
    """The refusal's start, with ``{snapshot}`` and ``{path}``."""
    prepare: Callable[["Run", WorkbookInputs], Prepared]


def _main_outputs(run: "Run") -> tuple[Path, ...]:
    places, annotators = run.places, run.config.labeling.annotators or 0
    return (
        places.records / ASSIGNMENT_FILE_NAME,
        places.records / MAIN_LAYOUT_FILE_NAME,
        *(places.workbooks / annotator_file_name(k) for k in range(1, annotators + 1)),
        places.results / MAIN_MANIFEST_FILE_NAME,
    )


STEPS = {
    "pilot": Step(
        name="pilot",
        manifest_name=PILOT_MANIFEST_FILE_NAME,
        outputs=lambda run: (
            run.places.records / PILOT_LAYOUT_FILE_NAME,
            run.places.pilot_workbook,
            run.places.results / PILOT_MANIFEST_FILE_NAME,
        ),
        exists="the pilot workbook of snapshot {snapshot} already exists ({path}) and may hold "
        "labels",
        prepare=prepare_pilot,
    ),
    "import-pilot": Step(
        name="import-pilot",
        manifest_name=PILOT_LABELS_MANIFEST_FILE_NAME,
        outputs=lambda run: (
            run.places.records / PILOT_LABELS_FILE_NAME,
            run.places.results / PILOT_LABELS_MANIFEST_FILE_NAME,
        ),
        exists="the pilot's agreed labels of snapshot {snapshot} are already stored ({path})",
        prepare=prepare_import,
    ),
    "main": Step(
        name="main",
        manifest_name=MAIN_MANIFEST_FILE_NAME,
        outputs=_main_outputs,
        exists="the main workbooks of snapshot {snapshot} already exist ({path}) and may hold "
        "labels",
        prepare=prepare_main,
    ),
}
"""The command's steps, in the order they run."""


@dataclass(frozen=True)
class Run:
    """What every step needs: the configuration, the places and the seams."""

    config: AppConfig
    places: Places
    dry_run: bool
    load: Callable[[AppConfig, Path, str], WorkbookInputs]
    git: Callable[[Path, str], GitState]
    clock: Callable[[], datetime]
    timer: Callable[[], float]
    choice: Callable[[str], str]


def _parse_arguments(argv: Sequence[str] | None) -> argparse.Namespace:
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument(
        "--project-root",
        type=Path,
        default=Path(),
        help=f"project root that holds {CONFIG_PATH.as_posix()} (default: current directory)",
    )
    common.add_argument("--snapshot-id", help="snapshot to use (default: the harvested last)")
    common.add_argument(
        "--dry-run",
        action="store_true",
        help="compute and check everything, apply the name and privacy guards, and write nothing",
    )
    parser = argparse.ArgumentParser(
        prog="python -m thematic_redundancy.labeling.pair_workbooks",
        description=(
            "Make the labeled set's workbooks of a snapshot (T20b, D34) from its draw, and import "
            "the pilot's agreed labels, in three steps run in this order: pilot, import-pilot, "
            "main. Each step happens once: its workbooks and records, which name theses, go under "
            "labeling.records_dir/<snapshot>/ (gitignored), and its manifest, counts and digests "
            "only, under results/labeled_set/<snapshot>/workbooks/ (tracked by git). Nothing is "
            "ever overwritten."
        ),
    )
    steps = parser.add_subparsers(dest="step", required=True, metavar="STEP")
    steps.add_parser(
        "pilot",
        parents=[common],
        help="make the pilot workbook and give every drawn pair its blind code",
        description=(
            f"Make the pilot workbook ({PILOT_WORKBOOK_FILE_NAME}): its pairs in a random order "
            "with random sides, the four fields of both fichas side by side, a drop-down label "
            "and a justification; and give every drawn pair its blind code."
        ),
    )
    steps.add_parser(
        "import-pilot",
        parents=[common],
        help="store the agreed labels of the pilot workbook returned from the calibration meeting",
        description=(
            "Read the pilot workbook back, saved under its name in its folder, and store its "
            "agreed labels and justifications once it holds its pairs unchanged, a label on every "
            "pair and at least one pair of each level (D34 (7))."
        ),
    )
    steps.add_parser(
        "main",
        parents=[common],
        help="make one workbook per annotator, once labeling.annotators and the pilot are set",
        description=(
            "Make one workbook per annotator (etiquetado_anotador_<k>.xlsx) with the pairs of the "
            "annotator's pairs of annotators, spread evenly (D34 (8)), and the pilot's agreed "
            "labels as examples. It refuses until labeling.annotators (3 or 4) is set and the "
            "pilot's labels are imported."
        ),
    )
    return parser.parse_args(argv)


def _progress(line: str) -> None:
    print(line, file=sys.stderr, flush=True)


def _now() -> datetime:
    return datetime.now(UTC).replace(microsecond=0)


def main(
    argv: Sequence[str] | None = None,
    *,
    load: Callable[[AppConfig, Path, str], WorkbookInputs] = load_workbook_inputs,
    git: Callable[[Path, str], GitState] = read_git_state,
    clock: Callable[[], datetime] = _now,
    timer: Callable[[], float] = time.perf_counter,
    choice: Callable[[str], str] = secrets.choice,
) -> int:
    """Run the command line and return its exit code: 0 on success, 1 after an error, and 130
    after an interruption."""
    arguments = _parse_arguments(argv)
    root = arguments.project_root.resolve()
    try:
        config = load_config(root / CONFIG_PATH)
        paths = config.paths.resolve_against(root)
        snapshot_id = check_snapshot_id(
            arguments.snapshot_id or latest_snapshot_id(paths.data_dir / RAW_DIR_NAME)
        )
        run = Run(
            config=config,
            places=Places.of(config, root, snapshot_id),
            dry_run=arguments.dry_run,
            load=load,
            git=git,
            clock=clock,
            timer=timer,
            choice=choice,
        )
        lines = _run_step(run, STEPS[arguments.step])
    except KeyboardInterrupt:
        print(
            "interrupted: a step is complete only once its manifest is written; move its partial "
            "records away, if any, and run the command again",
            file=sys.stderr,
        )
        return 130
    except LabelingWorkbookError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    except (OSError, ValueError, yaml.YAMLError) as error:
        print(f"error: {' '.join(str(error).split())}", file=sys.stderr)
        return 1
    print("\n".join(lines))
    return 0


def _run_step(run: Run, step: Step) -> list[str]:
    """Run one step: wait for what it needs, refuse a dirty tree and earlier records (a real
    run), hold the lock (a real run), prepare it, check its manifest, and write it (a real
    run)."""
    places = run.places
    relative = places.relative(places.results)
    if step.name == "main":
        _wait_for_the_main_workbooks(run)
    outputs = step.outputs(run)
    if run.dry_run:
        try:
            state: GitState | None = run.git(places.root, relative)
        except (OSError, UnicodeError):  # what read_git_state raises when git cannot tell
            state = None
        prepared, _, _ = _prepare(run, step, state)
        earlier = next((path for path in outputs if path.exists()), None)
        return [
            f"Dry run of the step '{step.name}' on snapshot {places.snapshot_id}: everything is "
            "computed and checked, and nothing is written",
            f"Git: {_git_text(state, relative)}",
            *(
                []
                if earlier is None
                else [f"This step's records exist ({earlier.name}); a real run refuses them"]
            ),
            *prepared.lines,
            f"Privacy guard: {step.manifest_name} checked, no hit (it is checked again when "
            "written)",
            f"Command of the real run: {RUN_COMMAND} {step.name} --snapshot-id "
            f"{places.snapshot_id}",
        ]
    state = run.git(places.root, relative)
    if state.dirty:
        raise WorkbooksError(
            f"the git tree has {len(state.dirty)} changed or untracked paths outside {relative}; "
            "commit them first, since the manifest records the commit"
        )
    _refuse(step, outputs, places)
    places.results.mkdir(parents=True, exist_ok=True)
    busy = WorkbooksError(f"another run holds the lock of {relative}; let it finish first")
    with exclusive_run(places.results, busy):
        _refuse(step, outputs, places)
        prepared, manifest, guarded = _prepare(run, step, state)
        _progress(f"writing the records, {step.manifest_name} last")
        try:
            final = write_records(
                places, prepared.tables, prepared.workbooks, manifest, step.manifest_name, guarded
            )
        except (OSError, ValueError) as error:
            partial = [path for path in outputs if path.exists()]
            if not partial:
                raise
            hint = _partial_hint(partial, step.manifest_name)
            raise PartialWorkbooksError(f"{error}; {hint}") from error
    return [
        f"Snapshot {places.snapshot_id}: step '{step.name}', commit {final.git_commit}",
        *prepared.lines,
        f"Wrote {prepared.written}, and {relative}/{step.manifest_name} last",
    ]


def _wait_for_the_main_workbooks(run: Run) -> None:
    """Refuse the main workbooks until the number of annotators is set and the pilot's agreed
    labels are stored (D34 (7), (8))."""
    problems = []
    if run.config.labeling.annotators is None:
        problems.append(
            "labeling.annotators is not set: set it to 3 or 4 in config/default.yaml once the "
            "number of annotators is confirmed (U3, D34 (8))"
        )
    if not (run.places.results / PILOT_LABELS_MANIFEST_FILE_NAME).is_file():
        problems.append(
            f"the pilot's agreed labels are not stored: run {RUN_COMMAND} import-pilot first "
            "(D34 (7))"
        )
    if problems:
        raise WorkbooksError("the main workbooks wait: " + "; ".join(problems))


def _prepare(
    run: Run, step: Step, state: GitState | None
) -> tuple[Prepared, WorkbookManifest, frozenset[str]]:
    """Load the inputs, prepare the step, and check its manifest without its files' digests."""
    started_at, started = run.clock(), run.timer()
    _progress("loading the draw, the fichas and the snapshot")
    inputs = run.load(run.config, run.places.root, run.places.snapshot_id)
    _progress(f"preparing the step '{step.name}'")
    prepared = step.prepare(run, inputs)
    seconds = run.timer() - started
    labeling = run.config.labeling
    manifest = WorkbookManifest(
        step=step.name,
        snapshot_id=run.places.snapshot_id,
        git_commit="unknown" if state is None else state.commit,
        seed=labeling.seed,
        settings={
            "seed": labeling.seed,
            "annotators": labeling.annotators,
            "records_dir": labeling.records_dir.as_posix(),
        },
        counts=prepared.counts,
        details=prepared.details,
        inputs={
            **inputs.input_sha256,
            **prepared.inputs,
            CONFIG_PATH.name: _sha256(run.places.root / CONFIG_PATH),
        },
        data_files={},
        versions=library_versions(),
        started_at=started_at,
        finished_at=run.clock(),
        seconds=round(max(seconds, 0.0), 3),
    )
    guarded = guarded_strings(inputs, prepared.codes)
    check_privacy({step.manifest_name: manifest.model_dump_json(indent=2)}, guarded)
    return prepared, manifest, guarded


def guarded_strings(inputs: WorkbookInputs, codes: Mapping[str, str]) -> frozenset[str]:
    """Return the strings no tracked manifest may hold: the snapshot's, the included theses' doc
    codes, the pair ids and the blind codes."""
    return (
        inputs.sensitive
        | frozenset(inputs.fichas)
        | frozenset(pair.pair_id for pair in inputs.pairs)
        | frozenset(codes.values())
    )


def _refuse(step: Step, outputs: Sequence[Path], places: Places) -> None:
    earlier = next((path for path in outputs if path.exists()), None)
    if earlier is not None:
        raise FileExistsError(
            step.exists.format(snapshot=places.snapshot_id, path=earlier)
            + "; a step happens once, so it is never overwritten: move its records away first "
            "to run it again"
        )


def _git_text(state: GitState | None, relative: str) -> str:
    if state is None:
        return "not readable here; a real run needs a repository with a commit"
    if not state.dirty:
        return f"commit {state.commit}, clean"
    count = len(state.dirty)
    return (
        f"commit {state.commit}, {count} changed path{'s' if count > 1 else ''} outside "
        f"{relative}; a real run refuses them until they are committed"
    )


def _guard_line(checked: int, inputs: WorkbookInputs) -> str:
    return (
        f"Name guard: {checked:,} distinct texts checked against {len(inputs.guard):,} strings "
        "(the snapshot's person-name forms and ORCID iDs), without word boundaries; no hit"
    )


def _partial_hint(partial: Sequence[Path], manifest_name: str) -> str:
    names = [path.name for path in partial]
    listed = ", ".join(names[:-1]) + f" and {names[-1]}" if len(names) > 1 else names[0]
    folders = ", ".join(dict.fromkeys(str(path.parent) for path in partial))
    return (
        f"the step stopped after writing {listed}, and a step is complete only once "
        f"{manifest_name} is written: move the partial records away ({folders}) and run the "
        "command again"
    )


def _manifest_at(path: Path) -> WorkbookManifest:
    try:
        return WorkbookManifest.model_validate_json(path.read_bytes())
    except ValidationError as error:
        raise ValueError(f"{path} is not a valid workbook manifest") from error


def _read_recorded[M: BaseModel](
    model: type[M], records: Path, name: str, manifest: WorkbookManifest, manifest_path: Path
) -> tuple[M, ...]:
    digest = manifest.data_files.get(name)
    if digest is None:
        raise ValueError(f"{manifest_path} records no SHA-256 of {name}")
    return _read_verified(model, records / name, digest, manifest_path.name)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


if __name__ == "__main__":
    raise SystemExit(main())
