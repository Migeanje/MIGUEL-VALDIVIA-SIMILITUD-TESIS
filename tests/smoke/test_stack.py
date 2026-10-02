"""Smoke checks that the locked third-party stack works together on this machine.

These tests exercise library compatibility, not project code. They are marked ``smoke``,
the default run excludes them, and ``uv run pytest -m smoke`` runs them. The first run
downloads the first configured embedding model (about 0.5 GB) into the Hugging Face cache.

Heavy libraries are imported inside fixtures and tests, so collecting this module during
the default run stays fast and offline.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, NamedTuple

import numpy as np
import pytest

from thematic_redundancy.shared.config import MODEL_MAX_SEQ_LENGTH, AppConfig, load_config

if TYPE_CHECKING:
    from bertopic import BERTopic
    from sentence_transformers import SentenceTransformer
    from sklearn.base import ClusterMixin
    from sklearn.feature_extraction.text import CountVectorizer

pytestmark = pytest.mark.smoke

DEFAULT_CONFIG_PATH = Path(__file__).resolve().parents[2] / "config" / "default.yaml"

SPACY_MODEL = "es_core_news_md"

EMBEDDING_DIMENSION = 384
"""Output size of the first configured model, paraphrase-multilingual-MiniLM-L12-v2."""

TOP_WORDS = 10
"""c-TF-IDF terms kept per topic; NPMI coherence is computed on them (D12)."""

TOPIC_INFO_COLUMNS = ("Topic", "Count", "Name", "Representation", "Representative_Docs")


class ThemeTemplate(NamedTuple):
    """Phrases that are combined into the documents of one synthetic theme.

    Document ``i`` joins ``subjects[i % 5]``, ``effects[i % 3]`` and ``settings[i // 5]``.
    Five and three are coprime, so the fifteen documents of a theme are all distinct.
    """

    subjects: tuple[str, str, str, str, str]
    effects: tuple[str, str, str]
    settings: tuple[str, str, str]


DOCUMENTS_PER_THEME = 15

# The themes share no word once spaCy's Spanish stop words are removed, so every
# c-TF-IDF top word belongs to exactly one theme.
THEME_TEMPLATES: dict[str, ThemeTemplate] = {
    "mining_blasting": ThemeTemplate(
        subjects=(
            "El diseño de la malla de voladura",
            "La carga de explosivo ANFO",
            "El modelo de fragmentación de Kuz-Ram",
            "La secuencia de retardos de los detonadores",
            "El ajuste del burden y el espaciamiento",
        ),
        effects=(
            "reduce la fragmentación gruesa del mineral",
            "controla la vibración del macizo rocoso",
            "disminuye la proyección de rocas tras el disparo",
        ),
        settings=(
            "en la mina a tajo abierto",
            "en la galería subterránea",
            "en la cantera de caliza",
        ),
    ),
    "industrial_maintenance": ThemeTemplate(
        subjects=(
            "El plan de mantenimiento preventivo",
            "El monitoreo de condición de los rodamientos",
            "La gestión de repuestos críticos",
            "El indicador de eficiencia global OEE",
            "El mantenimiento centrado en la confiabilidad",
        ),
        effects=(
            "aumenta la disponibilidad de los equipos",
            "evita paradas no programadas de la producción",
            "prolonga la vida útil de los motores eléctricos",
        ),
        settings=(
            "en la planta de envasado",
            "en la línea de producción de cemento",
            "en el taller de máquinas industriales",
        ),
    ),
    "erp_software": ThemeTemplate(
        subjects=(
            "La implementación del sistema ERP",
            "El módulo de inventarios del ERP",
            "La migración de datos contables al ERP",
            "La integración del ERP con la facturación electrónica",
            "La parametrización del módulo de compras",
        ),
        effects=(
            "automatiza los procesos administrativos",
            "integra la información de ventas y almacenes",
            "mejora la trazabilidad de los pedidos de clientes",
        ),
        settings=(
            "en la empresa comercializadora",
            "en la pyme textil",
            "en la distribuidora de alimentos",
        ),
    ),
}


@dataclass(frozen=True)
class SyntheticCorpus:
    """Short Spanish documents, each tagged with the theme that produced it."""

    documents: tuple[str, ...]
    themes: tuple[str, ...]


def build_corpus() -> SyntheticCorpus:
    """Build the deterministic corpus: fifteen documents for each theme, in theme order."""
    documents: list[str] = []
    themes: list[str] = []
    for theme, template in THEME_TEMPLATES.items():
        for index in range(DOCUMENTS_PER_THEME):
            subject = template.subjects[index % len(template.subjects)]
            effect = template.effects[index % len(template.effects)]
            setting = template.settings[index // len(template.subjects)]
            documents.append(f"{subject} {effect} {setting}.")
            themes.append(theme)
    return SyntheticCorpus(documents=tuple(documents), themes=tuple(themes))


def spanish_vectorizer() -> CountVectorizer:
    """Return the bag-of-words model behind c-TF-IDF, without spaCy's Spanish stop words."""
    from sklearn.feature_extraction.text import CountVectorizer
    from spacy.lang.es.stop_words import STOP_WORDS

    return CountVectorizer(stop_words=sorted(STOP_WORDS))


def fit_topic_model(
    corpus: SyntheticCorpus,
    embeddings: np.ndarray,
    cluster_model: ClusterMixin,
    config: AppConfig,
    seed: int,
) -> BERTopic:
    """Reduce precomputed embeddings with UMAP, then cluster them inside BERTopic."""
    from bertopic import BERTopic
    from umap import UMAP

    umap_model = UMAP(
        n_neighbors=config.umap.n_neighbors,
        n_components=config.umap.n_components,
        min_dist=config.umap.min_dist,
        metric=config.umap.metric,
        random_state=seed,
        n_jobs=1,  # A seeded UMAP runs single-threaded anyway; asking for it avoids a warning.
    )
    topic_model = BERTopic(
        embedding_model=None,  # D11: BERTopic only receives precomputed embeddings.
        # With no embedding model, the default language "english" makes BERTopic strip every
        # character outside [A-Za-z0-9 ] before c-TF-IDF, so "fragmentación" would become
        # "fragmentacin". Any other language keeps accents and ñ.
        language="spanish",
        umap_model=umap_model,
        hdbscan_model=cluster_model,
        vectorizer_model=spanish_vectorizer(),
        top_n_words=TOP_WORDS,
    )
    return topic_model.fit(list(corpus.documents), embeddings=embeddings)


def fit_kmeans_topic_model(
    corpus: SyntheticCorpus, embeddings: np.ndarray, config: AppConfig
) -> BERTopic:
    """Fit the K-means topic map with one cluster per theme and the first configured seed."""
    from sklearn.cluster import KMeans

    seed = config.seeds[0]
    cluster_model = KMeans(n_clusters=len(THEME_TEMPLATES), random_state=seed)
    return fit_topic_model(corpus, embeddings, cluster_model, config, seed)


def top_words(topic_model: BERTopic, topic: int) -> list[str]:
    """Return the c-TF-IDF top words of ``topic``, without BERTopic's empty padding."""
    return [word for word, _ in topic_model.get_topic(topic) if word]


def themes_containing(words: list[str], vocabularies: dict[str, frozenset[str]]) -> list[str]:
    """Return the themes whose vocabulary holds every one of ``words``."""
    return [theme for theme, vocabulary in vocabularies.items() if set(words) <= vocabulary]


@pytest.fixture(scope="module")
def config() -> AppConfig:
    """Return the declared run configuration."""
    return load_config(DEFAULT_CONFIG_PATH)


@pytest.fixture(scope="module")
def sentence_model(config: AppConfig) -> SentenceTransformer:
    """Load the first configured model; the first run downloads it from the Hugging Face Hub."""
    from sentence_transformers import SentenceTransformer

    return SentenceTransformer(config.embedding.models[0], device="cpu")


@pytest.fixture(scope="module")
def corpus() -> SyntheticCorpus:
    """Return the synthetic corpus of three clearly separated themes."""
    built = build_corpus()
    assert len(set(built.documents)) == len(built.documents), "documents must be distinct"
    return built


@pytest.fixture(scope="module")
def corpus_embeddings(sentence_model: SentenceTransformer, corpus: SyntheticCorpus) -> np.ndarray:
    """Return one L2-normalized embedding per document."""
    return sentence_model.encode(list(corpus.documents), normalize_embeddings=True)


@pytest.fixture(scope="module")
def theme_vocabularies(corpus: SyntheticCorpus) -> dict[str, frozenset[str]]:
    """Map each theme to the c-TF-IDF tokens of its documents; the themes share none."""
    analyze = spanish_vectorizer().build_analyzer()
    vocabularies: dict[str, set[str]] = {theme: set() for theme in THEME_TEMPLATES}
    for document, theme in zip(corpus.documents, corpus.themes, strict=True):
        vocabularies[theme].update(analyze(document))
    every_token = [token for vocabulary in vocabularies.values() for token in vocabulary]
    assert len(every_token) == len(set(every_token)), "themes must not share tokens"
    return {theme: frozenset(vocabulary) for theme, vocabulary in vocabularies.items()}


@pytest.fixture(scope="module")
def kmeans_topic_model(
    corpus: SyntheticCorpus, corpus_embeddings: np.ndarray, config: AppConfig
) -> BERTopic:
    """Return the K-means topic map of the synthetic corpus."""
    return fit_kmeans_topic_model(corpus, corpus_embeddings, config)


@pytest.fixture(scope="module")
def hdbscan_topic_model(
    corpus: SyntheticCorpus, corpus_embeddings: np.ndarray, config: AppConfig
) -> BERTopic:
    """Return the HDBSCAN topic map, using the smallest configured ``min_cluster_size``."""
    from hdbscan import HDBSCAN

    cluster_model = HDBSCAN(
        min_cluster_size=min(config.hdbscan.min_cluster_sizes),
        metric="euclidean",
        cluster_selection_method="eom",
        prediction_data=True,
    )
    return fit_topic_model(corpus, corpus_embeddings, cluster_model, config, config.seeds[0])


def _plotly_chart_page() -> None:
    """Tiny Streamlit page; AppTest runs the body of this function as a script."""
    import plotly.graph_objects as go
    import streamlit as st

    st.title("Mapa temático")
    st.plotly_chart(go.Figure(go.Bar(x=["voladura", "mantenimiento", "ERP"], y=[15, 15, 15])))


def test_spacy_spanish_model_lemmatizes_and_keeps_accents_and_enye() -> None:
    import spacy

    nlp = spacy.load(SPACY_MODEL)
    doc = nlp("Los ingenieros diseñaron la simulación del mantenimiento en la compañía minera.")
    lemmas = [token.lemma_ for token in doc if not token.is_punct]

    assert lemmas
    assert all(lemma.strip() for lemma in lemmas)
    assert {"ingeniero", "diseñar", "simulación", "compañía", "minero"} <= set(lemmas)


def test_first_configured_embedding_model_has_the_expected_limits(
    sentence_model: SentenceTransformer,
) -> None:
    assert sentence_model.max_seq_length == MODEL_MAX_SEQ_LENGTH
    assert sentence_model.get_embedding_dimension() == EMBEDDING_DIMENSION


def test_normalized_embeddings_rank_a_paraphrase_above_an_unrelated_sentence(
    sentence_model: SentenceTransformer,
) -> None:
    sentences = [
        "El mantenimiento preventivo reduce las fallas de las máquinas.",
        "Las averías de los equipos disminuyen gracias al mantenimiento planificado.",
        "La receta de la abuela lleva maíz morado y canela.",
    ]

    embeddings = sentence_model.encode(sentences, normalize_embeddings=True)

    assert embeddings.shape == (len(sentences), EMBEDDING_DIMENSION)
    np.testing.assert_allclose(np.linalg.norm(embeddings, axis=1), 1.0, rtol=0, atol=1e-5)
    anchor, paraphrase, unrelated = embeddings
    assert anchor @ paraphrase > anchor @ unrelated


def test_kmeans_topic_map_has_one_topic_per_theme(
    kmeans_topic_model: BERTopic,
    corpus: SyntheticCorpus,
    theme_vocabularies: dict[str, frozenset[str]],
) -> None:
    import pandas as pd

    info = kmeans_topic_model.get_topic_info()

    assert isinstance(info, pd.DataFrame)
    assert set(TOPIC_INFO_COLUMNS) <= set(info.columns)
    topics = sorted(info["Topic"].tolist())
    assert topics == list(range(len(THEME_TEMPLATES)))
    assert info["Count"].sum() == len(corpus.documents)
    words = {topic: top_words(kmeans_topic_model, topic) for topic in topics}
    matches = {topic: themes_containing(words[topic], theme_vocabularies) for topic in topics}
    assert all(len(themes) == 1 for themes in matches.values()), words
    assert sorted(themes[0] for themes in matches.values()) == sorted(THEME_TEMPLATES)


def test_hdbscan_topic_map_finds_at_least_two_topics(
    hdbscan_topic_model: BERTopic, corpus: SyntheticCorpus
) -> None:
    import pandas as pd

    info = hdbscan_topic_model.get_topic_info()

    assert isinstance(info, pd.DataFrame)
    assert set(TOPIC_INFO_COLUMNS) <= set(info.columns)
    assert (info["Topic"] >= 0).sum() >= 2  # Topic -1, when present, holds the outliers.
    assert info["Count"].sum() == len(corpus.documents)


def test_clustering_metrics_and_npmi_coherence_are_finite(
    kmeans_topic_model: BERTopic, corpus: SyntheticCorpus, corpus_embeddings: np.ndarray
) -> None:
    from gensim.corpora import Dictionary
    from gensim.models.coherencemodel import CoherenceModel
    from sklearn.metrics import calinski_harabasz_score, davies_bouldin_score, silhouette_score
    from sklearn.preprocessing import normalize

    labels = np.asarray(kmeans_topic_model.topics_)
    unit_embeddings = normalize(corpus_embeddings)  # D12: DB and CH use L2-normalized vectors.
    analyze = kmeans_topic_model.vectorizer_model.build_analyzer()
    texts = [analyze(document) for document in corpus.documents]
    coherence_model = CoherenceModel(
        topics=[top_words(kmeans_topic_model, topic) for topic in sorted(set(labels))],
        texts=texts,
        dictionary=Dictionary(texts),
        coherence="c_npmi",
        topn=TOP_WORDS,
        processes=1,  # The default spawns cpu_count - 1 workers, which costs seconds on Windows.
    )

    scores = {
        "silhouette": silhouette_score(corpus_embeddings, labels, metric="cosine"),
        "davies_bouldin": davies_bouldin_score(unit_embeddings, labels),
        "calinski_harabasz": calinski_harabasz_score(unit_embeddings, labels),
        "npmi": coherence_model.get_coherence(),
    }

    assert all(math.isfinite(score) for score in scores.values()), scores
    assert -1.0 <= scores["silhouette"] <= 1.0
    # gensim adds an epsilon inside its logarithms, so NPMI can overshoot 1 by about 1e-11.
    assert abs(scores["npmi"]) <= 1.0 + 1e-9
    assert scores["davies_bouldin"] >= 0.0
    assert scores["calinski_harabasz"] > 0.0


def test_bertopic_barchart_is_a_plotly_figure(kmeans_topic_model: BERTopic) -> None:
    import plotly.graph_objects as go

    figure = kmeans_topic_model.visualize_barchart()

    assert isinstance(figure, go.Figure)
    assert len(figure.data) == len(THEME_TEMPLATES)  # One bar trace per topic.


def test_streamlit_renders_a_plotly_chart_under_apptest() -> None:
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_function(_plotly_chart_page, default_timeout=60)
    app.run()

    assert not app.exception
    assert len(app.get("plotly_chart")) == 1


def test_pingouin_cronbach_alpha_matches_the_textbook_formula() -> None:
    import pandas as pd
    import pingouin as pg

    answers = pd.DataFrame(
        {
            "item_1": [4, 5, 3, 4, 2, 5],
            "item_2": [4, 4, 3, 5, 2, 5],
            "item_3": [3, 5, 2, 4, 3, 4],
            "item_4": [4, 5, 3, 4, 2, 4],
        }
    )
    item_count = answers.shape[1]
    item_variance = answers.var(ddof=1).sum()
    total_variance = answers.sum(axis=1).var(ddof=1)
    expected = item_count / (item_count - 1) * (1 - item_variance / total_variance)

    alpha, (low, high) = pg.cronbach_alpha(data=answers)

    assert alpha == pytest.approx(expected)
    assert low <= alpha <= high


def test_statsmodels_fleiss_kappa_matches_a_hand_computed_table() -> None:
    from statsmodels.stats.inter_rater import fleiss_kappa

    # Four pairs, three annotators, two categories. Pair agreement is 1, 1, 1/3 and 1/3,
    # so the mean is 2/3; both categories hold half of the ratings, so chance agreement
    # is 1/2. Kappa is (2/3 - 1/2) / (1 - 1/2) = 1/3.
    ratings = np.array([[3, 0], [0, 3], [2, 1], [1, 2]])

    assert fleiss_kappa(ratings) == pytest.approx(1 / 3)


def test_scipy_bootstrap_interval_brackets_the_sample_mean(config: AppConfig) -> None:
    from scipy.stats import bootstrap

    sample = np.array([0.62, 0.71, 0.68, 0.75, 0.66, 0.73, 0.70, 0.64, 0.69, 0.72])

    result = bootstrap(
        (sample,),
        np.mean,
        confidence_level=0.95,
        n_resamples=999,
        rng=np.random.default_rng(config.seeds[0]),
    )

    low, high = result.confidence_interval
    assert math.isfinite(low) and math.isfinite(high)
    assert low < sample.mean() < high


def test_openpyxl_round_trips_a_list_data_validation(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    choices = '"redundante,relacionado,sin relación"'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "pairs"
    sheet.append(["pair_id", "label"])
    sheet.append(["P001", "sin relación"])
    validation = DataValidation(type="list", formula1=choices, allow_blank=False)
    validation.add("B2:B121")
    sheet.add_data_validation(validation)
    path = tmp_path / "annotator.xlsx"
    workbook.save(path)

    loaded = load_workbook(path)["pairs"]

    (loaded_validation,) = loaded.data_validations.dataValidation
    assert loaded_validation.type == "list"
    assert loaded_validation.formula1 == choices
    assert str(loaded_validation.sqref) == "B2:B121"
    assert loaded["B2"].value == "sin relación"


def test_pandas_parquet_round_trip_through_pyarrow(tmp_path: Path) -> None:
    import pandas as pd

    frame = pd.DataFrame(
        {
            "doc_id": ["T0001", "T0002"],
            "title": ["Diseño de voladura en la mina", "Gestión del mantenimiento en la compañía"],
            "year": [2021, 2025],
            "similarity": [0.71, 0.42],
        }
    )
    path = tmp_path / "fichas.parquet"

    frame.to_parquet(path, engine="pyarrow", index=False)

    pd.testing.assert_frame_equal(pd.read_parquet(path, engine="pyarrow"), frame)


def test_kmeans_topic_map_is_reproducible_with_the_same_seed(
    kmeans_topic_model: BERTopic,
    corpus: SyntheticCorpus,
    corpus_embeddings: np.ndarray,
    config: AppConfig,
) -> None:
    rerun = fit_kmeans_topic_model(corpus, corpus_embeddings, config)

    np.testing.assert_array_equal(
        rerun.umap_model.embedding_, kmeans_topic_model.umap_model.embedding_
    )
    assert rerun.topics_ == kmeans_topic_model.topics_
