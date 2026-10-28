"""The labeled set's workbooks (T20b): the scale, the blind codes, the display of a ficha, the
order and sides, the sheets, the name guard and the reader of a returned workbook.

Every ficha, name and ORCID iD here is synthetic.
"""

import io
import random
import struct
import unicodedata
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import synthetic_files as files

from thematic_redundancy.corpus.profile import person_name_forms, person_strings
from thematic_redundancy.labeling.objectives_check import _NameRedactor
from thematic_redundancy.labeling.workbook import (
    BINARY,
    CODE_PATTERN,
    EXAMPLES_SHEET,
    FIELDS,
    GUIDE_SHEET,
    LABELING_SHEET,
    LEVEL_DEFINITIONS,
    LEVELS,
    MAIN_LAYOUT,
    NO_ABSTRACT,
    NO_KEYWORDS,
    NO_OBJECTIVES,
    PILOT_LAYOUT,
    FichaView,
    LabelingWorkbookError,
    NameGuardError,
    PairView,
    ReadLabel,
    SheetContent,
    build_workbook,
    check_blind_codes,
    ficha_view,
    guard_workbooks,
    missing_levels,
    name_guard_hits,
    new_blind_codes,
    read_labels,
    save_workbook,
    shuffled,
    swapped,
)
from thematic_redundancy.shared.config import load_config

CONFIG = load_config(files.PROJECT_ROOT / "config" / "default.yaml")
NAMES = {program.key: program.name for program in CONFIG.snapshot.programs}
AUTHOR = "Inventadoapellido Ficticiosegundo, Personaejemplo Sintética"
ORCID = "https://orcid.org/0000-0002-1825-0097"
GUIDE = (("Guía sintética", ""), ("Sección", ""), ("Propósito", "Un texto de la guía."))


def thesis(number: int, **changes: Any) -> files.Thesis:
    values: dict[str, Any] = {
        "number": number,
        "program": "sistemas",
        "title": f"Título sintético {number}",
        "abstract": f"Resumen sintético {number}.",
        "objectives": f"Objetivo sintético {number}.",
        "keywords": (f"clave {number}", "prueba"),
        "authors": (AUTHOR,),
        "advisor_orcid": ORCID,
    }
    return files.Thesis(**{**values, **changes})


def view(number: int, **changes: Any) -> FichaView:
    return ficha_view(files.ficha(thesis(number, **changes), NAMES))


ADVISOR = "Ñañezinventado Ficticiáez, Personaprueba"


def guard() -> frozenset[str]:
    return person_strings([files.snapshot_record(thesis(1, advisor=ADVISOR), CONFIG)])


def pairs(count: int, **labels: Any) -> list[PairView]:
    codes = ("GH234", "JK345", "MN456", "PQ567", "RS678", "TV789")
    return [
        PairView(n, codes[n - 1], view(2 * n), view(2 * n + 1), **labels)
        for n in range(1, count + 1)
    ]


def pilot_workbook(count: int = 3) -> openpyxl.Workbook:
    return build_workbook(GUIDE, [SheetContent(LABELING_SHEET, PILOT_LAYOUT, pairs(count), True)])


def texts(workbook: openpyxl.Workbook) -> str:
    return "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


# The scale


def test_the_scale_names_d34s_three_levels_their_meaning_and_the_binary_label() -> None:
    assert LEVELS == {
        "Redundante": "redundant",
        "Relacionado": "related",
        "No relacionado": "unrelated",
    }
    assert LEVEL_DEFINITIONS == {
        "redundant": "Mismo problema y mismo enfoque.",
        "related": "Mismo tema, distinto problema o enfoque.",
        "unrelated": "Distinto tema.",
    }
    assert BINARY == {"redundant": 1, "related": 0, "unrelated": 0}


# Blind codes, order and sides


def test_each_pair_gets_a_distinct_code_of_two_letters_and_three_digits() -> None:
    identifiers = [f"L-{number:08x}" for number in range(130)]

    codes = new_blind_codes(identifiers, choice=random.Random(5).choice)

    assert list(codes) == sorted(identifiers)
    assert len(set(codes.values())) == 130
    assert all(CODE_PATTERN.fullmatch(code) for code in codes.values())
    check_blind_codes(codes, identifiers)


def test_a_code_already_given_is_drawn_again() -> None:
    picks = iter("GH234GH234JK345")

    codes = new_blind_codes(["P-1", "P-2"], choice=lambda _: next(picks))

    assert codes == {"P-1": "GH234", "P-2": "JK345"}


def test_the_codes_come_from_secrets_so_they_differ_at_each_draw_and_tell_nothing() -> None:
    identifiers = [f"L-{number:08x}" for number in range(130)]

    first, second = new_blind_codes(identifiers), new_blind_codes(identifiers)

    assert first != second
    assert not set(first.values()) & set(identifiers)
    # No code letter is a hexadecimal digit, so no SHA-256 can hold a code by chance.
    assert all(
        any(char not in "0123456789abcdef" for char in code.casefold()) for code in first.values()
    )


@pytest.mark.parametrize(
    ("codes", "message"),
    [
        ({"P-1": "GH234"}, "1 drawn pairs have no code"),
        ({"P-1": "GH234", "P-2": "GH234"}, "two pairs share a code"),
        ({"P-1": "GH234", "P-2": "AB234"}, "1 codes are not"),
        ({"P-1": "GH234", "P-2": "JK345", "P-3": "MN456"}, "1 codes belong to no drawn pair"),
    ],
)
def test_codes_that_do_not_fit_the_draw_are_refused(codes: dict[str, str], message: str) -> None:
    with pytest.raises(ValueError, match=message):
        check_blind_codes(codes, ["P-1", "P-2"])


def test_the_order_is_a_seeded_shuffle_of_the_sorted_pairs_for_each_workbook() -> None:
    identifiers = [f"L-{number:08x}" for number in range(40)]
    expected = sorted(identifiers)
    random.Random("7:order:pilot").shuffle(expected)

    assert shuffled(reversed(identifiers), seed=7, workbook="pilot") == expected
    assert shuffled(identifiers, seed=7, workbook="annotator-1") != expected
    assert shuffled(identifiers, seed=8, workbook="pilot") != expected


def test_the_sides_depend_on_the_seed_and_the_pair_alone() -> None:
    identifiers = [f"L-{number:08x}" for number in range(200)]

    sides = [swapped(identifier, seed=7) for identifier in identifiers]

    assert sides == [random.Random(f"7:sides:{i}").random() < 0.5 for i in identifiers]
    assert 60 < sum(sides) < 140
    assert sides != [swapped(identifier, seed=8) for identifier in identifiers]


# What a pair shows


def test_a_ficha_shows_its_four_fields_with_its_keywords_joined() -> None:
    shown = view(1, objectives="Objetivo general.\n\nObjetivo específico.")

    assert shown.texts == (
        "Título sintético 1",
        "Resumen sintético 1.",
        "Objetivo general.\n\nObjetivo específico.",
        "clave 1; prueba",
    )


def test_a_ficha_without_objectives_abstract_or_keywords_shows_a_fixed_text() -> None:
    shown = view(1, objectives=None, abstract=None, keywords=())

    assert (shown.abstract, shown.objectives, shown.keywords) == (
        NO_ABSTRACT,
        NO_OBJECTIVES,
        NO_KEYWORDS,
    )
    assert NO_OBJECTIVES == "(La ficha no registra objetivos.)"


def test_each_field_passes_through_the_name_redactor() -> None:
    named = "Personaejemplo Sintética Inventadoapellido Ficticiosegundo"
    redactor = _NameRedactor(person_name_forms([files.snapshot_record(thesis(1), CONFIG)]))

    shown = ficha_view(
        files.ficha(
            thesis(
                1, title=f"Método de {named}", abstract=named, objectives=named, keywords=(named,)
            ),
            NAMES,
        ),
        redactor,
    )

    assert shown.texts == (
        "Método de [nombre omitido]",
        "[nombre omitido]",
        "[nombre omitido]",
        "[nombre omitido]",
    )
    assert redactor.redacted == 4


# The sheets


def test_each_pair_is_a_block_of_its_four_fields_with_both_theses_side_by_side() -> None:
    workbook = pilot_workbook()
    sheet = workbook[LABELING_SHEET]

    assert workbook.sheetnames == [GUIDE_SHEET, LABELING_SHEET]
    assert workbook.active.title == LABELING_SHEET
    assert [cell.value for cell in sheet[1]] == list(PILOT_LAYOUT.columns)
    assert PILOT_LAYOUT.columns == (
        "N.º",
        "Código",
        "Campo",
        "Tesis A",
        "Tesis B",
        "Etiqueta acordada",
        "Justificación",
    )
    for index, pair in enumerate(pairs(3)):
        first = 2 + 4 * index
        assert (sheet.cell(first, 1).value, sheet.cell(first, 2).value) == (pair.n, pair.code)
        for offset, field in enumerate(FIELDS):
            row = first + offset
            assert sheet.cell(row, 3).value == field
            assert sheet.cell(row, 4).value == pair.left.texts[offset]
            assert sheet.cell(row, 5).value == pair.right.texts[offset]
        for column in "ABFG":
            assert f"{column}{first}:{column}{first + 3}" in {
                str(merged) for merged in sheet.merged_cells.ranges
            }
    assert sheet.max_row == 13
    assert sheet.freeze_panes == "D2"


def test_the_label_column_offers_the_three_levels_and_refuses_other_values() -> None:
    sheet = pilot_workbook()[LABELING_SHEET]

    lists = sheet.data_validations.dataValidation
    assert [(str(rule.sqref), rule.formula1) for rule in lists] == [
        ("F2:F13", '"Redundante,Relacionado,No relacionado"')
    ]
    assert lists[0].showErrorMessage is True


def test_the_examples_sheet_shows_the_agreed_labels_without_a_drop_down() -> None:
    examples = pairs(2, label="Relacionado", note="Comparten el tema.")
    workbook = build_workbook(
        GUIDE,
        [
            SheetContent(EXAMPLES_SHEET, PILOT_LAYOUT, examples, False),
            SheetContent(LABELING_SHEET, MAIN_LAYOUT, pairs(3), True),
        ],
    )

    sheet = workbook[EXAMPLES_SHEET]
    assert workbook.sheetnames == [GUIDE_SHEET, EXAMPLES_SHEET, LABELING_SHEET]
    assert (sheet.cell(2, 6).value, sheet.cell(2, 7).value) == ("Relacionado", "Comparten el tema.")
    assert (sheet.cell(6, 6).value, sheet.cell(6, 7).value) == ("Relacionado", "Comparten el tema.")
    assert not sheet.data_validations.dataValidation
    labeling = workbook[LABELING_SHEET]
    assert [cell.value for cell in labeling[1]][5:] == ["Etiqueta", "Notas (opcional)"]


def test_the_guide_sheet_holds_its_lines_and_no_cell_is_a_formula() -> None:
    formula = '=HYPERLINK("https://example.org")'
    workbook = build_workbook(
        GUIDE,
        [
            SheetContent(
                LABELING_SHEET,
                PILOT_LAYOUT,
                [PairView(1, "GH234", view(2, title=formula), view(3))],
                True,
            )
        ],
    )

    guide = workbook[GUIDE_SHEET]
    assert [(guide.cell(row, 1).value, guide.cell(row, 2).value) for row in (1, 2, 3)] == [
        ("Guía sintética", None),
        ("Sección", None),
        ("Propósito", "Un texto de la guía."),
    ]
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
    assert all(cell.data_type != "f" for cell in cells)
    assert any(cell.value == formula for cell in cells)
    assert not any(cell.hyperlink for cell in cells)


def test_a_saved_workbook_reads_back_whole(tmp_path: Path) -> None:
    path = tmp_path / "libro.xlsx"

    save_workbook(pilot_workbook(), path)

    assert openpyxl.load_workbook(path)[LABELING_SHEET].cell(2, 2).value == "GH234"
    assert [p.name for p in tmp_path.iterdir()] == ["libro.xlsx"]


# The name guard


@pytest.mark.parametrize(
    "written",
    [
        "Según lapersonaejemplo sintética inventadoapellido ficticiosegundo del ensayo",
        "SEGÚN INVENTADOAPELLIDO FICTICIOSEGUNDOS",
        "Según Inventadoapellido\n  Ficticiosegundo, en su planta",
        # Its ñ and á decomposed: only Unicode NFC makes it the advisor's family name.
        unicodedata.normalize("NFD", "Según laÑañezinventado Ficticiáez del ensayo"),
        "Ver 0000-0002-1825-0097 en el registro",
    ],
    ids=["glued", "upper-case-and-glued", "line-break", "decomposed-accents", "orcid"],
)
def test_the_guard_finds_a_name_or_orcid_without_word_boundaries_and_names_only_the_place(
    written: str,
) -> None:
    workbook = build_workbook(
        GUIDE,
        [
            SheetContent(
                LABELING_SHEET,
                PILOT_LAYOUT,
                [
                    PairView(1, "GH234", view(2), view(3)),
                    PairView(2, "JK345", view(4), view(5, objectives=written)),
                ],
                True,
            )
        ],
    )

    hits = name_guard_hits(workbook, guard(), name="libro.xlsx")

    assert hits == ["libro.xlsx, sheet 'Etiquetado', cell E8 (pair JK345)"]
    with pytest.raises(NameGuardError) as refused:
        guard_workbooks({"libro.xlsx": workbook}, guard())
    message = str(refused.value).casefold()
    assert "cell e8 (pair jk345)" in message and "no workbook was written" in message
    for leaked in ("inventadoapellido", "ñañez", "0000-0002"):
        assert leaked not in unicodedata.normalize("NFC", message)


def test_the_guard_also_reads_the_guide_the_sheet_names_and_the_drop_down_lists() -> None:
    workbook = pilot_workbook(1)
    workbook[GUIDE_SHEET].cell(3, 2).value = "Escrito por Inventadoapellido Ficticiosegundo."
    workbook[LABELING_SHEET].data_validations.dataValidation[
        0
    ].error = "Consulte a Inventadoapellido Ficticiosegundo"
    workbook.create_sheet("0000-0002-1825-0097")

    hits = name_guard_hits(workbook, guard(), name="libro.xlsx")

    assert hits == [
        "libro.xlsx, sheet 'Guía', cell B3",
        "libro.xlsx, sheet 'Etiquetado', the drop-down list of F2:F5",
        "libro.xlsx, sheet 3, its name",
    ]


def test_the_guard_lists_every_hit_of_every_workbook_and_passes_a_clean_one() -> None:
    clean = pilot_workbook(2)
    named = pilot_workbook(2)
    named[LABELING_SHEET].cell(2, 4).value = "Inventadoapellido Ficticiosegundo"
    named[LABELING_SHEET].cell(9, 5).value = ORCID

    with pytest.raises(NameGuardError) as refused:
        guard_workbooks({"a.xlsx": clean, "b.xlsx": named}, guard())

    assert refused.value.locations == (
        "b.xlsx, sheet 'Etiquetado', cell D2 (pair GH234)",
        "b.xlsx, sheet 'Etiquetado', cell E9 (pair JK345)",
    )
    cells = {
        cell.value
        for sheet in clean
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    }
    rule = clean[LABELING_SHEET].data_validations.dataValidation[0]
    looked_at = cells | set(clean.sheetnames) | {rule.formula1, rule.error, rule.errorTitle}
    assert guard_workbooks({"a.xlsx": clean}, guard()) == len(looked_at)


def test_the_guard_does_not_look_for_titles_or_abstracts() -> None:
    # The workbook shows them, so they are not in the guard's strings.
    record = files.snapshot_record(thesis(2), CONFIG)
    workbook = pilot_workbook(1)

    assert name_guard_hits(workbook, person_strings([record]), name="libro.xlsx") == []


# Reading a returned workbook


def saved(tmp_path: Path, workbook: openpyxl.Workbook | None = None) -> Path:
    path = tmp_path / "libro.xlsx"
    save_workbook(workbook or pilot_workbook(), path)
    return path


EXPECTED = {"GH234": 1, "JK345": 2, "MN456": 3}


def fill(path: Path, cells: dict[str, Any]) -> None:
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[LABELING_SHEET]
    for coordinate, value in cells.items():
        sheet[coordinate].value = value
    workbook.save(path)


def labeled(path: Path) -> None:
    fill(
        path,
        {
            "F2": "Redundante",
            "G2": "Mismo problema.",
            "F6": "  no RELACIONADO ",
            "F10": "Relacionado",
            "G10": "Mismo tema.",
        },
    )


def test_the_reader_returns_each_pairs_level_and_note_in_sheet_order(tmp_path: Path) -> None:
    path = saved(tmp_path)
    labeled(path)

    labels = read_labels(path, PILOT_LAYOUT, EXPECTED)

    assert labels == (
        ReadLabel(1, "GH234", "redundant", "Mismo problema."),
        ReadLabel(2, "JK345", "unrelated", None),
        ReadLabel(3, "MN456", "related", "Mismo tema."),
    )
    assert missing_levels(labels) == ()
    assert missing_levels(labels[:2]) == ("Relacionado",)


def test_the_reader_lists_every_problem_at_once(tmp_path: Path) -> None:
    path = saved(tmp_path)
    labeled(path)
    workbook = openpyxl.load_workbook(path)
    workbook[LABELING_SHEET].unmerge_cells("F10:F13")
    workbook.save(path)
    fill(
        path,
        {
            "F2": None,
            "F6": "Parecido",
            "A10": 9,
            "C11": "Resumen del par",
            "F12": "Redundante",
        },
    )

    with pytest.raises(LabelingWorkbookError) as refused:
        read_labels(path, PILOT_LAYOUT, EXPECTED)

    assert refused.value.problems == (
        "row 2 (pair GH234) has no 'Etiqueta acordada'",
        "row 6 (pair JK345): 'Parecido' is not an allowed 'Etiqueta acordada'; use one of: "
        "Redundante, Relacionado, No relacionado",
        "row 10 (pair MN456): 'N.º' was changed; it must be 3",
        "row 10 (pair MN456): the 'Campo' rows of the pair were changed; they must be Título, "
        "Resumen, Objetivos, Palabras clave",
        "row 12 (pair MN456): 'Etiqueta acordada' must be in the first row of the pair",
    )


def test_the_reader_refuses_a_lost_a_repeated_and_an_unknown_pair(tmp_path: Path) -> None:
    path = saved(tmp_path)
    labeled(path)
    fill(path, {"B6": "GH234", "B10": "ZZ999"})

    with pytest.raises(LabelingWorkbookError) as refused:
        read_labels(path, PILOT_LAYOUT, EXPECTED)

    assert refused.value.problems == (
        "row 6 (pair GH234): 'N.º' was changed; it must be 1",
        "row 10: the code 'ZZ999' is not one of this workbook's pairs",
        "the pair GH234 appears 2 times",
        "the pair JK345 (N.º 2) is missing",
        "the pair MN456 (N.º 3) is missing",
    )


def test_the_reader_refuses_a_changed_header_or_a_missing_sheet(tmp_path: Path) -> None:
    path = saved(tmp_path)
    fill(path, {"F1": "Etiqueta"})

    with pytest.raises(LabelingWorkbookError, match="header row of 'Etiquetado' changed"):
        read_labels(path, PILOT_LAYOUT, EXPECTED)
    with pytest.raises(LabelingWorkbookError, match="has no sheet 'Ejemplos'"):
        read_labels(path, PILOT_LAYOUT, EXPECTED, sheet=EXAMPLES_SHEET)


def test_the_reader_refuses_a_row_outside_every_pair(tmp_path: Path) -> None:
    path = saved(tmp_path)
    labeled(path)
    fill(path, {"D15": "Una nota suelta"})

    with pytest.raises(LabelingWorkbookError) as refused:
        read_labels(path, PILOT_LAYOUT, EXPECTED)

    assert refused.value.problems == ("row 15 belongs to no pair",)


def member_data(path: Path, name: str) -> tuple[int, int]:
    """Return where the stored bytes of the archive member ``name`` start and end."""
    raw = path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo(name)
    start = info.header_offset + 30 + sum(struct.unpack("<HH", raw[info.header_offset + 26 :][:4]))
    return start, start + info.compress_size


def corrupt_a_deflate_member(path: Path) -> None:
    # A reserved block type in the first deflate block: zlib refuses it (zlib.error).
    start, _ = member_data(path, "xl/worksheets/sheet2.xml")
    raw = bytearray(path.read_bytes())
    raw[start] = 0xFF
    path.write_bytes(bytes(raw))


def rewrite_parts(path: Path, change: Callable[[str, bytes], bytes]) -> None:
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            target.writestr(info, change(info.filename, source.read(info)))


@pytest.mark.parametrize(
    "damage",
    [
        lambda path: path.write_bytes(b"synthetic bytes that are not a workbook"),
        lambda path: path.write_bytes(path.read_bytes()[:200]),
        corrupt_a_deflate_member,
        lambda path: rewrite_parts(
            path,
            lambda name, data: (
                data[: len(data) // 2] if name.startswith("xl/worksheets/") else data
            ),
        ),
        lambda path: rewrite_parts(
            path,
            lambda name, data: (
                data.replace(b'sheetId="1"', b'sheetId="x"') if name == "xl/workbook.xml" else data
            ),
        ),
    ],
    ids=["garbage", "truncated-zip", "corrupted-deflate", "broken-xml", "invalid-value"],
)
def test_a_file_that_is_not_a_readable_workbook_is_refused_in_one_line(
    tmp_path: Path, damage: Callable[[Path], None]
) -> None:
    path = saved(tmp_path)
    damage(path)

    with pytest.raises(ValueError, match="libro.xlsx is not a readable xlsx workbook") as refused:
        read_labels(path, PILOT_LAYOUT, EXPECTED)

    assert "\n" not in str(refused.value)


def test_a_corrupted_deflate_member_raises_zlib_error_which_the_reader_refuses(
    tmp_path: Path,
) -> None:
    path = saved(tmp_path)
    corrupt_a_deflate_member(path)

    with pytest.raises(ValueError, match=r"\(error: Error -3 while decompressing data"):
        read_labels(path, PILOT_LAYOUT, EXPECTED)


def test_a_workbook_cut_short_while_it_is_read_raises_eof_error_which_the_reader_refuses(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # As when another program is still saving the file: the archive's directory was read, but
    # the bytes of a part end early. zipfile then raises EOFError.
    path = saved(tmp_path)
    start, end = member_data(path, "xl/worksheets/sheet2.xml")
    shared = zipfile._SharedFile  # type: ignore[attr-defined]
    read = shared.read

    def cut_short(self: Any, n: int = -1) -> bytes:
        return b"" if start <= self._pos < end else read(self, n)

    monkeypatch.setattr(shared, "read", cut_short)

    with pytest.raises(ValueError, match=r"is not a readable xlsx workbook \(EOFError"):
        read_labels(path, PILOT_LAYOUT, EXPECTED)


def test_a_workbook_under_another_extension_is_refused(tmp_path: Path) -> None:
    path = saved(tmp_path)
    renamed = path.with_suffix(".xls")
    renamed.write_bytes(path.read_bytes())

    with pytest.raises(ValueError, match="is not a readable xlsx workbook"):
        read_labels(renamed, PILOT_LAYOUT, EXPECTED)
