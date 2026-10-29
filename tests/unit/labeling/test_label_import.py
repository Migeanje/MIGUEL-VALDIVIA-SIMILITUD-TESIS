"""The importer of the labeled set's workbooks (T20c) end to end, through the command line: the
workbooks are made by T20b's steps on the synthetic files of ``synthetic_files.py``, labeled
here with synthetic choices, and imported back; then the adjudication of the disputed pairs.

Only git, the clocks and the random source of the blind codes are replaced; the loaders and the
configuration are the real ones.
"""

import hashlib
import itertools
import json
import random
import shutil
import zipfile
from collections import Counter
from collections.abc import Callable, Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pyarrow.parquet as pq
import pytest
import synthetic_files as files

from thematic_redundancy.evaluation.agreement import cohen_kappa, fleiss_kappa
from thematic_redundancy.labeling.consensus import LEVELS, ConsensusLabel, final_level
from thematic_redundancy.labeling.label_import import (
    ADJUDICATION_LABELS_MANIFEST_FILE_NAME,
    ADJUDICATION_LAYOUT_FILE_NAME,
    ADJUDICATION_MANIFEST_FILE_NAME,
    KAPPA_TARGET,
    MAIN_LABELS_MANIFEST_FILE_NAME,
    AnnotatorLabelRow,
    FleissInputRow,
    assign_adjudicators,
)
from thematic_redundancy.labeling.pair_workbooks import (
    AssignmentRow,
    BlindCodeRow,
    LayoutRow,
    WorkbookManifest,
    load_workbook_inputs,
    main,
)
from thematic_redundancy.labeling.workbook import LABELING_SHEET
from thematic_redundancy.shared.config import load_config
from thematic_redundancy.shared.storage import ParquetTableStore
from thematic_redundancy.topic_map.run_experiment import GitState

SNAPSHOT = files.SNAPSHOT_ID
RECORDS = Path("data") / "labels" / "labeled_set" / SNAPSHOT
WORKBOOKS = RECORDS / "workbooks"
LABELS = RECORDS / "labels"
RESULTS = Path("results") / "labeled_set" / SNAPSHOT / "workbooks"
SEED = 20261021


def clean(root: Path, path: str) -> GitState:
    return GitState(commit="c" * 40, dirty=())


def run(root: Path, step: str, *extra: str, **overrides: Any) -> int:
    ticks = itertools.count(10.0, 0.25)
    moments = itertools.count()
    arguments: dict[str, Any] = {
        "git": clean,
        "clock": lambda: datetime(2026, 10, 27, 9, 0, next(moments) % 60, tzinfo=UTC),
        "timer": lambda: next(ticks),
        "choice": random.Random(5).choice,
    }
    arguments.update(overrides)
    return main([step, "--project-root", str(root), "--snapshot-id", SNAPSHOT, *extra], **arguments)


def read(model: Any, root: Path, relative: Path) -> tuple[Any, ...]:
    return ParquetTableStore().read(model, root / relative)


def manifest_of(root: Path, name: str) -> WorkbookManifest:
    return WorkbookManifest.model_validate_json((root / RESULTS / name).read_text("utf-8"))


def files_under(root: Path) -> dict[str, str]:
    return {
        p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(root.rglob("*"))
        if p.is_file()
    }


def truth(pair_id: str) -> str:
    """The level that every annotator gives a pair unless a test says otherwise."""
    return LEVELS[int(hashlib.sha256(pair_id.encode()).hexdigest(), 16) % 3]


def fill(
    root: Path,
    layout_name: str,
    workbook_name: Callable[[int], str],
    label_of: Callable[[int, LayoutRow], str | None],
    note_of: Callable[[int, LayoutRow], str | None] = lambda annotator, row: None,
) -> None:
    """Label every workbook of a layout: annotator ``k``'s pair ``row`` gets ``label_of(k,
    row)``."""
    layout = read(LayoutRow, root, RECORDS / layout_name)
    for annotator in sorted({row.annotator for row in layout}):
        path = root / WORKBOOKS / workbook_name(annotator)
        workbook = openpyxl.load_workbook(path)
        sheet = workbook[LABELING_SHEET]
        for row in (r for r in layout if r.annotator == annotator):
            sheet.cell(2 + 4 * (row.n - 1), 6).value = label_of(annotator, row)
            sheet.cell(2 + 4 * (row.n - 1), 7).value = note_of(annotator, row)
        workbook.save(path)


def annotator_file(k: int) -> str:
    return f"etiquetado_anotador_{k}.xlsx"


def adjudicator_file(k: int) -> str:
    return f"desempate_anotador_{k}.xlsx"


def label_main(
    root: Path,
    second: Mapping[str, str] | None = None,
    notes: Mapping[tuple[int, str], str] | None = None,
    **blanks: Any,
) -> None:
    """Every annotator gives each pair its :func:`truth`, but a pair's second annotator gives
    the level in ``second``, if any; each note is ``Nota <k>.`` unless ``notes`` gives one by
    annotator and pair id."""
    assignment = {
        row.pair_id: row for row in read(AssignmentRow, root, RECORDS / "assignment.parquet")
    }

    def label_of(annotator: int, row: LayoutRow) -> str | None:
        if (annotator, row.pair_id) in blanks.get("blank", set()):
            return None
        if annotator == assignment[row.pair_id].annotator_b and row.pair_id in (second or {}):
            return (second or {})[row.pair_id]
        return truth(row.pair_id)

    fill(
        root,
        "main_layout.parquet",
        annotator_file,
        label_of,
        lambda k, row: (notes or {}).get((k, row.pair_id), f"Nota {k}."),
    )


def disputes(root: Path, binary: int = 3, levels_only: int = 3) -> dict[str, str]:
    """Choose ``binary`` pairs whose second label crosses the binary label (Redundante against
    Relacionado) and ``levels_only`` pairs that differ on the levels only."""
    labeled = sorted(
        row.pair_id for row in read(AssignmentRow, root, RECORDS / "assignment.parquet")
    )
    on_binary = [p for p in labeled if truth(p) == "Redundante"][:binary]
    on_levels = [p for p in labeled if truth(p) == "Relacionado"][:levels_only]
    return {**dict.fromkeys(on_binary, "Relacionado"), **dict.fromkeys(on_levels, "No relacionado")}


def build(base: Path, annotators: int) -> Path:
    root = base / f"made-{annotators}"
    (root / "config").mkdir(parents=True)
    shutil.copyfile(
        files.PROJECT_ROOT / "config" / "default.yaml", root / "config" / "default.yaml"
    )
    files.write_labeling_project(root, load_config(root / "config" / "default.yaml"))
    assert run(root, "pilot") == 0
    pilot = openpyxl.load_workbook(root / WORKBOOKS / "etiquetado_piloto.xlsx")
    for n in range(1, 11):
        pilot[LABELING_SHEET].cell(2 + 4 * (n - 1), 6).value = LEVELS[n % 3]
    pilot.save(root / WORKBOOKS / "etiquetado_piloto.xlsx")
    assert run(root, "import-pilot") == 0
    config = root / "config" / "default.yaml"
    text = config.read_bytes().decode("utf-8")
    config.write_bytes(text.replace("annotators: null", f"annotators: {annotators}").encode())
    assert run(root, "main") == 0
    return root


@pytest.fixture(scope="module")
def made(tmp_path_factory: pytest.TempPathFactory) -> dict[int, Path]:
    """Projects whose annotators' workbooks are made, for 3 and 4 annotators."""
    base = tmp_path_factory.mktemp("made")
    return {annotators: build(base, annotators) for annotators in (3, 4)}


@pytest.fixture
def project(made: dict[int, Path], tmp_path: Path) -> Path:
    """A copy of the project of 3 annotators, its workbooks made and not yet labeled."""
    return Path(shutil.copytree(made[3], tmp_path / "project"))


@pytest.fixture
def project4(made: dict[int, Path], tmp_path: Path) -> Path:
    return Path(shutil.copytree(made[4], tmp_path / "project4"))


def secrets_of(root: Path) -> list[str]:
    """Every doc code, pair id and blind code, none of which a tracked manifest may hold."""
    codes = [row.code for row in read(BlindCodeRow, root, RECORDS / "blind_codes.parquet")]
    rows = read(files.PairRow, root, RECORDS / "pairs.parquet")
    return [
        *codes,
        *(r.pair_id for r in rows),
        *(c for r in rows for c in (r.doc_code_a, r.doc_code_b)),
    ]


# import-main


def test_import_main_stores_each_annotators_labels_and_with_no_dispute_the_consensus(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    label_main(project)
    capsys.readouterr()

    code = run(project, "import-main")

    out = capsys.readouterr().out
    assert code == 0
    labels = read(AnnotatorLabelRow, project, LABELS / "labels.parquet")
    layout = read(LayoutRow, project, RECORDS / "main_layout.parquet")
    assert len(labels) == 240 == len(layout)
    assert {(r.annotator, r.pair_id, r.code) for r in labels} == {
        (r.annotator, r.pair_id, r.code) for r in layout
    }
    for row in labels:
        assert row.label_3 == truth(row.pair_id)
        assert row.label == int(row.label_3 == "Redundante")
        assert row.note == f"Nota {row.annotator}."
    consensus = read(ConsensusLabel, project, LABELS / "consensus.parquet")
    labeled = [
        p.pair_id
        for p in read(files.PairRow, project, RECORDS / "pairs.parquet")
        if p.phase == "labeled"
    ]
    assert [row.pair_id for row in consensus] == labeled
    assert all(row.source == "agreement" and row.label_3 == truth(row.pair_id) for row in consensus)
    assert [f.name for f in pq.read_schema(project / LABELS / "consensus.parquet")] == [
        "pair_id",
        "label_3",
        "label",
        "source",
    ]
    manifest = manifest_of(project, MAIN_LABELS_MANIFEST_FILE_NAME)
    assert manifest.counts == {
        "annotators": 3,
        "pairs": 120,
        "labels": 240,
        "agreed": 120,
        "disputed": 0,
        "consensus_rows": 120,
        "redacted_notes": 0,
    }
    assert manifest.details["fleiss_binary"]["kappa"] == 1.0
    assert manifest.details["meets_target"] is True
    assert set(manifest.data_files) == {
        "labels/labels.parquet",
        "labels/fleiss_inputs.parquet",
        "labels/consensus.parquet",
    }
    text = (project / RESULTS / MAIN_LABELS_MANIFEST_FILE_NAME).read_text("utf-8")
    for secret in secrets_of(project):
        assert secret not in text and secret not in out
    assert "Fleiss' kappa on the binary labels: 1.000 over 120 pairs (target 0.6: met)" in out


def test_the_kappas_are_those_of_the_agreement_module_on_the_first_two_labels(
    project: Path,
) -> None:
    second = disputes(project, binary=8, levels_only=10)
    label_main(project, second)

    assert run(project, "import-main") == 0

    assignment = read(AssignmentRow, project, RECORDS / "assignment.parquet")
    first = [truth(p.pair_id) for p in assignment]
    other = [second.get(p.pair_id, truth(p.pair_id)) for p in assignment]
    binary = [int(level == "Redundante") for level in first]
    binary2 = [int(level == "Redundante") for level in other]
    details = manifest_of(project, MAIN_LABELS_MANIFEST_FILE_NAME).details
    expected = fleiss_kappa(list(zip(binary, binary2, strict=True)), categories=(1, 0))
    assert details["fleiss_binary"]["kappa"] == expected.kappa
    assert details["fleiss_binary"]["observed_agreement"] == expected.observed_agreement
    levels = fleiss_kappa(list(zip(first, other, strict=True)), categories=LEVELS)
    assert details["fleiss_levels"]["kappa"] == levels.kappa
    assert details["meets_target"] == (expected.kappa >= KAPPA_TARGET)
    for record in details["cohen"]:
        group = tuple(record["annotators"])
        mine = [i for i, p in enumerate(assignment) if (p.annotator_a, p.annotator_b) == group]
        assert record["pairs"] == len(mine) == 40
        on_levels = cohen_kappa(
            [first[i] for i in mine], [other[i] for i in mine], categories=LEVELS, weights="none"
        )
        assert record["levels"]["kappa"] == on_levels.kappa
        assert record["levels"]["confusion"] == [list(r) for r in on_levels.confusion]
    assert [record["annotators"] for record in details["cohen"]] == [[1, 2], [1, 3], [2, 3]]


def test_the_tracked_combinations_give_fleiss_kappa_again_without_a_pair_id(
    project: Path,
) -> None:
    label_main(project, disputes(project, binary=5, levels_only=7))
    assert run(project, "import-main") == 0
    details = manifest_of(project, MAIN_LABELS_MANIFEST_FILE_NAME).details

    for kind, categories in (("binary", (1, 0)), ("levels", LEVELS)):
        ratings = []
        for key, count in details["combinations"][kind].items():
            pair = [int(v) if kind == "binary" else v for v in key.split("+")]
            ratings += [pair] * count
        again = fleiss_kappa(ratings, categories=categories)
        assert again.kappa == pytest.approx(details[f"fleiss_{kind}"]["kappa"], abs=1e-15)
    assert sum(details["combinations"]["binary"].values()) == 120


def test_the_fleiss_inputs_count_each_pairs_first_two_labels(project: Path) -> None:
    second = disputes(project)
    label_main(project, second)

    assert run(project, "import-main") == 0

    rows = read(FleissInputRow, project, LABELS / "fleiss_inputs.parquet")
    assert len(rows) == 120
    for row in rows:
        levels = Counter((truth(row.pair_id), second.get(row.pair_id, truth(row.pair_id))))
        assert (row.redundant, row.related, row.unrelated) == tuple(
            levels[level] for level in LEVELS
        )
        assert row.binary_1 == levels["Redundante"] and row.binary_0 == 2 - levels["Redundante"]


def test_a_pair_is_disputed_when_its_two_levels_differ_and_then_no_consensus_is_written(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    label_main(project, disputes(project, binary=3, levels_only=4))
    capsys.readouterr()

    assert run(project, "import-main") == 0

    out = capsys.readouterr().out
    manifest = manifest_of(project, MAIN_LABELS_MANIFEST_FILE_NAME)
    assert (manifest.counts["agreed"], manifest.counts["disputed"]) == (113, 7)
    assert manifest.counts["consensus_rows"] == 0
    assert not (project / LABELS / "consensus.parquet").exists()
    assert "disputed: 7 (3 of them on the binary label too)" in out
    assert "adjudication, for the 7 disputed pairs" in out


def test_an_incomplete_workbook_is_refused_and_nothing_is_written(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    layout = read(LayoutRow, project, RECORDS / "main_layout.parquet")
    missing = next(r for r in layout if r.annotator == 2 and r.n == 7)
    label_main(project, blank={(2, missing.pair_id)})
    before = {k: v for k, v in files_under(project).items() if not k.endswith(".lock")}
    capsys.readouterr()

    code = run(project, "import-main")

    err = capsys.readouterr().err
    assert code == 1
    assert f"etiquetado_anotador_2.xlsx: row 26 (pair {missing.code}) has no 'Etiqueta'" in err
    assert {k: v for k, v in files_under(project).items() if not k.endswith(".lock")} == before
    assert not (project / LABELS).exists()


def test_every_problem_of_every_workbook_is_listed_at_once(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    label_main(project)
    first = openpyxl.load_workbook(project / WORKBOOKS / annotator_file(1))
    first[LABELING_SHEET].cell(1, 6).value = "Mi etiqueta"
    first.save(project / WORKBOOKS / annotator_file(1))
    second = openpyxl.load_workbook(project / WORKBOOKS / annotator_file(2))
    sheet = second[LABELING_SHEET]
    sheet.cell(2, 6).value = "Parecido"
    sheet.cell(7, 3).value = "Resumen editado"
    code = sheet.cell(10, 2).value
    sheet.cell(10, 2).value = sheet.cell(14, 2).value
    second.save(project / WORKBOOKS / annotator_file(2))
    (project / WORKBOOKS / annotator_file(3)).unlink()
    capsys.readouterr()

    assert run(project, "import-main") == 1

    err = capsys.readouterr().err
    assert "etiquetado_anotador_1.xlsx: the header row of 'Etiquetado' changed" in err
    assert (
        "etiquetado_anotador_2.xlsx: row 2 (pair" in err and "'Parecido' is not an allowed" in err
    )
    assert "the 'Campo' rows of the pair were changed" in err
    assert f"etiquetado_anotador_2.xlsx: the pair {code} (N.º 3) is missing" in err
    assert "appears 2 times" in err
    assert "etiquetado_anotador_3.xlsx is missing; save the returned workbook as" in err
    assert not (project / LABELS).exists()


def labeling_part(path: Path) -> tuple[int, int]:
    """Return where the stored bytes of the labeling sheet's part start and end."""
    raw = path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo("xl/worksheets/sheet3.xml")
    lengths = raw[info.header_offset + 26 : info.header_offset + 30]
    start = info.header_offset + 30 + int.from_bytes(lengths[:2], "little")
    start += int.from_bytes(lengths[2:], "little")
    return start, start + info.compress_size


def corrupt(path: Path) -> None:
    """Give the labeling sheet's deflate stream a reserved block type: zlib.error."""
    start, _ = labeling_part(path)
    raw = bytearray(path.read_bytes())
    raw[start] = 0xFF
    path.write_bytes(bytes(raw))


def test_a_damaged_workbook_is_refused_with_the_others_problems(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    label_main(
        project, blank={(1, read(LayoutRow, project, RECORDS / "main_layout.parquet")[0].pair_id)}
    )
    corrupt(project / WORKBOOKS / annotator_file(3))
    capsys.readouterr()

    assert run(project, "import-main") == 1

    err = capsys.readouterr().err
    assert "etiquetado_anotador_3.xlsx is not a readable xlsx workbook (error: Error -3" in err
    assert "etiquetado_anotador_1.xlsx: row 2 (pair" in err
    assert not (project / LABELS).exists()


def test_a_workbook_cut_short_while_it_is_read_is_refused(
    project: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    label_main(project)
    path = (project / WORKBOOKS / annotator_file(2)).resolve()
    start, end = labeling_part(path)
    shared = zipfile._SharedFile  # type: ignore[attr-defined]
    real = shared.read

    def cut_short(self: Any, n: int = -1) -> bytes:
        # As when another program is still saving this workbook: its part ends early.
        inside = start <= self._pos < end and Path(self._file.name).resolve() == path
        return b"" if inside else real(self, n)

    monkeypatch.setattr(shared, "read", cut_short)

    assert run(project, "import-main") == 1
    assert "etiquetado_anotador_2.xlsx is not a readable xlsx workbook (EOFError" in (
        capsys.readouterr().err
    )


def test_a_note_with_a_name_the_redactor_leaves_is_refused_without_echo(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    layout = read(LayoutRow, project, RECORDS / "main_layout.parquet")
    one = next(r for r in layout if r.annotator == 1 and r.n == 1)
    two = next(r for r in layout if r.annotator == 2 and r.n == 2)
    label_main(
        project,
        notes={
            (1, one.pair_id): "Ver 0000-0002-1825-0097.",
            (2, two.pair_id): "Como en lainventadoapellido ficticiosegundo.",
        },
    )
    capsys.readouterr()

    assert run(project, "import-main") == 1

    err = capsys.readouterr().err
    for name, row in ((annotator_file(1), one), (annotator_file(2), two)):
        assert (
            f"{name}: the 'Notas (opcional)' of pair {row.code} holds a person's name or an "
            "ORCID iD of the snapshot" in err
        )
    assert "inventadoapellido" not in err.casefold() and "0000-0002" not in err
    assert not (project / LABELS).exists()


def test_a_note_with_a_whole_word_name_is_stored_redacted(project: Path) -> None:
    layout = read(LayoutRow, project, RECORDS / "main_layout.parquet")
    one = next(r for r in layout if r.annotator == 3 and r.n == 5)
    label_main(project, notes={(3, one.pair_id): "Como en Supuestoapellido Imaginario."})

    assert run(project, "import-main") == 0

    labels = read(AnnotatorLabelRow, project, LABELS / "labels.parquet")
    note = next(r.note for r in labels if (r.annotator, r.pair_id) == (3, one.pair_id))
    assert note == "Como en [nombre omitido]."
    assert manifest_of(project, MAIN_LABELS_MANIFEST_FILE_NAME).counts["redacted_notes"] == 1


def test_import_main_waits_for_the_workbooks_happens_once_and_has_a_dry_run(
    project: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    early = tmp_path / "early"
    (early / "config").mkdir(parents=True)
    shutil.copyfile(
        files.PROJECT_ROOT / "config" / "default.yaml", early / "config" / "default.yaml"
    )
    files.write_labeling_project(early, load_config(early / "config" / "default.yaml"))
    loaded: list[Any] = []

    def load(*arguments: Any) -> Any:
        loaded.append(arguments)
        return load_workbook_inputs(*arguments)

    for step, missing in (
        ("import-main", "main_manifest.json"),
        ("adjudication", "main_labels_manifest.json"),
        ("import-adjudication", "adjudication_manifest.json"),
    ):
        assert run(early, step, "--dry-run", load=load) == 1
        assert f"are missing ({missing}); run" in capsys.readouterr().err
    # Each step waits for the one before it before it reads anything.
    assert loaded == []
    label_main(project)
    before = files_under(project)

    assert run(project, "import-main", "--dry-run") == 0
    assert "Fleiss' kappa on the binary labels" in capsys.readouterr().out
    assert files_under(project) == before
    assert run(project, "import-main") == 0
    after = files_under(project)
    assert run(project, "import-main") == 1
    assert "already imported" in capsys.readouterr().err
    assert files_under(project) == after


# Adjudication


def test_each_disputed_pair_goes_to_the_annotator_who_did_not_label_it(project: Path) -> None:
    second = disputes(project, binary=4, levels_only=5)
    label_main(project, second)
    assert run(project, "import-main") == 0

    assert run(project, "adjudication") == 0

    layout = read(LayoutRow, project, RECORDS / ADJUDICATION_LAYOUT_FILE_NAME)
    assignment = {
        r.pair_id: r for r in read(AssignmentRow, project, RECORDS / "assignment.parquet")
    }
    assert {row.pair_id for row in layout} == set(second)
    for row in layout:
        assert row.annotator not in (
            assignment[row.pair_id].annotator_a,
            assignment[row.pair_id].annotator_b,
        )
        assert row.workbook == f"adjudication-{row.annotator}"
    main_layout = {r.pair_id: r for r in read(LayoutRow, project, RECORDS / "main_layout.parquet")}
    for row in layout:
        shown = main_layout[row.pair_id]
        assert (row.code, row.doc_code_left, row.doc_code_right) == (
            shown.code,
            shown.doc_code_left,
            shown.doc_code_right,
        )
    adjudicators = sorted({row.annotator for row in layout})
    assert sorted(p.name for p in (project / WORKBOOKS).glob("desempate_*.xlsx")) == [
        adjudicator_file(k) for k in adjudicators
    ]
    for k in adjudicators:
        workbook = openpyxl.load_workbook(project / WORKBOOKS / adjudicator_file(k))
        sheet = workbook[LABELING_SHEET]
        mine = [row for row in layout if row.annotator == k]
        assert [sheet.cell(2 + 4 * i, 2).value for i in range(len(mine))] == [r.code for r in mine]
        assert all(
            sheet.cell(2 + 4 * i, column).value is None
            for i in range(len(mine))
            for column in (6, 7)
        )
        guide = workbook["Guía"]
        assert guide.cell(1, 1).value == f"Etiquetado de pares de tesis: desempate, anotador {k}"
        text = " ".join(str(c.value) for row in guide.iter_rows() for c in row if c.value)
        assert "no muestra las etiquetas anteriores" in text
    manifest = manifest_of(project, ADJUDICATION_MANIFEST_FILE_NAME)
    assert manifest.counts["disputed"] == 9 and manifest.counts["disputed_binary"] == 4
    assert manifest.counts["texts_checked"] > 0


def test_with_four_annotators_the_two_others_share_the_adjudication(project4: Path) -> None:
    second = disputes(project4, binary=6, levels_only=8)
    label_main(project4, second)
    assert run(project4, "import-main") == 0

    assert run(project4, "adjudication") == 0

    layout = read(LayoutRow, project4, RECORDS / ADJUDICATION_LAYOUT_FILE_NAME)
    assignment = read(AssignmentRow, project4, RECORDS / "assignment.parquet")
    disputed = [p for p in assignment if p.pair_id in second]
    expected = assign_adjudicators(disputed, 4, seed=SEED)
    assert {row.pair_id: row.annotator for row in layout} == expected
    loads = Counter(expected.values())
    assert max(loads.values()) - min(loads.values()) <= 1


def test_the_adjudicators_follow_the_documented_rotation() -> None:
    groups = [(1, 2), (1, 3), (1, 4), (2, 3), (2, 4), (3, 4)] * 4
    pairs = [
        AssignmentRow(
            pair_id=f"L-{i:08x}",
            code=f"GH{200 + i}",
            pool="random",
            source=None,
            group=1 + i % 6,
            annotator_a=a,
            annotator_b=b,
        )
        for i, (a, b) in enumerate(groups)
    ]
    order = sorted(p.pair_id for p in pairs)
    random.Random("7:adjudication").shuffle(order)
    by_id = {p.pair_id: p for p in pairs}
    load: Counter[int] = Counter()
    expected = {}
    for identifier in order:
        p = by_id[identifier]
        k = min(
            (k for k in (1, 2, 3, 4) if k not in (p.annotator_a, p.annotator_b)),
            key=lambda k: (load[k], k),
        )
        load[k] += 1
        expected[identifier] = k

    assert assign_adjudicators(pairs, 4, seed=7) == expected
    three = [p for p in pairs if 4 not in (p.annotator_a, p.annotator_b)]
    for identifier, k in assign_adjudicators(three, 3, seed=7).items():
        assert {k, by_id[identifier].annotator_a, by_id[identifier].annotator_b} == {1, 2, 3}


def test_the_adjudication_waits_for_disputes(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert run(project, "adjudication") == 1
    assert "the annotators' labels of snapshot" in capsys.readouterr().err
    label_main(project)
    assert run(project, "import-main") == 0
    capsys.readouterr()

    assert run(project, "adjudication") == 1
    assert "no pair is disputed, so there is nothing to adjudicate" in capsys.readouterr().err
    assert not list((project / WORKBOOKS).glob("desempate_*.xlsx"))


def adjudicated(
    project: Path, second: Mapping[str, str], third: Callable[[str], str | None]
) -> None:
    label_main(project, second)
    assert run(project, "import-main") == 0
    assert run(project, "adjudication") == 0
    fill(
        project, ADJUDICATION_LAYOUT_FILE_NAME, adjudicator_file, lambda k, row: third(row.pair_id)
    )


def test_the_consensus_takes_the_majority_of_the_three_labels(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    second = disputes(project, binary=3, levels_only=3)
    pairs = sorted(second)
    # The third annotator sides with the first on two pairs, with the second on two, and gives a
    # third level on the last two, which the median settles.
    third_levels = {}
    for index, pair in enumerate(pairs):
        first, other = truth(pair), second[pair]
        third_levels[pair] = [
            first,
            first,
            other,
            other,
            *(level for level in LEVELS if level not in (first, other)),
        ][min(index, 4)]
    adjudicated(project, second, third_levels.get)
    capsys.readouterr()

    assert run(project, "import-adjudication") == 0

    consensus = read(ConsensusLabel, project, LABELS / "consensus.parquet")
    labeled = [
        p.pair_id
        for p in read(files.PairRow, project, RECORDS / "pairs.parquet")
        if p.phase == "labeled"
    ]
    assert [row.pair_id for row in consensus] == labeled
    for row in consensus:
        if row.pair_id in second:
            assert row.source == "adjudication"
            expected = final_level(
                (truth(row.pair_id), second[row.pair_id], third_levels[row.pair_id])
            )
            assert row.label_3 == expected
        else:
            assert (row.source, row.label_3) == ("agreement", truth(row.pair_id))
        assert row.label == int(row.label_3 == "Redundante")
    third = read(AnnotatorLabelRow, project, LABELS / "adjudication_labels.parquet")
    assert {(r.pair_id, r.label_3) for r in third} == set(third_levels.items())
    manifest = manifest_of(project, ADJUDICATION_LABELS_MANIFEST_FILE_NAME)
    assert manifest.counts["adjudicated"] == 6 and manifest.counts["consensus_rows"] == 120
    without = sum(len({truth(p), second[p], third_levels[p]}) == 3 for p in pairs)
    assert manifest.counts["without_majority"] == without >= 1
    out = capsys.readouterr().out
    assert "Consensus labels: 120 pairs" in out and "114 agreed, 6 adjudicated" in out
    text = (project / RESULTS / ADJUDICATION_LABELS_MANIFEST_FILE_NAME).read_text("utf-8")
    for secret in secrets_of(project):
        assert secret not in text


def test_an_incomplete_adjudication_workbook_is_refused_and_no_consensus_is_written(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    second = disputes(project)
    skipped = sorted(second)[0]
    adjudicated(project, second, lambda pair: None if pair == skipped else truth(pair))
    capsys.readouterr()

    assert run(project, "import-adjudication") == 1

    assert "desempate_anotador_" in capsys.readouterr().err
    assert not (project / LABELS / "consensus.parquet").exists()
    assert not (project / LABELS / "adjudication_labels.parquet").exists()


def test_an_adjudicators_note_with_a_glued_name_is_refused_without_echo(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    second = disputes(project)
    adjudicated(project, second, truth)
    layout = read(LayoutRow, project, RECORDS / ADJUDICATION_LAYOUT_FILE_NAME)
    path = project / WORKBOOKS / adjudicator_file(layout[0].annotator)
    workbook = openpyxl.load_workbook(path)
    workbook[LABELING_SHEET].cell(
        2 + 4 * (layout[0].n - 1), 7
    ).value = "SegúnSupuestoapellido Imaginario."
    workbook.save(path)
    capsys.readouterr()

    assert run(project, "import-adjudication") == 1

    err = capsys.readouterr().err
    assert f"the 'Notas (opcional)' of pair {layout[0].code} holds a person's name" in err
    assert "supuestoapellido" not in err.casefold()
    assert not (project / LABELS / "consensus.parquet").exists()


def test_the_adjudication_labels_are_imported_once(
    project: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    second = disputes(project)
    adjudicated(project, second, truth)
    assert run(project, "import-adjudication") == 0
    before = files_under(project)
    capsys.readouterr()

    assert run(project, "import-adjudication") == 1

    assert "already imported" in capsys.readouterr().err
    assert files_under(project) == before


def test_the_help_names_every_step(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit):
        main(["--help"])

    out = " ".join(capsys.readouterr().out.split())
    for step in (
        "pilot",
        "import-pilot",
        "main",
        "import-main",
        "adjudication",
        "import-adjudication",
    ):
        assert step in out
    with pytest.raises(SystemExit):
        main(["import-main", "--help"])
    out = " ".join(capsys.readouterr().out.split())
    for text in (
        "incomplete workbook is refused",
        "labels/labels.parquet",
        "target 0.6",
        "--dry-run",
    ):
        assert text in out


def test_the_manifests_are_plain_json_of_aggregates(project: Path) -> None:
    label_main(project, disputes(project))
    assert run(project, "import-main") == 0

    data = json.loads((project / RESULTS / MAIN_LABELS_MANIFEST_FILE_NAME).read_text("utf-8"))

    assert set(data["details"]) == {
        "kappa_target",
        "meets_target",
        "fleiss_binary",
        "fleiss_levels",
        "cohen",
        "combinations",
        "levels_by_annotator",
    }
    assert data["details"]["levels_by_annotator"].keys() == {"1", "2", "3"}
