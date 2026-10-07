"""Behavior of the objectives verification sample: drawing it, its workbook, and the import.

Every record here is synthetic: made-up uuids, handles, titles, objectives and names.
"""

import io
import json
import math
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
import pytest

from thematic_redundancy.corpus.snapshot import SnapshotRecord
from thematic_redundancy.extraction.objectives import ObjectivesRecord
from thematic_redundancy.labeling.objectives_check import (
    COLUMNS,
    GENERAL_VERDICTS,
    SPECIFIC_VERDICTS,
    WRONG_FILE_HANDLES,
    WorkbookValidationError,
    import_verification,
    main,
    make_sample,
    make_verification_workbook,
    verification_report,
    wilson_interval,
)
from thematic_redundancy.shared.config import AppConfig, load_config

DEFAULT_CONFIG_PATH = Path(__file__).resolve().parents[3] / "config" / "default.yaml"
DECLARED = load_config(DEFAULT_CONFIG_PATH)
PROGRAMS = tuple(program.key for program in DECLARED.snapshot.programs)
SNAPSHOT_ID = "20261002T150405Z"
HARVEST_TIME = datetime(2026, 10, 2, 15, 4, 5, tzinfo=UTC)
NOW = datetime(2026, 10, 7, 12, 0, 0, tzinfo=UTC)
AUTHOR = "Inventadoapellido Ficticiosegundo, Personaejemplo Sintética"
ADVISOR = "Supuestoapellido Imaginario, Asesorejemplo"
PAGES_SHA256 = "a" * 64

GENERAL_LABEL = "Veredicto objetivo general"
SPECIFIC_LABEL = "Veredicto objetivos específicos"


def item_uuid(number: int) -> UUID:
    return UUID(f"00000000-0000-4000-8000-{number:012d}")


def handle(number: int) -> str:
    return f"123456789/{1000 + number}"


def objectives_record(
    number: int, program: str, status: str = "extracted", *, general: str | None = None
) -> ObjectivesRecord:
    if status == "not_found":
        return ObjectivesRecord(
            item_uuid=item_uuid(number),
            handle=handle(number),
            program=program,
            status="not_found",
            objective_general=None,
            objectives_specific=None,
            page_start=None,
            page_end=None,
            pattern=None,
            flags=(),
            general_chars=0,
            specific_chars=0,
            pages=120,
            pages_sha256=PAGES_SHA256,
        )
    general = general or f"Determinar el efecto sintético número {number} en la planta de prueba."
    specific = f"• Medir la variable {number}.\n• Comparar los resultados {number}."
    return ObjectivesRecord(
        item_uuid=item_uuid(number),
        handle=handle(number),
        program=program,
        status="extracted",
        objective_general=general,
        objectives_specific=specific,
        page_start=19,
        page_end=20,
        pattern="objetivo_general",
        flags=(),
        general_chars=len(general),
        specific_chars=len(specific),
        pages=120,
        pages_sha256=PAGES_SHA256,
    )


def snapshot_record(
    number: int,
    program: str,
    *,
    title: str | None = None,
    record_handle: str | None = None,
    author: str = AUTHOR,
) -> SnapshotRecord:
    metadata: dict[str, Any] = {
        "dc.title": [{"value": title or f"Título sintético {number}"}],
        "dc.description.abstract": [{"value": title or f"Resumen sintético {number}"}],
        "dc.contributor.author": [{"value": author}],
        "dc.contributor.advisor": [{"value": ADVISOR}],
    }
    return SnapshotRecord(
        uuid=item_uuid(number),
        handle=record_handle or handle(number),
        program_key=program,
        collection_uuid=next(
            p.collection_uuid for p in DECLARED.snapshot.programs if p.key == program
        ),
        harvested_at=HARVEST_TIME,
        metadata=metadata,
    )


def population(
    sizes: dict[str, tuple[int, int]],
) -> tuple[list[ObjectivesRecord], list[SnapshotRecord]]:
    """Return records for ``sizes``: per program, (extracted, not_found) counts."""
    objectives, records = [], []
    number = 1
    for program, (extracted, not_found) in sizes.items():
        for index in range(extracted + not_found):
            status = "extracted" if index < extracted else "not_found"
            objectives.append(objectives_record(number, program, status))
            records.append(snapshot_record(number, program))
            number += 1
    return objectives, records


FULL = {program: (40, 10) for program in PROGRAMS}


def sample_of(
    sizes: dict[str, tuple[int, int]] = FULL, *, seed: int = 7, per_program: int = 12
) -> Any:
    objectives, records = population(sizes)
    return make_sample(objectives, records, programs=PROGRAMS, per_program=per_program, seed=seed)


# Wilson interval


@pytest.mark.parametrize(("successes", "total"), [(54, 60), (1, 12), (30, 60), (59, 60)])
def test_wilson_bounds_solve_the_score_equation(successes: int, total: int) -> None:
    z = 1.959963984540054
    observed = successes / total

    low, high = wilson_interval(successes, total)

    for bound in (low, high):
        residual = (observed - bound) ** 2 - z**2 * bound * (1 - bound) / total
        assert math.isclose(residual, 0.0, abs_tol=1e-12)
    assert low < observed < high


def test_wilson_bounds_reach_the_ends_for_all_or_nothing() -> None:
    assert wilson_interval(0, 10)[0] == 0.0
    assert wilson_interval(10, 10)[1] == 1.0
    assert wilson_interval(54, 60) == pytest.approx((0.7985, 0.9534), abs=1e-4)


def test_wilson_interval_needs_at_least_one_row() -> None:
    with pytest.raises(ValueError, match="at least one"):
        wilson_interval(0, 0)


# Sampling


def test_the_sample_draws_the_configured_number_per_program() -> None:
    sample = sample_of()

    assert len(sample.items) == 60
    assert {program: sum(i.program == program for i in sample.items) for program in PROGRAMS} == (
        dict.fromkeys(PROGRAMS, 12)
    )
    assert [item.n for item in sample.items] == list(range(1, 61))


def test_the_sample_is_reproducible_under_a_fixed_seed() -> None:
    first, second = sample_of(seed=7), sample_of(seed=7)

    assert first.items == second.items
    assert [i.item_uuid for i in sample_of(seed=8).items] != [i.item_uuid for i in first.items]


def test_the_workbook_order_mixes_the_programs() -> None:
    programs = [item.program for item in sample_of().items]

    assert programs != sorted(programs, key=PROGRAMS.index)


def test_each_program_is_drawn_independently_of_the_others() -> None:
    larger = dict(FULL) | {"minas": (60, 10)}

    def drawn(sizes: dict[str, tuple[int, int]], program: str) -> set[UUID]:
        return {i.item_uuid for i in sample_of(sizes).items if i.program == program}

    assert drawn(FULL, "sistemas") == drawn(larger, "sistemas")


def test_not_found_documents_enter_in_proportion_within_each_program() -> None:
    sizes = {"sistemas": (100, 25), "industrial": (27, 3), "minas": (30, 0)}

    sample = sample_of(sizes)

    counts = {
        (program, status): sum(i.program == program and i.status == status for i in sample.items)
        for program in sizes
        for status in ("extracted", "not_found")
    }
    assert counts == {
        ("sistemas", "extracted"): 10,
        ("sistemas", "not_found"): 2,
        ("industrial", "extracted"): 11,
        ("industrial", "not_found"): 1,
        ("minas", "extracted"): 12,
        ("minas", "not_found"): 0,
    }
    assert sample.population["sistemas"] == {"extracted": 100, "not_found": 25}


def test_a_program_with_fewer_eligible_theses_gives_all_of_them() -> None:
    sample = sample_of({"electronica": (5, 2), "minas": (20, 0)})

    assert sum(i.program == "electronica" for i in sample.items) == 7


def test_the_wrong_file_item_is_never_sampled() -> None:
    objectives, records = population({"sistemas": (12, 0)})
    objectives.append(objectives_record(99, "sistemas"))
    wrong = next(iter(WRONG_FILE_HANDLES))
    records.append(snapshot_record(99, "sistemas", record_handle=wrong))
    objectives[-1] = objectives[-1].model_copy(update={"handle": wrong})

    sample = make_sample(objectives, records, programs=PROGRAMS, per_program=13, seed=1)

    assert item_uuid(99) not in {item.item_uuid for item in sample.items}
    assert sample.excluded_wrong_file == (wrong,)
    assert len(sample.items) == 12


def test_only_the_smaller_uuid_of_an_exact_duplicate_pair_is_eligible() -> None:
    objectives, records = population({"mecanica": (3, 0)})
    objectives += [objectives_record(51, "mecanica"), objectives_record(50, "mecanica")]
    # Same title and abstract, written with other spacing and case: one thesis (D24).
    records += [
        snapshot_record(51, "mecanica", title="Diseño  de un  equipo de prueba"),
        snapshot_record(50, "mecanica", title="diseño de un equipo de prueba"),
    ]

    sample = make_sample(objectives, records, programs=PROGRAMS, per_program=12, seed=1)

    sampled = {item.item_uuid for item in sample.items}
    assert item_uuid(50) in sampled and item_uuid(51) not in sampled
    assert sample.excluded_duplicates == (handle(51),)
    assert sample.population["mecanica"] == {"extracted": 4, "not_found": 0}


# Workbook


@pytest.fixture
def config() -> AppConfig:
    return DECLARED


def workbook_dir(project_root: Path) -> Path:
    return project_root / "data" / "labels" / "objectives_check" / SNAPSHOT_ID


def make(project_root: Path, sizes: dict[str, tuple[int, int]], **changes: Any) -> Any:
    objectives, records = population(sizes)
    for number, general in changes.get("generals", {}).items():
        objectives[number - 1] = objectives_record(
            number, objectives[number - 1].program, general=general
        )
    return make_verification_workbook(
        DECLARED,
        project_root,
        SNAPSHOT_ID,
        objectives,
        records,
        provenance={"locator_version": 1, "table_sha256": "b" * 64},
        clock=lambda: NOW,
    )


def check_sheet(path: Path) -> Any:
    return openpyxl.load_workbook(path)["Verificación"]


def fill(path: Path, verdicts: dict[str, tuple[str, str]] | None = None, **default: str) -> None:
    """Judge every row: by handle from ``verdicts``, otherwise by status."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Verificación"]
    columns = {name: index + 1 for index, name in enumerate(COLUMNS)}
    for row in range(2, sheet.max_row + 1):
        row_handle = sheet.cell(row, columns["Handle"]).value
        found = sheet.cell(row, columns["Estado"]).value == "extraído"
        general, specific = (verdicts or {}).get(
            row_handle,
            (
                default.get("general", "Correcto") if found else "No existe en la tesis",
                default.get("specific", "Correcto") if found else "No aplica",
            ),
        )
        sheet.cell(row, columns[GENERAL_LABEL]).value = general
        sheet.cell(row, columns[SPECIFIC_LABEL]).value = specific
    workbook.save(path)


def test_the_workbook_has_a_spanish_guide_and_one_row_per_sampled_thesis(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 1), "minas": (2, 0)})

    workbook = openpyxl.load_workbook(made.workbook_path)
    assert workbook.sheetnames == ["Guía", "Verificación"]
    guide = "\n".join(
        str(cell.value) for row in workbook["Guía"].iter_rows() for cell in row if cell.value
    )
    for label in [*GENERAL_VERDICTS, *SPECIFIC_VERDICTS, "4 a 6 minutos"]:
        assert label in guide
    sheet = workbook["Verificación"]
    assert [cell.value for cell in sheet[1]] == list(COLUMNS)
    assert sheet.max_row == 1 + 6
    assert sheet.freeze_panes == "A2"
    assert made.workbook_path == workbook_dir(tmp_path) / "objectives_check.xlsx"


def test_each_row_links_to_the_repository_item_and_to_the_local_pdf(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (1, 0)})

    sheet = check_sheet(made.workbook_path)
    columns = {name: index + 1 for index, name in enumerate(COLUMNS)}
    handle_cell = sheet.cell(2, columns["Handle"])
    pdf_cell = sheet.cell(2, columns["PDF local"])
    assert handle_cell.value == handle(1)
    assert handle_cell.hyperlink.target == f"https://repositorio.ucsm.edu.pe/handle/{handle(1)}"
    pdf = (workbook_dir(tmp_path) / pdf_cell.hyperlink.target).resolve()
    assert (
        pdf == (tmp_path / "data" / "raw" / SNAPSHOT_ID / "pdfs" / f"{item_uuid(1)}.pdf").resolve()
    )
    assert sheet.cell(2, columns["Páginas del PDF"]).value == "20–21"
    assert sheet.cell(2, columns["Programa"]).value == "Ingeniería de Sistemas"
    assert sheet.cell(2, columns["Estado"]).value == "extraído"
    assert sheet.cell(2, columns["Objetivo general extraído"]).value.startswith("Determinar")


def test_the_verdict_columns_offer_dropdown_lists(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (2, 1)})

    sheet = check_sheet(made.workbook_path)
    lists = {
        str(validation.sqref): validation.formula1
        for validation in sheet.data_validations.dataValidation
    }
    assert lists == {
        "I2:I4": '"' + ",".join(GENERAL_VERDICTS) + '"',
        "J2:J4": '"' + ",".join(SPECIFIC_VERDICTS) + '"',
    }


def test_the_workbook_shows_no_person_name_and_no_formula(tmp_path: Path) -> None:
    named = (
        "Aplicar el método de Personaejemplo Sintética Inventadoapellido Ficticiosegundo en la "
        "planta."
    )
    formula = '=HYPERLINK("https://example.org")'
    made = make(tmp_path, {"sistemas": (2, 0)}, generals={1: named, 2: formula})

    workbook = openpyxl.load_workbook(made.workbook_path)
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
    text = "\n".join(str(cell.value) for cell in cells if cell.value is not None).casefold()
    assert "inventadoapellido" not in text and "supuestoapellido" not in text
    assert "[nombre omitido]" in text
    assert made.redacted == 1
    assert all(cell.data_type != "f" for cell in cells)
    assert any(cell.value == formula for cell in cells)


def test_a_name_is_redacted_only_where_it_stands_as_whole_words(tmp_path: Path) -> None:
    # The author's family name also ends and begins longer words here; "ñ" is a letter, so
    # "ficticiosegundoña" is one word, which an ASCII-only word boundary would split.
    inside = (
        "Comparar la preinventadoapellido ficticiosegundo con la inventadoapellido "
        "ficticiosegundoña del ensayo."
    )
    standalone = "Aplicar el método de Inventadoapellido Ficticiosegundo en la planta."
    # A footnote number, plain or superscript, does not make the name part of a word.
    footnoted = (
        "Seguir a Inventadoapellido Ficticiosegundo1 y a Inventadoapellido Ficticiosegundo²."
    )
    made = make(
        tmp_path,
        {"sistemas": (3, 0)},
        generals={1: inside, 2: standalone, 3: footnoted},
    )

    sheet = check_sheet(made.workbook_path)
    handles, generals = (
        COLUMNS.index(name) + 1 for name in ("Handle", "Objetivo general extraído")
    )
    by_handle = {
        sheet.cell(row, handles).value: sheet.cell(row, generals).value for row in (2, 3, 4)
    }
    assert by_handle[handle(1)] == inside
    assert by_handle[handle(2)] == "Aplicar el método de [nombre omitido] en la planta."
    assert by_handle[handle(3)] == "Seguir a [nombre omitido]1 y a [nombre omitido]²."
    assert made.redacted == 2


def test_a_name_that_ends_with_an_initial_is_redacted(tmp_path: Path) -> None:
    # A plain word boundary after the final period would need a letter to follow it.
    author = "Inventadoapellido, Personaejemplo J."
    general = f"Aplicar el método de {author} en la planta."
    made = make_verification_workbook(
        DECLARED,
        tmp_path,
        SNAPSHOT_ID,
        [objectives_record(1, "sistemas", general=general)],
        [snapshot_record(1, "sistemas", author=author)],
        provenance={},
        clock=lambda: NOW,
    )

    cell = check_sheet(made.workbook_path).cell(2, COLUMNS.index("Objetivo general extraído") + 1)
    assert cell.value == "Aplicar el método de [nombre omitido] en la planta."
    assert made.redacted == 1


def test_the_sample_file_records_the_draw_without_text(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 1)})

    content = json.loads((workbook_dir(tmp_path) / "sample.json").read_text(encoding="utf-8"))
    assert content["seed"] == 20261007 and content["per_program"] == 12
    assert content["provenance"] == {"locator_version": 1, "table_sha256": "b" * 64}
    assert len(content["items"]) == 4
    assert "Determinar" not in json.dumps(content)
    assert made.sample.items[0].n == 1


def test_an_existing_workbook_is_never_overwritten(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (2, 0)})
    before = made.workbook_path.read_bytes()

    with pytest.raises(FileExistsError, match="already exists"):
        make(tmp_path, {"sistemas": (2, 0)})

    assert made.workbook_path.read_bytes() == before


# Import


def test_a_filled_workbook_gives_a_strict_accuracy_with_its_interval(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (9, 1), "minas": (10, 0)})
    fill(
        made.workbook_path,
        {
            handle(1): ("Parcial", "Correcto"),
            handle(2): ("Incorrecto", "Incorrecto"),
            handle(12): ("Correcto", "No aplica"),
        },
    )

    judgements = import_verification(made.workbook_path, made.sample)
    report = verification_report(judgements, made.sample, imported_at=NOW)

    general = report.general
    assert (general.judged, general.correct) == (20, 17)
    assert general.counts == {
        "correct": 17,
        "partial": 1,
        "incorrect": 1,
        "missed": 0,
        "absent": 1,
    }
    assert general.accuracy == pytest.approx(17 / 20)
    assert general.ci95 == pytest.approx(wilson_interval(17, 20))
    assert general.accuracy_with_true_absences == pytest.approx(18 / 20)
    assert report.by_program["sistemas"].accuracy == pytest.approx(7 / 10)
    assert report.by_program["minas"].accuracy == pytest.approx(10 / 10)
    assert report.specific.applicable == 18
    assert report.specific.correct == 17
    assert report.meets_target is False


def test_the_report_holds_numbers_only(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 0)})
    fill(made.workbook_path)

    report = verification_report(
        import_verification(made.workbook_path, made.sample), made.sample, imported_at=NOW
    )

    dumped = report.model_dump_json()
    assert "Determinar" not in dumped and handle(1) not in dumped
    assert report.general.accuracy == 1.0 and report.meets_target is True


@pytest.mark.parametrize(
    ("verdicts", "message"),
    [
        ({handle(1): (None, "Correcto")}, "row 2 .* has no 'Veredicto objetivo general'"),
        ({handle(1): ("Correcto", "")}, "has no 'Veredicto objetivos específicos'"),
        ({handle(1): ("Bien", "Correcto")}, "'Bien' is not an allowed"),
        ({handle(4): ("Correcto", "No aplica")}, "not allowed for a thesis whose status"),
        ({handle(4): ("No existe en la tesis", "Correcto")}, "must be 'No aplica'"),
    ],
    ids=["blank-general", "blank-specific", "unknown-value", "correct-not-found", "specific"],
)
def test_validation_refuses_blank_or_unknown_verdicts(
    tmp_path: Path, verdicts: dict[str, tuple[Any, Any]], message: str
) -> None:
    made = make(tmp_path, {"sistemas": (3, 1)})
    rows = {cell.value: cell.row for cell in check_sheet(made.workbook_path)["C"][1:]}
    fill(made.workbook_path, verdicts)
    message = message.replace("row 2", f"row {rows[next(iter(verdicts))]}")

    with pytest.raises(WorkbookValidationError, match=message):
        import_verification(made.workbook_path, made.sample)


def test_verdicts_are_read_without_regard_to_case_or_surrounding_spaces(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (2, 0)})
    fill(made.workbook_path, general="  correcto ", specific="NO APLICA")

    judgements = import_verification(made.workbook_path, made.sample)

    assert {(j.general, j.specific) for j in judgements} == {("correct", "not_applicable")}


def test_validation_lists_every_problem_at_once(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 0)})
    fill(made.workbook_path, general="", specific="")

    with pytest.raises(WorkbookValidationError) as caught:
        import_verification(made.workbook_path, made.sample)

    assert len(caught.value.problems) == 6


def test_validation_refuses_a_deleted_or_repeated_row(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 0)})
    fill(made.workbook_path)
    workbook = openpyxl.load_workbook(made.workbook_path)
    sheet = workbook["Verificación"]
    sheet.cell(3, COLUMNS.index("Handle") + 1).value = sheet.cell(
        2, COLUMNS.index("Handle") + 1
    ).value
    workbook.save(made.workbook_path)

    with pytest.raises(WorkbookValidationError) as caught:
        import_verification(made.workbook_path, made.sample)

    problems = "\n".join(caught.value.problems)
    assert "appears 2 times" in problems and "is missing" in problems


def edit_cell(path: Path, row: int, name: str, value: str) -> None:
    """Write ``value`` into column ``name`` of ``row`` of the check sheet."""
    workbook = openpyxl.load_workbook(path)
    workbook["Verificación"].cell(row, COLUMNS.index(name) + 1).value = value
    workbook.save(path)


def test_validation_refuses_a_row_whose_handle_is_not_in_the_sample(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 0)})
    fill(made.workbook_path)
    edit_cell(made.workbook_path, 2, "Handle", "123456789/9999")

    with pytest.raises(WorkbookValidationError) as caught:
        import_verification(made.workbook_path, made.sample)

    assert "row 2: handle '123456789/9999' is not in the sample" in caught.value.problems


def test_validation_refuses_a_row_whose_status_was_edited(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (3, 0)})
    fill(made.workbook_path)
    row_handle = check_sheet(made.workbook_path).cell(2, COLUMNS.index("Handle") + 1).value
    edit_cell(made.workbook_path, 2, "Estado", "no encontrado")

    with pytest.raises(WorkbookValidationError) as caught:
        import_verification(made.workbook_path, made.sample)

    assert caught.value.problems == (
        f"row 2 ({row_handle}): 'Estado' was changed; it must be 'extraído'",
    )


def test_validation_refuses_a_changed_header(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (1, 0)})
    workbook = openpyxl.load_workbook(made.workbook_path)
    workbook["Verificación"].cell(1, 9).value = "Veredicto"
    workbook.save(made.workbook_path)

    with pytest.raises(WorkbookValidationError, match="header row"):
        import_verification(made.workbook_path, made.sample)


def test_a_workbook_saved_under_an_unsupported_extension_is_refused(tmp_path: Path) -> None:
    made = make(tmp_path, {"sistemas": (1, 0)})
    fill(made.workbook_path)
    renamed = made.workbook_path.with_suffix(".xls")
    renamed.write_bytes(made.workbook_path.read_bytes())

    with pytest.raises(ValueError, match="is not a readable xlsx workbook"):
        import_verification(renamed, made.sample)


def project(tmp_path: Path) -> Path:
    """Return a project root that holds a copy of the declared configuration."""
    root = tmp_path / "project"
    (root / "config").mkdir(parents=True)
    (root / "config" / "default.yaml").write_bytes(DEFAULT_CONFIG_PATH.read_bytes())
    return root


def test_the_import_command_writes_the_report(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    root = project(tmp_path)
    made = make(root, {"sistemas": (4, 0)})
    fill(made.workbook_path)
    report_path = root / "results" / "objectives" / SNAPSHOT_ID / "verification.json"

    code = main(
        [
            "--project-root",
            str(root),
            "--snapshot-id",
            SNAPSHOT_ID,
            "--report-out",
            str(report_path),
        ]
    )

    assert code == 0
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["general"]["accuracy"] == 1.0
    assert "Accuracy" in capsys.readouterr().out


def test_the_import_command_reports_validation_problems(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    root = project(tmp_path)
    make(root, {"sistemas": (2, 0)})

    code = main(["--project-root", str(root), "--snapshot-id", SNAPSHOT_ID])

    assert code == 1
    assert "has no 'Veredicto objetivo general'" in capsys.readouterr().err


def rewrite_parts(path: Path, change: Callable[[str, bytes], bytes]) -> None:
    """Write the xlsx archive at ``path`` again, passing each part through ``change``."""
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            target.writestr(info, change(info.filename, source.read(info)))


def write_garbage(path: Path) -> None:
    path.write_bytes(b"synthetic bytes that are not a workbook")


def truncate_zip(path: Path) -> None:
    path.write_bytes(path.read_bytes()[:200])


def write_zip_without_workbook(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("note.txt", "synthetic")


def break_sheet_xml(path: Path) -> None:
    rewrite_parts(
        path,
        lambda name, data: data[: len(data) // 2] if name.startswith("xl/worksheets/") else data,
    )


def give_an_invalid_value(path: Path) -> None:
    rewrite_parts(
        path,
        lambda name, data: (
            data.replace(b'sheetId="1"', b'sheetId="x"') if name == "xl/workbook.xml" else data
        ),
    )


@pytest.mark.parametrize(
    "damage",
    [
        write_garbage,
        truncate_zip,
        write_zip_without_workbook,
        break_sheet_xml,
        give_an_invalid_value,
    ],
    ids=["garbage", "truncated-zip", "zip-without-workbook", "broken-xml", "invalid-value"],
)
def test_the_import_command_reports_an_unreadable_workbook_in_one_line(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], damage: Callable[[Path], None]
) -> None:
    root = project(tmp_path)
    made = make(root, {"sistemas": (2, 0)})
    fill(made.workbook_path)
    damage(made.workbook_path)

    code = main(["--project-root", str(root), "--snapshot-id", SNAPSHOT_ID])

    assert code == 1
    error = capsys.readouterr().err
    assert error.startswith("error: ") and error.count("\n") == 1
    assert "objectives_check.xlsx is not a readable xlsx workbook" in error
