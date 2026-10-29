"""The labeled set's workbooks (T20b) through their command line, on synthetic files that the
real loaders read: a snapshot, the fichas and a draw of T20a (``synthetic_files.py``), in a
project root under pytest's temporary directory.

Only git, the clocks and the random source of the blind codes are replaced; the configuration is
the real file.
"""

import hashlib
import itertools
import json
import random
import shutil
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import synthetic_files as files

import thematic_redundancy.labeling.pair_workbooks as command
from thematic_redundancy.labeling.pair_workbooks import (
    ASSIGNMENT_FILE_NAME,
    BLIND_CODES_FILE_NAME,
    MAIN_LAYOUT_FILE_NAME,
    MAIN_MANIFEST_FILE_NAME,
    PILOT_LABELS_FILE_NAME,
    PILOT_LABELS_MANIFEST_FILE_NAME,
    PILOT_LAYOUT_FILE_NAME,
    PILOT_MANIFEST_FILE_NAME,
    PILOT_WORKBOOK_FILE_NAME,
    RUN_COMMAND,
    AssignmentRow,
    BlindCodeRow,
    LayoutRow,
    PilotLabelRow,
    WorkbookManifest,
    load_workbook_inputs,
    main,
)
from thematic_redundancy.labeling.workbook import (
    CODE_PATTERN,
    EXAMPLES_SHEET,
    FIELDS,
    GUIDE_SHEET,
    LABELING_SHEET,
    NO_ABSTRACT,
    NO_KEYWORDS,
    NO_OBJECTIVES,
    ficha_view,
    shuffled,
    swapped,
)
from thematic_redundancy.shared.config import AppConfig, load_config
from thematic_redundancy.shared.filesystem import exclusive_run
from thematic_redundancy.shared.storage import ParquetTableStore
from thematic_redundancy.topic_map.run_experiment import GitState

SNAPSHOT = files.SNAPSHOT_ID
RECORDS = Path("data") / "labels" / "labeled_set" / SNAPSHOT
WORKBOOKS = RECORDS / "workbooks"
RESULTS = Path("results") / "labeled_set" / SNAPSHOT / "workbooks"
PILOT_WORKBOOK = WORKBOOKS / PILOT_WORKBOOK_FILE_NAME
STARTED = datetime(2026, 10, 27, 9, 0, tzinfo=UTC)
SEED = 20261021


def clean(root: Path, path: str) -> GitState:
    return GitState(commit="c" * 40, dirty=())


def clock() -> Any:
    moments = itertools.count()
    return lambda: datetime(2026, 10, 27, 9, 0, next(moments) % 60, tzinfo=UTC)


def timer() -> Any:
    ticks = itertools.count(10.0, 0.25)
    return lambda: next(ticks)


def config_of(root: Path) -> AppConfig:
    return load_config(root / "config" / "default.yaml")


@pytest.fixture
def project(project_root: Path) -> files.LabelingProject:
    return files.write_labeling_project(project_root, config_of(project_root))


def run(root: Path, step: str, *extra: str, **overrides: Any) -> int:
    arguments: dict[str, Any] = {
        "git": clean,
        "clock": clock(),
        "timer": timer(),
        "choice": random.Random(5).choice,
    }
    arguments.update(overrides)
    return main([step, "--project-root", str(root), "--snapshot-id", SNAPSHOT, *extra], **arguments)


def files_under(root: Path) -> set[str]:
    if not root.exists():
        return set()
    return {p.relative_to(root).as_posix() for p in root.rglob("*") if p.is_file()}


def written(root: Path) -> set[str]:
    """Return the files of the workbooks' steps, the draw's left out."""
    found = {f"results/{name}" for name in files_under(root / RESULTS)}
    records = files_under(root / RECORDS) - {"pairs.parquet"}
    return found | {f"data/{name}" for name in records}


def digests(root: Path) -> dict[str, str]:
    return {
        name: hashlib.sha256((root / name).read_bytes()).hexdigest() for name in files_under(root)
    }


def read(model: Any, root: Path, name: str) -> tuple[Any, ...]:
    return ParquetTableStore().read(model, root / RECORDS / name)


def manifest_of(root: Path, name: str) -> WorkbookManifest:
    return WorkbookManifest.model_validate_json((root / RESULTS / name).read_text("utf-8"))


def every_text(path: Path) -> str:
    workbook = openpyxl.load_workbook(path)
    return "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def set_annotators(root: Path, annotators: int) -> None:
    path = root / "config" / "default.yaml"
    text = path.read_text("utf-8")
    assert "annotators: null" in text
    path.write_text(text.replace("annotators: null", f"annotators: {annotators}"), "utf-8")


LEVELS_CYCLE = ("Redundante", "Relacionado", "No relacionado")


def fill_pilot(root: Path, labels: dict[int, tuple[str | None, str | None]] | None = None) -> None:
    """Label the pilot workbook's pairs, by ``N.º``: from ``labels``, otherwise in turn."""
    path = root / PILOT_WORKBOOK
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[LABELING_SHEET]
    for n in range(1, 11):
        label, note = (labels or {}).get(n, (LEVELS_CYCLE[n % 3], f"Justificación {n}."))
        sheet.cell(2 + 4 * (n - 1), 6).value = label
        sheet.cell(2 + 4 * (n - 1), 7).value = note
    workbook.save(path)


def through_the_pilot(root: Path, annotators: int | None = None) -> None:
    """Make the pilot workbook, label it and import it; then set the number of annotators."""
    assert run(root, "pilot") == 0
    fill_pilot(root)
    assert run(root, "import-pilot") == 0
    if annotators is not None:
        set_annotators(root, annotators)


# The pilot workbook


def test_the_pilot_writes_its_workbook_and_records_and_its_tracked_manifest_last(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    code = run(project.root, "pilot")

    out = capsys.readouterr().out
    assert code == 0
    assert written(project.root) == {
        f"data/workbooks/{PILOT_WORKBOOK_FILE_NAME}",
        f"data/{BLIND_CODES_FILE_NAME}",
        f"data/{PILOT_LAYOUT_FILE_NAME}",
        f"results/{PILOT_MANIFEST_FILE_NAME}",
        "results/.lock",
    }
    pilot = [p for p in project.pairs if p.phase == "pilot"]
    theses = len({code for p in pilot for code in (p.doc_code_a, p.doc_code_b)})
    for text in (
        f"Pilot workbook: 10 pairs over {theses} theses",
        "Name guard",
        PILOT_MANIFEST_FILE_NAME,
    ):
        assert text in out
    codes = {row.code for row in read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME)}
    for secret in (*project.fichas, *(p.pair_id for p in project.pairs), *codes):
        assert secret not in out


def test_the_pilot_shows_its_pairs_blind_in_the_recorded_order_and_sides(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")
    layout = read(LayoutRow, project.root, PILOT_LAYOUT_FILE_NAME)
    pilot = [p for p in project.pairs if p.phase == "pilot"]
    by_id = {p.pair_id: p for p in pilot}

    assert [row.pair_id for row in layout] == shuffled(by_id, seed=SEED, workbook="pilot")
    assert [row.n for row in layout] == list(range(1, 11))
    sheet = openpyxl.load_workbook(project.root / PILOT_WORKBOOK)[LABELING_SHEET]
    for row in layout:
        pair = by_id[row.pair_id]
        assert row.swapped == swapped(row.pair_id, seed=SEED)
        expected = (
            (pair.doc_code_b, pair.doc_code_a)
            if row.swapped
            else (pair.doc_code_a, pair.doc_code_b)
        )
        assert (row.doc_code_left, row.doc_code_right) == expected
        first = 2 + 4 * (row.n - 1)
        assert (sheet.cell(first, 1).value, sheet.cell(first, 2).value) == (row.n, row.code)
        left = ficha_view(project.fichas[row.doc_code_left]).texts
        right = ficha_view(project.fichas[row.doc_code_right]).texts
        for offset in range(len(FIELDS)):
            assert sheet.cell(first + offset, 4).value == left[offset]
            assert sheet.cell(first + offset, 5).value == right[offset]
    assert 0 < sum(row.swapped for row in layout) < 10


def test_the_pilot_shows_no_score_pool_link_doc_code_pair_id_or_person(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")

    workbook = openpyxl.load_workbook(project.root / PILOT_WORKBOOK)
    text = every_text(project.root / PILOT_WORKBOOK).casefold()
    hidden = [
        *project.fichas,
        *(p.pair_id for p in project.pairs),
        "sbert_high",
        "tfidf_high",
        "mid_range",
        "between",
        "minilm",
        "mpnet",
        "0.81",
        "0.995",
        "http",
        "handle",
        "inventadoapellido",
        "supuestoapellido",
        "0000-0002-1825-0097",
        "ingeniería",
    ]
    for item in hidden:
        assert item.casefold() not in text, item
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
    assert not any(cell.hyperlink for cell in cells)
    assert workbook.sheetnames == [GUIDE_SHEET, LABELING_SHEET]


def test_every_drawn_pair_gets_one_blind_code_that_is_not_its_pair_id(
    project: files.LabelingProject, tmp_path: Path
) -> None:
    run(project.root, "pilot", choice=__import__("secrets").choice)
    rows = read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME)
    codes = {row.pair_id: row.code for row in rows}

    assert set(codes) == {p.pair_id for p in project.pairs}
    assert len(set(codes.values())) == 130
    assert all(CODE_PATTERN.fullmatch(code) for code in codes.values())
    assert not any(code in pair_id or pair_id in code for pair_id, code in codes.items())
    layout = read(LayoutRow, project.root, PILOT_LAYOUT_FILE_NAME)
    assert all(row.code == codes[row.pair_id] for row in layout)
    # Another project with the same draw draws other codes: nothing derives them.
    other = tmp_path / "other"
    shutil.copytree(project.root, other, ignore=shutil.ignore_patterns("workbooks", "*.lock"))
    for name in (BLIND_CODES_FILE_NAME, PILOT_LAYOUT_FILE_NAME):
        (other / RECORDS / name).unlink()
    (other / RESULTS / PILOT_MANIFEST_FILE_NAME).unlink(missing_ok=True)
    assert run(other, "pilot", choice=__import__("secrets").choice) == 0
    again = {r.pair_id: r.code for r in read(BlindCodeRow, other, BLIND_CODES_FILE_NAME)}
    assert again != codes


def test_the_pilot_guide_names_the_scale_its_binary_label_and_the_blind_rules(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")

    guide = openpyxl.load_workbook(project.root / PILOT_WORKBOOK)[GUIDE_SHEET]
    lines = {
        guide.cell(row, 1).value: guide.cell(row, 2).value for row in range(1, guide.max_row + 1)
    }
    assert guide.cell(1, 1).value == "Etiquetado de pares de tesis: piloto"
    assert lines["Redundante"] == "Mismo problema y mismo enfoque."
    assert lines["Relacionado"] == "Mismo tema, distinto problema o enfoque."
    assert lines["No relacionado"] == "Distinto tema."
    assert "solo «Redundante» cuenta como par redundante" in lines["Etiqueta binaria"]
    assert "10 pares del piloto" in lines["Propósito"]
    assert "No busque las tesis" in lines["Solo este libro"]
    assert "no muestra autores, asesores, jurados, enlaces" in lines["Sin identidades"]
    for fixed in (NO_OBJECTIVES, NO_ABSTRACT, NO_KEYWORDS):
        assert fixed in lines["Campos ausentes"]
    assert "«N.º», «Código» y «Campo»" in lines["No modifique"]
    assert lines["4."].startswith("Etiquete los 10 pares y guarde el archivo")
    assert "exige al menos un par de cada nivel" in lines["4."]
    assert "vos" not in " ".join(str(text) for text in lines.values()).split()


def test_a_thesis_without_objectives_shows_the_fixed_text(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")
    layout = read(LayoutRow, project.root, PILOT_LAYOUT_FILE_NAME)
    sheet = openpyxl.load_workbook(project.root / PILOT_WORKBOOK)[LABELING_SHEET]
    missing = {f.doc_code for f in project.fichas.values() if f.objectives is None}

    shown = [
        sheet.cell(2 + 4 * (row.n - 1) + 2, column).value
        for row in layout
        for column, code in ((4, row.doc_code_left), (5, row.doc_code_right))
        if code in missing
    ]

    assert shown and set(shown) == {NO_OBJECTIVES}
    assert manifest_of(project.root, PILOT_MANIFEST_FILE_NAME).counts["without_objectives"] == len(
        {code for row in layout for code in (row.doc_code_left, row.doc_code_right)} & missing
    )


def with_objectives(project_root: Path, number: int, objectives: str) -> files.LabelingProject:
    config = config_of(project_root)
    theses = [
        files.Thesis(**{**t.__dict__, "objectives": objectives}) if t.number == number else t
        for t in files.labeling_theses(config)
    ]
    return files.write_labeling_project(project_root, config, theses)


def test_a_name_glued_to_a_word_stops_the_pilot_before_anything_is_written(
    project_root: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    # Thesis 6 is in the first pilot pair. The redactor replaces a name that stands as whole
    # words only, so it lets this one pass, glued to the word before it; the guard stops it.
    project = with_objectives(
        project_root, 6, "Seguir el método deInventadoapellido Ficticiosegundo en la planta."
    )

    code = run(project.root, "pilot")

    err = capsys.readouterr().err
    assert code == 1
    assert "the name guard found a person's name or an ORCID iD" in err
    assert f"{PILOT_WORKBOOK_FILE_NAME}, sheet 'Etiquetado', cell" in err and "(pair " in err
    assert "inventadoapellido" not in err.casefold() and "personaejemplo" not in err.casefold()
    assert written(project.root) == {"results/.lock"}


def test_a_name_that_stands_as_whole_words_is_redacted_and_counted(project_root: Path) -> None:
    project = with_objectives(
        project_root,
        6,
        "Seguir el método de Personaejemplo Sintética Inventadoapellido Ficticiosegundo.",
    )

    assert run(project.root, "pilot") == 0

    text = every_text(project.root / PILOT_WORKBOOK)
    assert "Seguir el método de [nombre omitido]." in text
    assert manifest_of(project.root, PILOT_MANIFEST_FILE_NAME).counts["redacted_cells"] == 1


def test_the_pilot_manifest_holds_counts_and_digests_and_no_code_id_or_thesis(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")
    manifest = manifest_of(project.root, PILOT_MANIFEST_FILE_NAME)
    text = (project.root / RESULTS / PILOT_MANIFEST_FILE_NAME).read_text("utf-8")
    codes = {row.code for row in read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME)}

    assert (manifest.step, manifest.snapshot_id, manifest.git_commit) == (
        "pilot",
        SNAPSHOT,
        "c" * 40,
    )
    assert manifest.settings == {
        "seed": SEED,
        "annotators": None,
        "records_dir": "data/labels/labeled_set",
    }
    assert manifest.counts["pairs"] == 10 and manifest.counts["drawn_pairs"] == 130
    assert manifest.counts["guard_strings"] > 0 and manifest.counts["texts_checked"] > 0
    assert manifest.data_files == {
        name: hashlib.sha256((project.root / RECORDS / name).read_bytes()).hexdigest()
        for name in (
            BLIND_CODES_FILE_NAME,
            PILOT_LAYOUT_FILE_NAME,
            f"workbooks/{PILOT_WORKBOOK_FILE_NAME}",
        )
    }
    assert set(manifest.inputs) == {
        "draw/manifest.json",
        "pairs.parquet",
        "fichas_manifest.json",
        "metadata.jsonl",
        "default.yaml",
    }
    for secret in (
        *project.fichas,
        *(p.pair_id for p in project.pairs),
        *codes,
        "Título sintético",
    ):
        assert secret not in text


@pytest.mark.parametrize(
    "left", [PILOT_LAYOUT_FILE_NAME, PILOT_WORKBOOK_FILE_NAME, PILOT_MANIFEST_FILE_NAME]
)
def test_an_existing_pilot_is_never_overwritten(
    project: files.LabelingProject, left: str, capsys: pytest.CaptureFixture[str]
) -> None:
    assert run(project.root, "pilot") == 0
    for name in {PILOT_LAYOUT_FILE_NAME, PILOT_WORKBOOK_FILE_NAME, PILOT_MANIFEST_FILE_NAME} - {
        left
    }:
        next(project.root.rglob(name)).unlink()
    before = digests(project.root)
    capsys.readouterr()

    code = run(project.root, "pilot")

    assert code == 1
    assert "already exists" in capsys.readouterr().err
    assert digests(project.root) == before


def test_the_codes_of_a_pilot_that_stopped_before_its_manifest_are_kept(
    project: files.LabelingProject,
) -> None:
    assert run(project.root, "pilot") == 0
    codes = read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME)
    for path in (
        project.root / RECORDS / PILOT_LAYOUT_FILE_NAME,
        project.root / PILOT_WORKBOOK,
        project.root / RESULTS / PILOT_MANIFEST_FILE_NAME,
    ):
        path.unlink()

    assert run(project.root, "pilot", choice=random.Random(99).choice) == 0

    assert read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME) == codes


def test_the_pilot_refuses_a_dirty_tree_and_another_run_that_holds_the_lock(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    seen: list[str] = []

    def dirty(root: Path, path: str) -> GitState:
        seen.append(path)
        return GitState(commit="c" * 40, dirty=("src/x.py",))

    assert run(project.root, "pilot", git=dirty) == 1
    assert "1 changed or untracked paths" in capsys.readouterr().err
    assert seen == [RESULTS.as_posix()]
    assert written(project.root) == set()
    (project.root / RESULTS).mkdir(parents=True)
    with exclusive_run(project.root / RESULTS, OSError("held")):
        assert run(project.root, "pilot") == 1
    assert "another run holds the lock" in capsys.readouterr().err
    assert written(project.root) == {"results/.lock"}


def test_the_pilot_dry_run_checks_everything_and_writes_nothing(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    before = digests(project.root)
    dirty = lambda root, path: GitState(commit="c" * 40, dirty=("odd/notes.md",))  # noqa: E731

    code = run(project.root, "pilot", "--dry-run", git=dirty)

    out = capsys.readouterr().out
    assert code == 0
    assert digests(project.root) == before
    for text in (
        "Dry run of the step 'pilot'",
        "nothing is written",
        "1 changed path",
        "Pilot workbook: 10 pairs",
        "Name guard",
        "Privacy guard",
        f"{RUN_COMMAND} pilot --snapshot-id {SNAPSHOT}",
    ):
        assert text in out


def test_the_dry_run_outside_a_repository_says_git_is_not_readable_here(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    from thematic_redundancy.topic_map.run_experiment import read_git_state

    assert run(project.root, "pilot", "--dry-run", git=read_git_state) == 0
    assert "Git: not readable here" in capsys.readouterr().out


def test_a_missing_draw_and_pairs_that_differ_from_the_draw_are_refused(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    pairs = project.root / RECORDS / "pairs.parquet"
    pairs.write_bytes(pairs.read_bytes() + b" ")

    assert run(project.root, "pilot") == 1
    assert "differs from what the draw's manifest.json records" in capsys.readouterr().err
    (project.root / "results" / "labeled_set" / SNAPSHOT / "draw" / "manifest.json").unlink()
    assert run(project.root, "pilot", "--dry-run") == 1
    assert "there is no draw of snapshot" in capsys.readouterr().err
    assert written(project.root) == {"results/.lock"}


def test_a_write_failure_tells_how_to_move_the_partial_records_away(
    project: files.LabelingProject,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    real = command.write_records

    def failing(*arguments: Any, **keywords: Any) -> Any:
        def write_text(path: Path, text: str) -> None:
            raise OSError("disk full")

        return real(*arguments, **keywords, write_text=write_text)

    monkeypatch.setattr(command, "write_records", failing)

    code = run(project.root, "pilot")

    last = capsys.readouterr().err.splitlines()[-1]
    assert code == 1
    assert last.startswith("error: disk full") and "move the partial records away" in last
    assert PILOT_LAYOUT_FILE_NAME in last and PILOT_WORKBOOK_FILE_NAME in last
    assert f"results/{PILOT_MANIFEST_FILE_NAME}" not in written(project.root)


def test_a_string_of_the_snapshot_in_the_manifest_refuses_every_write(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    # The records directory, which the manifest names, is taken for a string of the snapshot.
    def leaky(config: AppConfig, root: Path, snapshot_id: str) -> Any:
        inputs = load_workbook_inputs(config, root, snapshot_id)
        return command.WorkbookInputs(
            **{**inputs.__dict__, "sensitive": inputs.sensitive | {"data/labels/labeled_set"}}
        )

    code = run(project.root, "pilot", load=leaky)

    assert code == 1
    assert f"{PILOT_MANIFEST_FILE_NAME} (1)" in capsys.readouterr().err
    assert written(project.root) == {"results/.lock"}


def test_a_privacy_hit_on_the_final_manifest_says_what_was_written(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    # Only the final manifest names the records it vouches for, so only its check finds this.
    def leaky(config: AppConfig, root: Path, snapshot_id: str) -> Any:
        inputs = load_workbook_inputs(config, root, snapshot_id)
        return command.WorkbookInputs(
            **{**inputs.__dict__, "sensitive": inputs.sensitive | {PILOT_LAYOUT_FILE_NAME}}
        )

    code = run(project.root, "pilot", load=leaky)

    last = capsys.readouterr().err.splitlines()[-1]
    assert code == 1
    assert last.startswith(f"error: {PILOT_MANIFEST_FILE_NAME} would hold 1 string taken")
    assert "so it was not written, though the records before it were" in last
    assert "move the partial records away" in last and "nothing was written" not in last
    assert f"results/{PILOT_MANIFEST_FILE_NAME}" not in written(project.root)
    assert f"data/{PILOT_LAYOUT_FILE_NAME}" in written(project.root)


# The pilot's agreed labels


def test_the_import_stores_the_agreed_labels_and_justifications(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root, {1: ("  redundante ", "Mismo problema."), 2: ("Relacionado", None)})
    capsys.readouterr()

    code = run(project.root, "import-pilot")

    out = capsys.readouterr().out
    assert code == 0
    layout = {row.n: row for row in read(LayoutRow, project.root, PILOT_LAYOUT_FILE_NAME)}
    labels = read(PilotLabelRow, project.root, PILOT_LABELS_FILE_NAME)
    assert [row.n for row in labels] == list(range(1, 11))
    assert labels[0] == PilotLabelRow(
        n=1,
        code=layout[1].code,
        pair_id=layout[1].pair_id,
        level="redundant",
        binary=1,
        justification="Mismo problema.",
    )
    assert (labels[1].level, labels[1].binary, labels[1].justification) == ("related", 0, None)
    levels = Counter(row.level for row in labels)
    manifest = manifest_of(project.root, PILOT_LABELS_MANIFEST_FILE_NAME)
    assert manifest.details == {
        "levels": {
            "Redundante": levels["redundant"],
            "Relacionado": levels["related"],
            "No relacionado": levels["unrelated"],
        }
    }
    assert manifest.counts["binary_redundant"] == levels["redundant"]
    assert (
        manifest.inputs[PILOT_WORKBOOK_FILE_NAME]
        == hashlib.sha256((project.root / PILOT_WORKBOOK).read_bytes()).hexdigest()
    )
    assert manifest.data_files == {
        PILOT_LABELS_FILE_NAME: hashlib.sha256(
            (project.root / RECORDS / PILOT_LABELS_FILE_NAME).read_bytes()
        ).hexdigest()
    }
    assert "Pilot labels: 10 pairs" in out


@pytest.mark.parametrize(
    "justification",
    ["Lo aplica deInventadoapellido Ficticiosegundo en la planta.", "Ver 0000-0002-1825-0097."],
    ids=["glued-name", "orcid"],
)
def test_a_justification_with_a_name_the_redactor_leaves_is_refused_without_echo(
    project: files.LabelingProject, justification: str, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root, {2: ("Relacionado", justification)})
    code_of_2 = read(LayoutRow, project.root, PILOT_LAYOUT_FILE_NAME)[1].code
    capsys.readouterr()

    code = run(project.root, "import-pilot")

    err = capsys.readouterr().err
    assert code == 1
    assert (
        f"the 'Justificación' of pair {code_of_2} holds a person's name or an ORCID iD of the "
        "snapshot" in err
    )
    assert "inventadoapellido" not in err.casefold() and "0000-0002" not in err
    assert not (project.root / RECORDS / PILOT_LABELS_FILE_NAME).exists()


def test_a_justification_with_a_whole_word_name_is_stored_redacted(
    project: files.LabelingProject,
) -> None:
    run(project.root, "pilot")
    fill_pilot(
        project.root, {3: ("Relacionado", "Como propone Inventadoapellido Ficticiosegundo.")}
    )

    assert run(project.root, "import-pilot") == 0

    labels = read(PilotLabelRow, project.root, PILOT_LABELS_FILE_NAME)
    assert labels[2].justification == "Como propone [nombre omitido]."
    assert manifest_of(project.root, PILOT_LABELS_MANIFEST_FILE_NAME).counts["redacted_notes"] == 1


def test_the_import_lists_every_problem_and_stores_nothing(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root, {3: (None, "Sin acuerdo."), 4: ("Parecido", None)})
    capsys.readouterr()

    code = run(project.root, "import-pilot")

    err = capsys.readouterr().err
    assert code == 1
    assert "\nerror: 2 problem(s) in the workbook:\n" in err
    assert "has no 'Etiqueta acordada'" in err and "'Parecido' is not an allowed" in err
    assert not (project.root / RECORDS / PILOT_LABELS_FILE_NAME).exists()
    assert not (project.root / RESULTS / PILOT_LABELS_MANIFEST_FILE_NAME).exists()


def test_the_import_needs_a_pair_of_each_level(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root, {n: ("Relacionado", None) for n in range(1, 11)})
    capsys.readouterr()

    code = run(project.root, "import-pilot")

    err = capsys.readouterr().err
    assert code == 1
    assert "no pair is labeled 'Redundante'" in err and "no pair is labeled 'No relacionado'" in err
    assert "at least one pair of each level (D34 (7))" in err
    assert not (project.root / RECORDS / PILOT_LABELS_FILE_NAME).exists()


def test_a_damaged_pilot_workbook_is_refused_in_one_line(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root)
    path = project.root / PILOT_WORKBOOK
    path.write_bytes(path.read_bytes()[:300])
    capsys.readouterr()

    code = run(project.root, "import-pilot")

    err = capsys.readouterr().err
    assert code == 1
    last = err.splitlines()[-1]
    assert last.startswith("error: ") and err.count("error:") == 1
    assert last.endswith(")") and "etiquetado_piloto.xlsx is not a readable xlsx workbook" in last


def test_the_labels_are_imported_once_and_only_after_the_pilot(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    assert run(project.root, "import-pilot") == 1
    assert "the pilot workbook of snapshot" in capsys.readouterr().err
    through_the_pilot(project.root)
    before = digests(project.root)
    capsys.readouterr()

    assert run(project.root, "import-pilot") == 1
    assert "already stored" in capsys.readouterr().err
    assert digests(project.root) == before


def test_the_import_dry_run_reads_and_checks_the_workbook_and_writes_nothing(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    run(project.root, "pilot")
    fill_pilot(project.root)
    before = digests(project.root)
    capsys.readouterr()

    code = run(project.root, "import-pilot", "--dry-run")

    assert code == 0
    assert "Pilot labels: 10 pairs" in capsys.readouterr().out
    assert digests(project.root) == before


# The main workbooks


@pytest.mark.parametrize(
    ("annotators", "imported", "messages"),
    [
        (
            None,
            False,
            ("labeling.annotators is not set", "the pilot's agreed labels are not stored"),
        ),
        (3, False, ("the pilot's agreed labels are not stored",)),
        (None, True, ("labeling.annotators is not set",)),
    ],
)
def test_the_main_workbooks_wait_for_the_annotators_and_the_pilots_labels(
    project: files.LabelingProject,
    annotators: int | None,
    imported: bool,
    messages: tuple[str, ...],
    capsys: pytest.CaptureFixture[str],
) -> None:
    assert run(project.root, "pilot") == 0
    if imported:
        fill_pilot(project.root)
        assert run(project.root, "import-pilot") == 0
    if annotators is not None:
        set_annotators(project.root, annotators)
    before = digests(project.root)
    capsys.readouterr()

    for extra in ((), ("--dry-run",)):
        assert run(project.root, "main", *extra) == 1
        err = capsys.readouterr().err
        assert "the main workbooks wait" in err
        for message in messages:
            assert message in err
    assert digests(project.root) == before


@pytest.mark.parametrize(
    ("annotators", "groups", "per_group", "per_annotator"), [(3, 3, 40, 80), (4, 6, 20, 60)]
)
def test_each_annotator_gets_the_pairs_of_their_pairs_of_annotators(
    project: files.LabelingProject,
    annotators: int,
    groups: int,
    per_group: int,
    per_annotator: int,
    capsys: pytest.CaptureFixture[str],
) -> None:
    through_the_pilot(project.root, annotators)

    code = run(project.root, "main")

    assert code == 0
    assignment = read(AssignmentRow, project.root, ASSIGNMENT_FILE_NAME)
    labeled = {p.pair_id for p in project.pairs if p.phase == "labeled"}
    assert {row.pair_id for row in assignment} == labeled and len(assignment) == 120
    assert Counter(row.group for row in assignment) == {g: per_group for g in range(1, groups + 1)}
    layout = read(LayoutRow, project.root, MAIN_LAYOUT_FILE_NAME)
    for annotator in range(1, annotators + 1):
        mine = [row for row in layout if row.annotator == annotator]
        assert len(mine) == per_annotator
        expected = {
            row.pair_id for row in assignment if annotator in (row.annotator_a, row.annotator_b)
        }
        assert {row.pair_id for row in mine} == expected
        assert [row.pair_id for row in mine] == shuffled(
            expected, seed=SEED, workbook=f"annotator-{annotator}"
        )
        sheet = openpyxl.load_workbook(
            project.root / WORKBOOKS / f"etiquetado_anotador_{annotator}.xlsx"
        )[LABELING_SHEET]
        assert sheet.max_row == 1 + 4 * per_annotator
        assert [sheet.cell(2 + 4 * i, 2).value for i in range(per_annotator)] == [
            row.code for row in mine
        ]
    for pair_id, count in Counter(row.pair_id for row in layout).items():
        assert count == 2, pair_id
    manifest = manifest_of(project.root, MAIN_MANIFEST_FILE_NAME)
    assert manifest.counts["annotators"] == annotators and manifest.counts["pairs"] == 120
    assert [group["pairs"] for group in manifest.details["groups"]] == [per_group] * groups
    assert manifest.details["annotators"] == {
        str(k): per_annotator for k in range(1, annotators + 1)
    }
    assert "Main workbooks: " in capsys.readouterr().out


def test_each_main_workbook_shows_the_pilots_examples_with_their_agreed_labels(
    project: files.LabelingProject,
) -> None:
    through_the_pilot(project.root, 3)
    labels = {row.code: row for row in read(PilotLabelRow, project.root, PILOT_LABELS_FILE_NAME)}
    pilot = openpyxl.load_workbook(project.root / PILOT_WORKBOOK)[LABELING_SHEET]

    assert run(project.root, "main") == 0

    for annotator in (1, 2, 3):
        workbook = openpyxl.load_workbook(
            project.root / WORKBOOKS / f"etiquetado_anotador_{annotator}.xlsx"
        )
        assert workbook.sheetnames == [GUIDE_SHEET, EXAMPLES_SHEET, LABELING_SHEET]
        examples = workbook[EXAMPLES_SHEET]
        for row in range(2, 42):
            for column in (1, 2, 3, 4, 5):
                assert examples.cell(row, column).value == pilot.cell(row, column).value
        for first in range(2, 42, 4):
            agreed = labels[examples.cell(first, 2).value]
            assert (
                examples.cell(first, 6).value
                == {
                    "redundant": "Redundante",
                    "related": "Relacionado",
                    "unrelated": "No relacionado",
                }[agreed.level]
            )
            assert examples.cell(first, 7).value == agreed.justification
        assert not examples.data_validations.dataValidation
        guide = workbook[GUIDE_SHEET]
        assert guide.cell(1, 1).value == f"Etiquetado de pares de tesis: anotador {annotator}"
        texts = " ".join(str(guide.cell(r, 2).value) for r in range(1, guide.max_row + 1))
        assert "La hoja «Ejemplos» muestra los 10 pares del piloto" in texts
        assert "80 pares que le corresponden" in texts


def test_a_pair_keeps_its_code_and_sides_in_every_workbook(
    project: files.LabelingProject,
) -> None:
    through_the_pilot(project.root, 4)

    assert run(project.root, "main") == 0

    codes = {r.pair_id: r.code for r in read(BlindCodeRow, project.root, BLIND_CODES_FILE_NAME)}
    rows = read(LayoutRow, project.root, MAIN_LAYOUT_FILE_NAME)
    by_pair: dict[str, set[tuple[str, str, str]]] = {}
    for row in rows:
        by_pair.setdefault(row.pair_id, set()).add(
            (row.code, row.doc_code_left, row.doc_code_right)
        )
    assert all(len(shown) == 1 for shown in by_pair.values())
    assert all(next(iter(shown))[0] == codes[pair] for pair, shown in by_pair.items())


def test_the_main_workbooks_are_never_overwritten(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    through_the_pilot(project.root, 3)
    assert run(project.root, "main") == 0
    before = digests(project.root)
    capsys.readouterr()

    assert run(project.root, "main") == 1

    assert "already exist" in capsys.readouterr().err
    assert digests(project.root) == before


def test_the_main_dry_run_checks_everything_and_writes_nothing(
    project: files.LabelingProject, capsys: pytest.CaptureFixture[str]
) -> None:
    through_the_pilot(project.root, 3)
    before = digests(project.root)
    capsys.readouterr()

    code = run(project.root, "main", "--dry-run")

    out = capsys.readouterr().out
    assert code == 0
    assert digests(project.root) == before
    assert "Main workbooks: 3 annotators, 120 pairs over 3 pairs of annotators" in out
    assert "pairs per annotator: 1 80, 2 80, 3 80" in out


def test_a_glued_name_in_one_main_workbook_stops_them_all(
    project_root: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    config = config_of(project_root)
    pilot_theses = set()
    draw = files.labeling_draw(config, files.labeling_theses(config))
    for pair in draw[:10]:
        pilot_theses |= {pair.doc_code_a, pair.doc_code_b}
    outside = next(
        t.number
        for t in files.labeling_theses(config)
        if t.include and t.doc_code not in pilot_theses
    )
    project = with_objectives(project_root, outside, "Según lainventadoapellido ficticiosegundo.")
    through_the_pilot(project.root, 3)
    capsys.readouterr()

    code = run(project.root, "main")

    err = capsys.readouterr().err
    assert code == 1
    assert "the name guard found" in err and "etiquetado_anotador_" in err
    assert "inventadoapellido" not in err.casefold()
    assert not (project.root / WORKBOOKS / "etiquetado_anotador_1.xlsx").exists()
    assert not (project.root / RECORDS / ASSIGNMENT_FILE_NAME).exists()


# The real loader and the help


def test_load_workbook_inputs_reads_the_draw_the_fichas_and_the_snapshot(
    project: files.LabelingProject,
) -> None:
    inputs = load_workbook_inputs(config_of(project.root), project.root, SNAPSHOT)

    assert inputs.pairs == project.pairs
    assert dict(inputs.fichas) == project.fichas
    assert "inventadoapellido ficticiosegundo" in inputs.guard
    assert "0000-0002-1825-0097" in inputs.guard
    assert "Inventadoapellido Ficticiosegundo" in inputs.names
    assert not any("título sintético" in text for text in inputs.guard)
    assert any("título sintético" in text for text in inputs.sensitive)
    draw = project.root / "results" / "labeled_set" / SNAPSHOT / "draw" / "manifest.json"
    assert (
        inputs.input_sha256["draw/manifest.json"] == hashlib.sha256(draw.read_bytes()).hexdigest()
    )
    assert (
        inputs.input_sha256["pairs.parquet"]
        == json.loads(draw.read_text("utf-8"))["data_files"]["pairs.parquet"]
    )


def test_the_help_names_the_steps_and_where_they_write(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as stopped:
        main(["--help"])

    out = " ".join(capsys.readouterr().out.split())
    assert stopped.value.code == 0
    for text in (
        "pilot",
        "import-pilot",
        "main",
        "labeling.records_dir/<snapshot>/",
        "results/labeled_set/<snapshot>/workbooks/",
        "once",
    ):
        assert text in out
    with pytest.raises(SystemExit):
        main(["main", "--help"])
    out = " ".join(capsys.readouterr().out.split())
    for text in ("--dry-run", "--snapshot-id", "labeling.annotators (3 or 4)"):
        assert text in out
